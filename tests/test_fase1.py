import ctypes
import json
import os

import openpyxl
import pytest

import configuracion
import rutas_app
import app_automatizacion_optimizada as app


# ---------- configuración ----------

@pytest.fixture
def config_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(configuracion, "CONFIG_FILE", str(tmp_path / "config.json"))
    return tmp_path / "config.json"


def test_guardar_y_cargar_configuracion(config_tmp):
    configuracion.guardar_configuracion({"usuario": "ñandú", "num_hilos": 3})
    assert configuracion.cargar_configuracion() == {"usuario": "ñandú", "num_hilos": 3}
    assert not os.path.exists(str(config_tmp) + ".tmp")


def test_segunda_escritura_conserva_copia_bak(config_tmp):
    configuracion.guardar_configuracion({"v": 1})
    configuracion.guardar_configuracion({"v": 2})
    assert json.loads((config_tmp.parent / "config.json.bak").read_text(encoding="utf-8")) == {"v": 1}


def test_configuracion_corrupta_recupera_bak(config_tmp):
    configuracion.guardar_configuracion({"v": 1})
    configuracion.guardar_configuracion({"v": 2})
    config_tmp.write_text("{ esto no es json", encoding="utf-8")
    assert configuracion.cargar_configuracion() == {"v": 1}


def test_configuracion_inexistente_devuelve_vacio(config_tmp):
    assert configuracion.cargar_configuracion() == {}


# ---------- rutas ----------

def test_ruta_datos_migra_desde_directorio_de_trabajo(tmp_path, monkeypatch):
    carpeta_programa = tmp_path / "programa"
    carpeta_programa.mkdir()
    antiguo = tmp_path / "trabajo"
    antiguo.mkdir()
    (antiguo / "datos.json").write_text("{}", encoding="utf-8")
    monkeypatch.setattr(rutas_app, "CARPETA_DATOS", carpeta_programa)
    monkeypatch.chdir(antiguo)
    destino = rutas_app.ruta_datos("datos.json")
    assert destino == str(carpeta_programa / "datos.json")
    assert os.path.exists(destino)


def test_ruta_datos_no_pisa_lo_existente(tmp_path, monkeypatch):
    carpeta_programa = tmp_path / "programa"
    carpeta_programa.mkdir()
    (carpeta_programa / "datos.json").write_text('{"nuevo": 1}', encoding="utf-8")
    antiguo = tmp_path / "trabajo"
    antiguo.mkdir()
    (antiguo / "datos.json").write_text('{"viejo": 1}', encoding="utf-8")
    monkeypatch.setattr(rutas_app, "CARPETA_DATOS", carpeta_programa)
    monkeypatch.chdir(antiguo)
    rutas_app.ruta_datos("datos.json")
    assert json.loads((carpeta_programa / "datos.json").read_text(encoding="utf-8")) == {"nuevo": 1}


# ---------- avisos ----------

def crear_hoja_pedido(ruta, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PSA"
    ws["B2"], ws["E2"] = "PEDIDO Nº", "6003"
    ws.append([])
    ws.append([None, "REFERENCIAS A PEDIR", None, None, "ESTADO", "SUSTITUCIÓN", "CANTIDAD", "CLIENTE DMS"])
    for fila in filas:
        ws.append(fila)
    wb.save(ruta)


def test_extraer_avisos_agrupa_por_cliente(tmp_path):
    ruta = tmp_path / "23-09-2026.xlsx"
    crear_hoja_pedido(ruta, [
        [None, "A1", None, None, "NO SUMINISTRA", None, 1, 100],
        [None, "A2", None, None, " no  suministra ", None, 1, 100.0],
        [None, "B1", None, None, None, "B1X", 2, 200],
        [None, "C1", None, None, "OK", None, 1, 300],
        [None, "D1", None, None, "NO SUMINISTRA", "D1X", 1, 400],
        [None, "E1", None, None, None, "E1X", 1, None],
    ])
    no_suministra, sustituciones = app.extraer_avisos(str(ruta))
    assert no_suministra == {"100": ["A1", "A2"], "400": ["D1"]}
    assert sustituciones == {"200": [("B1", "B1X")]}


def test_extraer_avisos_no_recorre_millones_de_filas(tmp_path):
    import time
    ruta = tmp_path / "grande.xlsx"
    crear_hoja_pedido(ruta, [[None, "A1", None, None, "NO SUMINISTRA", None, 1, 100]])
    wb = openpyxl.load_workbook(ruta)
    wb.active["J1048576"] = "OK"  # dimensión declarada de más de un millón de filas
    wb.save(ruta)
    inicio = time.time()
    no_suministra, _ = app.extraer_avisos(str(ruta))
    assert no_suministra == {"100": ["A1"]}
    assert time.time() - inicio < 10


def test_cuerpos_de_los_avisos():
    assert "no suministra." in app.construir_cuerpo_aviso(["A1"], "23-09-2026")
    assert "- A1\n- A2" in app.construir_cuerpo_aviso(["A1", "A2"], "23-09-2026")
    assert "ha sido sustituida por la referencia X2." in app.construir_cuerpo_sustitucion([("X1", "X2")], "23-09-2026")
    assert "- X1 → X2\n- Y1 → Y2" in app.construir_cuerpo_sustitucion([("X1", "X2"), ("Y1", "Y2")], "23-09-2026")


# ---------- exclusión mutua y archivos en uso ----------

def test_tarea_activa_cubre_todas_las_tareas(monkeypatch):
    assert app.tarea_activa() is None
    for bandera in ("automation_running", "filtro_rotacion_running", "pedidos_running", "hoja_pedidos_running",
                    "avisos_running", "comparador_running", "mopar_running", "stock_fiat_running", "madd_running"):
        monkeypatch.setattr(app, bandera, True)
        assert app.tarea_activa() is not None
        monkeypatch.setattr(app, bandera, False)


def test_archivo_bloqueado_detecta_excel_abierto(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    ruta.write_bytes(b"x")
    assert app.archivo_bloqueado(str(ruta)) is False
    assert app.archivo_bloqueado(str(tmp_path / "no_existe.xlsx")) is False
    kernel32 = ctypes.windll.kernel32
    kernel32.CreateFileW.restype = ctypes.c_void_p
    GENERIC_READ_WRITE = 0x80000000 | 0x40000000
    FILE_SHARE_READ = 0x1
    OPEN_EXISTING = 3
    manejador = kernel32.CreateFileW(str(ruta), GENERIC_READ_WRITE, FILE_SHARE_READ, None, OPEN_EXISTING, 0x80, None)
    assert manejador not in (None, ctypes.c_void_p(-1).value)
    try:
        assert app.archivo_bloqueado(str(ruta)) is True
    finally:
        kernel32.CloseHandle(ctypes.c_void_p(manejador))
    assert app.archivo_bloqueado(str(ruta)) is False


def test_guardar_resultado_excel_escribe_columna_18(tmp_path):
    ruta = tmp_path / "quiter.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["h"] * 18)
    wb.active.append([None, "REF"] + [None] * 16)
    wb.save(ruta)
    app.guardar_resultado_excel(str(ruta), 2, "Pieza disponible")
    assert openpyxl.load_workbook(ruta).active.cell(row=2, column=18).value == "Pieza disponible"


# ---------- modo carpeta ----------

def test_modo_carpeta_guarda_los_datos_un_nivel_por_encima(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    carpeta_exe = dist / "app"
    (carpeta_exe / "_internal").mkdir(parents=True)
    monkeypatch.setattr(rutas_app.sys, "frozen", True, raising=False)
    monkeypatch.setattr(rutas_app.sys, "executable", str(carpeta_exe / "app.exe"))
    monkeypatch.setattr(rutas_app.sys, "_MEIPASS", str(carpeta_exe / "_internal"), raising=False)
    assert rutas_app.es_modo_carpeta() is True
    assert rutas_app.carpeta_programa() == dist.resolve()


def test_un_solo_exe_guarda_los_datos_junto_al_ejecutable(tmp_path, monkeypatch):
    monkeypatch.setattr(rutas_app.sys, "frozen", True, raising=False)
    monkeypatch.setattr(rutas_app.sys, "executable", str(tmp_path / "dist" / "app.exe"))
    monkeypatch.setattr(rutas_app.sys, "_MEIPASS", str(tmp_path / "temp" / "_MEI123"), raising=False)
    assert rutas_app.es_modo_carpeta() is False
    assert rutas_app.carpeta_programa() == (tmp_path / "dist").resolve()
