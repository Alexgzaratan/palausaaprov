import os
import time

import openpyxl
import pandas as pd
import pytest

import app_automatizacion_optimizada as app_mod


def ejecutar_hasta(app, condicion, limite=30):
    """Ejecuta el bucle de eventos real (los hilos usan after) hasta que se cumple la condición."""
    fin = time.time() + limite

    def revisar():
        if condicion() or time.time() > fin:
            app.quit()
        else:
            app.after(100, revisar)

    app.after(100, revisar)
    app.mainloop()


@pytest.fixture
def escenario(app, tmp_path, monkeypatch):
    quiter = tmp_path / "quiter.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["h"] * 18)
    for i, precio in enumerate([200, 100, 300, 50, 150, 400]):
        fila = [None] * 18
        fila[1], fila[5] = f"REF{i}", precio
        wb.active.append(fila)
    wb.save(quiter)
    propuesta = tmp_path / "propuesta.xlsx"
    pd.DataFrame({"Articulo": [f"REF{i}" for i in range(6)], "Media": [0.0] * 6}).to_excel(propuesta, sheet_name="Informe", index=False)

    app.config["usuario"], app.config["password"] = "u", "x"  # viven en la pestaña Configuración, aún sin construir
    app.construir_pestana("Automatización")
    for entrada, valor in ((app.excel_auto_entry, quiter), (app.excel_rot_entry, propuesta), (app.folder_rot_entry, tmp_path),
                           (app.precio_entry, "150"), (app.umbral_entry, "0.12")):
        entrada.delete(0, "end")
        entrada.insert(0, str(valor))
    preguntas = []
    monkeypatch.setattr(app_mod.messagebox, "askyesno", lambda *a, **k: preguntas.append(a[1]) or app.respuesta_si)
    app.respuesta_si = True
    app.preguntas = preguntas
    app.salida_rotacion = tmp_path / "quiter_ROTACION.xlsx"
    app.capturas = []
    monkeypatch.setattr(app, "mostrar_captura", lambda *datos: app.capturas.append(datos))
    return app


def worker_completo(worker_id, sub_filas, ruta_excel, usuario, password, log_callback, update_progress, stop_event):
    for fila_idx, referencia in sub_filas:
        app_mod.guardar_resultado_excel(ruta_excel, fila_idx, "Pieza disponible")
        update_progress()


def worker_que_deja_filas_sin_procesar(worker_id, sub_filas, ruta_excel, usuario, password, log_callback, update_progress, stop_event):
    if worker_id == 1:
        return
    worker_completo(worker_id, sub_filas, ruta_excel, usuario, password, log_callback, update_progress, stop_event)


def lanzar_flujo(app):
    app.iniciar_flujo_completo()
    assert app.btn_flujo.cget("state") == "disabled"
    ejecutar_hasta(app, lambda: app.btn_flujo.cget("state") == "normal")


def test_flujo_completo_lanza_la_rotacion_tras_la_automatizacion(escenario, monkeypatch):
    monkeypatch.setattr(app_mod, "worker_playwright", worker_completo)
    lanzar_flujo(escenario)
    assert escenario.salida_rotacion.exists()
    ws = openpyxl.load_workbook(escenario.salida_rotacion).active
    colores = [ws.cell(row=r, column=1).fill.start_color.rgb if ws.cell(row=r, column=1).fill.fill_type else None for r in range(1, 7)]
    assert colores == ["00FF0000", None, "00FF0000", None, "00FF0000", "00FF0000"]  # bajo 150 EUR queda sin marcar
    assert escenario.capturas == [(4, 6, "quiter")]  # «¡Capturado!» con las líneas marcadas
    assert ("ok", "Rotación completada: quiter_ROTACION.xlsx") in escenario.avisos.historial


def test_flujo_con_incidencias_pregunta_y_continua_si_se_acepta(escenario, monkeypatch):
    monkeypatch.setattr(app_mod, "worker_playwright", worker_que_deja_filas_sin_procesar)
    escenario.respuesta_si = True
    lanzar_flujo(escenario)
    assert any("incidencias" in p for p in escenario.preguntas[1:])
    assert escenario.salida_rotacion.exists()


def test_flujo_con_incidencias_no_lanza_la_rotacion_si_se_rechaza(escenario, monkeypatch):
    monkeypatch.setattr(app_mod, "worker_playwright", worker_que_deja_filas_sin_procesar)
    llamadas = []

    def responder(*a, **k):
        llamadas.append(a[1])
        return len(llamadas) == 1  # acepta el inicio, rechaza continuar tras las incidencias

    monkeypatch.setattr(app_mod.messagebox, "askyesno", responder)
    lanzar_flujo(escenario)
    assert len(llamadas) == 2
    assert not escenario.salida_rotacion.exists()
    assert escenario.capturas == []


def test_flujo_con_error_avisa_y_no_lanza_la_rotacion(escenario):
    ruta = escenario.excel_auto_entry.get()
    with open(ruta, "wb") as f:
        f.write(b"esto no es un Excel")  # la automatizacion falla al leer el fichero
    lanzar_flujo(escenario)
    assert not escenario.salida_rotacion.exists()
    assert escenario.mensajes  # se avisa con un mensaje visible, no solo en el registro


def test_hilo_que_falla_deja_el_flujo_como_incompleto(escenario, monkeypatch):
    def worker_roto(*args):
        raise RuntimeError("fallo simulado")

    monkeypatch.setattr(app_mod, "worker_playwright", worker_roto)
    preguntas = []
    monkeypatch.setattr(app_mod.messagebox, "askyesno", lambda *a, **k: preguntas.append(a[1]) or len(preguntas) == 1)
    lanzar_flujo(escenario)
    assert any("incidencias" in p for p in preguntas)
    assert not escenario.salida_rotacion.exists()


def test_rotacion_suelta_no_muestra_capturado(escenario):
    escenario.excel_quiter_entry.delete(0, "end")
    escenario.excel_quiter_entry.insert(0, escenario.excel_auto_entry.get())
    assert escenario.iniciar_filtro_rotacion()
    ejecutar_hasta(escenario, lambda: escenario.btn_rotacion.cget("state") == "normal")
    assert escenario.salida_rotacion.exists()
    assert escenario.capturas == []
