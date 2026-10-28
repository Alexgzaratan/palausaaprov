"""worker_playwright contra un portal falso en localhost que imita los pasos del de Citroën
(idioma, menús, formulario, resultado y «Otro Articulo"), con autenticación básica."""
import base64
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse

import openpyxl
import pytest

import app_automatizacion_optimizada as app_mod

USUARIO, CLAVE = "prueba", "clave"


class Portal:
    def __init__(self):
        self.visitas = {}          # peticiones con la contraseña correcta, por ruta
        self.rechazadas = 0        # peticiones sin contraseña o con una incorrecta (respuesta 401)
        self.fallos_acceso = 0     # nº de accesos que responden 503 antes de funcionar
        self.navegador_bloqueado = None  # nº de navegador (por orden de llegada) al que nunca se deja entrar
        self.navegadores = 0
        self.lentas = set()        # referencias cuyo resultado tarda 30 s
        self.rotas = set()         # referencias cuya página de resultado no trae la tabla
        portal = self

        class Manejador(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def responder(self, cuerpo, estado=200, cabeceras=()):
                datos = f"<html><body>{cuerpo}</body></html>".encode("utf-8")
                self.send_response(estado)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(datos)))
                for clave, valor in cabeceras:
                    self.send_header(clave, valor)
                self.end_headers()
                self.wfile.write(datos)

            def do_GET(self):
                url = urlparse(self.path)
                cookie = self.headers.get("Cookie") or ""
                cabeceras = []
                if "navegador=" not in cookie:  # cada navegador recibe su número en su primera petición
                    portal.navegadores += 1
                    cookie = f"navegador={portal.navegadores}"
                    cabeceras.append(("Set-Cookie", cookie))
                esperado = "Basic " + base64.b64encode(f"{USUARIO}:{CLAVE}".encode()).decode()
                if self.headers.get("Authorization") != esperado:
                    portal.rechazadas += 1
                    return self.responder("No autorizado", 401, cabeceras + [("WWW-Authenticate", 'Basic realm="portal"')])
                portal.visitas[url.path] = portal.visitas.get(url.path, 0) + 1
                if url.path == "/zzlogpl0":
                    if portal.navegador_bloqueado is not None and f"navegador={portal.navegador_bloqueado}" in cookie:
                        return self.responder("Servicio no disponible", 503)
                    if portal.fallos_acceso > 0:
                        portal.fallos_acceso -= 1
                        return self.responder("Servicio no disponible", 503)
                    return self.responder('<a href="/menu">Español</a>')
                if url.path == "/menu":
                    return self.responder('<a href="/submenu">Referencia</a>')
                if url.path == "/submenu":
                    return self.responder('<a href="/formulario">Mostrar disponibilidad pieza</a>')
                if url.path == "/formulario":
                    return self.responder(
                        '<form action="/resultado"><input type="text" name="mara-matnr[1]" value="">'
                        '<input type="radio" name="r" value="1" checked><input type="radio" name="r" value="2">'
                        '<a href="#" onclick="document.forms[0].submit()">Ejecutar programa</a></form>')
                if url.path == "/resultado":
                    ref = parse_qs(url.query).get("mara-matnr[1]", [""])[0]
                    if ref in portal.lentas:
                        time.sleep(30)
                    if ref in portal.rotas:
                        return self.responder("Error interno del portal")
                    estado = "Sin stock" if ref.endswith("X") else "Pieza disponible"
                    return self.responder(f'<table><tr class="moduloTableauBlanc"><td>{estado}</td></tr></table>'
                                          '<a href="/formulario">Otro Articulo</a>')
                self.responder("No encontrado", 404)

        self.servidor = ThreadingHTTPServer(("127.0.0.1", 0), Manejador)
        self.servidor.daemon_threads = True
        threading.Thread(target=self.servidor.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.servidor.server_port}/zzlogpl0"


@pytest.fixture
def portal(monkeypatch):
    portal = Portal()
    monkeypatch.setattr(app_mod, "URL_PORTAL", portal.url)
    monkeypatch.setattr(app_mod, "PAUSA_REINTENTO_ACCESO", 0.1)
    yield portal
    portal.servidor.shutdown()


def excel(tmp_path, referencias):
    ruta = tmp_path / "S99P.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["h"] * 18)
    for ref in referencias:
        fila = [None] * 18
        fila[1], fila[5] = ref, 500
        wb.active.append(fila)
    wb.save(ruta)
    return ruta


def lanzar(ruta, hilos=1, clave=CLAVE, stop_event=None, detener_a_los=None):
    registro = []
    stop_event = stop_event or threading.Event()
    if detener_a_los is not None:
        threading.Timer(detener_a_los, stop_event.set).start()
    inicio = time.monotonic()
    app_mod.ejecutar_automatizacion_task(str(ruta), USUARIO, clave, 0, hilos, log_callback=lambda t, tag=None: registro.append(t),
                                         stop_event=stop_event)
    return registro, time.monotonic() - inicio


def resultados(ruta):
    ws = openpyxl.load_workbook(ruta).active
    return [ws.cell(row=r, column=18).value for r in range(2, ws.max_row + 1)]


def test_consulta_todo_y_solo_pasa_por_los_menus_una_vez(portal, tmp_path):
    ruta = excel(tmp_path, ["A1", "B2X", "C3", "D4"])
    registro, _ = lanzar(ruta)
    assert resultados(ruta) == ["Pieza disponible", "Sin stock", "Pieza disponible", "Pieza disponible"]
    assert app_mod.auto_resultado == "completa"
    assert portal.visitas["/submenu"] == 1  # tras «Otro Articulo» el formulario ya está: no se repiten los menús
    assert portal.visitas["/resultado"] == 4


def test_contrasena_incorrecta_para_enseguida_con_un_mensaje_claro(portal, tmp_path):
    ruta = excel(tmp_path, ["A1", "B2", "C3"])
    registro, duracion = lanzar(ruta, hilos=3, clave="mal")
    assert app_mod.auto_resultado == "credenciales"
    assert any("rechazado el usuario o la contraseña" in t for t in registro)
    assert duracion < 15 and resultados(ruta) == [None, None, None]
    assert portal.rechazadas <= 6 and not portal.visitas  # un intento por hilo (Chromium pide dos veces), sin reintentos


def test_si_el_portal_falla_al_entrar_se_reintenta(portal, tmp_path):
    portal.fallos_acceso = 2
    ruta = excel(tmp_path, ["A1", "B2"])
    registro, _ = lanzar(ruta)
    assert resultados(ruta) == ["Pieza disponible", "Pieza disponible"]
    assert sum("intento 1 de 3" in t for t in registro) == 1 and sum("intento 2 de 3" in t for t in registro) == 1


def test_si_un_hilo_no_consigue_entrar_los_demas_hacen_su_parte(portal, tmp_path):
    portal.navegador_bloqueado = 1  # el primer navegador en llegar agota sus 3 intentos
    ruta = excel(tmp_path, [f"R{i}" for i in range(6)])
    registro, _ = lanzar(ruta, hilos=2)
    assert resultados(ruta) == ["Pieza disponible"] * 6
    assert app_mod.auto_resultado == "completa"
    assert any("Los demás hilos harán sus referencias" in t for t in registro)


def test_tras_dos_errores_seguidos_se_recarga_el_portal(portal, tmp_path, monkeypatch):
    portal.rotas = {"MAL1", "MAL2"}
    monkeypatch.setattr(app_mod, "ESCALA_ESPERAS", 0.15)  # una página rota hace esperar a los tiempos máximos
    ruta = excel(tmp_path, ["A1", "MAL1", "MAL2", "B2"])
    registro, _ = lanzar(ruta)
    assert resultados(ruta) == ["Pieza disponible", None, None, "Pieza disponible"]
    assert any("se vuelve a cargar el portal" in t for t in registro)
    assert portal.visitas["/zzlogpl0"] == 2
    assert app_mod.auto_resultado == "completa"  # las dos rotas cuentan como procesadas (con error en el registro)


def test_detener_responde_enseguida_aunque_el_portal_tarde(portal, tmp_path):
    portal.lentas = {"LENTA"}
    ruta = excel(tmp_path, ["LENTA", "B2", "C3"])
    registro, duracion = lanzar(ruta, detener_a_los=3)
    assert app_mod.auto_resultado == "detenida"
    assert duracion < 8  # antes esperaba hasta 15 s por la tabla de resultado
