import re

import openpyxl
import pandas as pd
import pytest

import app_automatizacion_optimizada as app_mod

# (referencia, pedidas, precio unitario, resultado, media)
FILAS = [
    ("CARA", 1, 1850.0, "Pieza disponible", 0.5),     # unidad > 1600 y quedaría en blanco -> azul
    ("CARA_ROJA", 1, 2400.0, "Pieza disponible", 0.0),  # baja rotación: manda el rojo
    ("CARA_AMARILLA", 1, 1700.0, "Sin stock", 0.5),     # no disponible: manda el amarillo
    ("MUCHAS", 20, 150.0, "Pieza disponible", 0.5),     # total 3000 pero la unidad es barata -> blanco
    ("JUSTO", 1, 1600.0, "Pieza disponible", 0.5),      # igual al límite: no lo supera -> blanco
    ("BARATA", 1, 90.0, "Pieza disponible", 0.5),       # total < precio mínimo -> sin marcar
]


@pytest.fixture
def archivos(tmp_path):
    quiter = tmp_path / "S40P.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Marca", "Referencia", "Descripción", "Pedidas", "P. compra", "Total"] + [None] * 11 + [None])
    for ref, pedidas, unitario, resultado, _ in FILAS:
        fila = [None] * 18
        fila[0], fila[1], fila[3], fila[4], fila[5], fila[17] = 2, ref, pedidas, unitario, round(pedidas * unitario, 2), resultado
        wb.active.append(fila)
    wb.save(quiter)
    propuesta = tmp_path / "S40PR.xlsx"
    pd.DataFrame({"Articulo": [f[0] for f in FILAS], "Media": [f[4] for f in FILAS]}).to_excel(propuesta, sheet_name="Informe", index=False)
    return quiter, propuesta, tmp_path


def filtrar(quiter, propuesta, carpeta, limite=1600.0):
    registro = []
    app_mod.filtrar_por_rotacion_task(str(quiter), str(propuesta), str(carpeta), 0.12, "FF0000", 150.0,
                                      log_callback=lambda t, tag=None: registro.append(t), limite_revisar=limite)
    return registro, carpeta / "S40P_ROTACION.xlsx"


def test_regla_de_piezas_caras():
    clasificar = app_mod.clasificar_fila_rotacion
    assert clasificar(1850, 0.5, "pieza disponible", 0.12, 150, 1850, 1600) == "revisar"
    assert clasificar(2400, 0.0, "pieza disponible", 0.12, 150, 2400, 1600) == "rojo"
    assert clasificar(1700, 0.5, "sin stock", 0.12, 150, 1700, 1600) == "amarillo"
    assert clasificar(3000, 0.5, "", 0.12, 150, 150, 1600) == "blanco"  # se mira la unidad, no el total
    assert clasificar(1850, 0.5, "", 0.12, 150, 1850, 0) == "blanco"    # 0 = no marcar
    assert clasificar(1850, 0.5, "", 0.12, 150, None, 1600) == "blanco"


def test_el_filtro_pinta_de_azul_solo_las_caras_que_quedarian_en_blanco(archivos):
    quiter, propuesta, carpeta = archivos
    registro, salida = filtrar(quiter, propuesta, carpeta)
    ws = openpyxl.load_workbook(salida).active
    colores = {ws.cell(row=f, column=2).value: (ws.cell(row=f, column=2).fill.start_color.rgb if ws.cell(row=f, column=2).fill.fill_type else None)
               for f in range(1, ws.max_row + 1)}
    assert colores == {"CARA": "009DC3E6", "CARA_ROJA": "00FF0000", "CARA_AMARILLA": "00FFFF00", "MUCHAS": None, "JUSTO": None, "BARATA": None}
    assert "Piezas en azul para revisar (más de 1600 EUR la unidad): 1." in " ".join(registro)
    resumen = next(t for t in registro if "Resumen rotacion" in t)
    assert re.search(r"rojas por umbral 1, amarillas por disponibilidad 1, .* EUR 1", resumen)


def test_sin_limite_no_hay_azules(archivos):
    quiter, propuesta, carpeta = archivos
    _, salida = filtrar(quiter, propuesta, carpeta, limite=0)
    assert app_mod.lineas_a_revisar_rotacion(str(salida)) == []


def test_la_vista_previa_cuenta_las_azules_igual_que_el_filtro(archivos):
    quiter, propuesta, carpeta = archivos
    previsto = app_mod.contar_rotacion(app_mod.leer_filas_quiter(str(quiter)), app_mod.leer_medias_propuesta(str(propuesta)), 0.12, 150.0, 1600.0)
    assert previsto == {"rojo": 1, "amarillo": 1, "revisar": 1, "precio": 1, "blanco": 2}


def test_las_azules_no_van_a_poner_a_cero_sino_aparte(archivos):
    quiter, propuesta, carpeta = archivos
    _, salida = filtrar(quiter, propuesta, carpeta)
    marcadas, total = app_mod.lineas_marcadas_rotacion(str(salida))
    assert [ref for _, ref, _ in marcadas] == ["CARA_ROJA", "CARA_AMARILLA"] and total == 6
    revisar = app_mod.lineas_a_revisar_rotacion(str(salida))
    assert revisar == [(1, "CARA", 1850.0)]
    texto = app_mod.texto_lista_quiter("S40P", marcadas, total, revisar)
    assert texto.splitlines() == ["rotación S40P: 2 de 6 líneas a poner a 0", "Lín. 2 · CARA_ROJA (rojo)", "Lín. 3 · CARA_AMARILLA (amarillo)",
                                  "Para decidir tú (piezas caras en azul, no poner a 0):", "Lín. 1 · CARA (1.850,00 €/ud)"]
    assert app_mod.texto_lista_quiter("S40P", [], 6, revisar).splitlines()[0] == "rotación S40P: ninguna línea marcada de 6"


def test_formato_euros():
    assert app_mod.euros(1850) == "1.850,00 €"
    assert app_mod.euros(12345.5) == "12.345,50 €"
    assert app_mod.euros(99.9) == "99,90 €"


@pytest.mark.parametrize("escrito,valor", [("1600", 1600.0), ("1.600", 1600.0), ("1.600,50", 1600.5), ("1600,5", 1600.5),
                                           ("2000 €", 2000.0), ("", 0.0), ("0", 0.0), ("1.5", 1.5), ("mucho", None)])
def test_importe_para_revisar(app, escrito, valor):
    app.limite_revisar_entry.delete(0, "end")
    app.limite_revisar_entry.insert(0, escrito)
    assert app.limite_revisar() == valor


def test_copiar_lista_incluye_las_azules_aparte(app, archivos):
    quiter, propuesta, carpeta = archivos
    filtrar(quiter, propuesta, carpeta)
    for campo, valor in ((app.excel_quiter_entry, quiter), (app.folder_rot_entry, carpeta)):
        campo.delete(0, "end")
        campo.insert(0, str(valor))
    texto = app.copiar_lista_quiter()
    assert "Para decidir tú" in texto and "CARA (1.850,00 €/ud)" in texto
    assert app.avisos.historial[-1] == ("ok", "Lista copiada: 2 de 6 líneas a poner a 0 (y 1 en azul para que decidas tú). Pégala en Claude.")
