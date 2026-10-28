import threading

import openpyxl

import app_automatizacion_optimizada as app


def test_filas_a_consultar_aplica_el_precio_minimo():
    filas = [(2, "A", 200), (3, "B", "99,5"), (4, "C", "150"), (5, "D", None), (6, "E", "abc"), (7, "F", "150,0")]
    assert app.filas_a_consultar(filas, 150) == [(2, "A"), (4, "C"), (7, "F")]
    assert app.filas_a_consultar(filas, 0) == [(2, "A"), (3, "B"), (4, "C"), (7, "F")]


def test_cola_compartida_no_pierde_ni_repite_nada():
    filas = [(i, f"R{i}") for i in range(200)]
    cola = app.ColaReferencias(filas)
    cogidas = [[] for _ in range(4)]
    hilos = [threading.Thread(target=lambda destino=destino: destino.extend(cola)) for destino in cogidas]
    for hilo in hilos:
        hilo.start()
    for hilo in hilos:
        hilo.join()
    assert sorted(f for lista in cogidas for f in lista) == filas
    assert list(app.ColaReferencias([])) == []


def crear_quiter(ruta, precios):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["h"] * 18)
    for i, precio in enumerate(precios):
        fila = [None] * 18
        fila[1], fila[5] = f"REF{i}", precio
        ws.append(fila)
    wb.save(ruta)


def lanzar(ruta, precio_minimo, num_hilos, worker):
    registro, progresos, fin = [], [], []
    original = app.worker_playwright
    app.worker_playwright = worker
    try:
        app.ejecutar_automatizacion_task(
            str(ruta), "u", "p", precio_minimo, num_hilos,
            log_callback=lambda t, tag=None: registro.append(t),
            progreso_callback=progresos.append,
            finalizado_callback=lambda: fin.append(True),
            stop_event=threading.Event(),
        )
    finally:
        app.worker_playwright = original
    return registro, progresos, fin


def worker_correcto(worker_id, sub_filas, ruta_excel, usuario, password, log_callback, update_progress, stop_event):
    for fila_idx, referencia in sub_filas:
        app.guardar_resultado_excel(ruta_excel, fila_idx, "Pieza disponible")
        update_progress()


def test_automatizacion_consulta_solo_lo_que_supera_el_precio(tmp_path):
    ruta = tmp_path / "q.xlsx"
    crear_quiter(ruta, [200, 100, 300, 50, 150])
    registro, progresos, fin = lanzar(ruta, 150, 3, worker_correcto)
    ws = openpyxl.load_workbook(ruta).active
    consultadas = [ws.cell(row=r, column=18).value for r in range(2, 7)]
    assert consultadas == ["Pieza disponible", None, "Pieza disponible", None, "Pieza disponible"]
    assert progresos[-1] == 1.0 and fin == [True]
    assert app.auto_completada_ok is True
    assert app.automation_running is False


def test_automatizacion_con_todo_bajo_el_minimo_cuenta_como_correcta(tmp_path):
    ruta = tmp_path / "q.xlsx"
    crear_quiter(ruta, [10, 20])
    registro, progresos, fin = lanzar(ruta, 150, 3, worker_correcto)
    assert app.auto_completada_ok is True and fin == [True]
    assert any("precio mínimo" in t for t in registro)


def test_automatizacion_incompleta_no_se_marca_como_correcta(tmp_path):
    ruta = tmp_path / "q.xlsx"
    crear_quiter(ruta, [200, 300, 400])

    def worker_que_falla_en_login(worker_id, sub_filas, *args):
        return  # un hilo que no consigue entrar no procesa nada

    registro, progresos, fin = lanzar(ruta, 0, 2, worker_que_falla_en_login)
    assert app.auto_completada_ok is False
    assert app.automation_running is False
    assert fin == [True]
    assert any("incidencias" in t for t in registro)


def test_automatizacion_sin_referencias_no_es_correcta(tmp_path):
    ruta = tmp_path / "q.xlsx"
    crear_quiter(ruta, [])
    registro, progresos, fin = lanzar(ruta, 0, 2, worker_correcto)
    assert app.auto_completada_ok is False and fin == [True]
