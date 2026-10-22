import re

import openpyxl
import pandas as pd
import pytest

import app_automatizacion_optimizada as app_mod

# (referencia, precio, resultado de la automatización); medias de la propuesta abajo
FILAS = [
    ("REF0", 200, "Pieza disponible"),   # media 0.30 -> blanco
    ("2/REF1", 300, None),               # media 0.05 -> rojo (el prefijo 2/ se ignora)
    ("REF2", 100, "Sin stock"),          # precio < 150 -> sin marcar
    ("REF3", 250, "Sin stock"),          # media 0.20 -> amarillo
    ("REF4", "1.000,5", "disponible"),   # precio no numérico -> no se excluye; no está en la propuesta -> rojo
    ("REF5", 400, "Pieza  DISPONIBLE"),  # media 0.10 -> rojo con umbral 0.12, blanco con 0.08
]
MEDIAS = {"REF0": 0.30, "REF1": 0.05, "REF2": 0.0, "REF3": 0.20, "REF5": 0.10}


@pytest.fixture
def archivos(tmp_path):
    quiter = tmp_path / "S39P.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["cabecera"] * 18)
    for ref, precio, resultado in FILAS:
        fila = [None] * 18
        fila[1], fila[5], fila[17] = ref, precio, resultado
        wb.active.append(fila)
    wb.save(quiter)
    propuesta = tmp_path / "S39PR.xlsx"
    pd.DataFrame({"Articulo": list(MEDIAS), "Media": list(MEDIAS.values())}).to_excel(propuesta, sheet_name="Informe", index=False)
    return quiter, propuesta, tmp_path


def ejecutar_filtro(quiter, propuesta, carpeta, umbral=0.12, precio_minimo=150.0):
    registro = []
    app_mod.filtrar_por_rotacion_task(str(quiter), str(propuesta), str(carpeta), umbral, "FF0000", precio_minimo, log_callback=lambda t, tag=None: registro.append(t))
    resumen = next(t for t in registro if "Resumen rotacion" in t)
    rojas, amarillas, precio = map(int, re.findall(r"rojas por umbral (\d+), amarillas por disponibilidad (\d+), .* EUR (\d+)", resumen)[0])
    return {"rojo": rojas, "amarillo": amarillas, "precio": precio}


def test_regla_de_clasificacion():
    clasificar = app_mod.clasificar_fila_rotacion
    assert clasificar(100, 0.0, "", 0.12, 150) == "precio"       # el precio manda sobre todo lo demás
    assert clasificar(None, 0.0, "", 0.12, 150) == "rojo"
    assert clasificar(200, 0.05, "sin stock", 0.12, 150) == "rojo"  # el rojo manda sobre el amarillo
    assert clasificar(200, 0.20, "sin stock", 0.12, 150) == "amarillo"
    assert clasificar(200, 0.20, "pieza disponible", 0.12, 150) == "blanco"
    assert clasificar(200, 0.20, "", 0.12, 150) == "blanco"


@pytest.mark.parametrize("umbral", [0.0, 0.08, 0.12, 0.5])
def test_la_vista_previa_coincide_con_el_filtro(archivos, umbral):
    quiter, propuesta, carpeta = archivos
    previsto = app_mod.contar_rotacion(app_mod.leer_filas_quiter(str(quiter)), app_mod.leer_medias_propuesta(str(propuesta)), umbral, 150.0)
    real = ejecutar_filtro(quiter, propuesta, carpeta, umbral)
    assert {k: previsto[k] for k in real} == real
    assert sum(previsto.values()) == len(FILAS)


def test_subir_el_umbral_marca_mas_filas_en_rojo(archivos):
    quiter, propuesta, _ = archivos
    filas, medias = app_mod.leer_filas_quiter(str(quiter)), app_mod.leer_medias_propuesta(str(propuesta))
    assert app_mod.contar_rotacion(filas, medias, 0.08, 150)["rojo"] == 2
    assert app_mod.contar_rotacion(filas, medias, 0.12, 150)["rojo"] == 3


def test_lista_para_quiter_sale_del_excel_generado(archivos):
    quiter, propuesta, carpeta = archivos
    ejecutar_filtro(quiter, propuesta, carpeta)
    marcadas, total = app_mod.lineas_marcadas_rotacion(str(carpeta / "S39P_ROTACION.xlsx"))
    assert total == len(FILAS)
    assert marcadas == [(2, "2/REF1", "rojo"), (4, "REF3", "amarillo"), (5, "REF4", "rojo"), (6, "REF5", "rojo")]
    texto = app_mod.texto_lista_quiter("S39P", marcadas, total)
    assert texto.splitlines()[0] == "rotación S39P: 4 de 6 líneas a poner a 0"
    assert "Lín. 4 · REF3 (amarillo)" in texto
    assert app_mod.texto_lista_quiter("S39P", [], 6) == "rotación S39P: ninguna línea marcada de 6"


def poner(entrada, valor):
    entrada.delete(0, "end")
    entrada.insert(0, str(valor))


def test_vista_previa_en_la_pestana(app, archivos):
    quiter, propuesta, carpeta = archivos
    poner(app.excel_quiter_entry, quiter)
    poner(app.excel_rot_entry, propuesta)
    poner(app.umbral_entry, "0.12")
    poner(app.precio_entry, "150")
    app.actualizar_vista_previa_umbral(en_segundo_plano=False)
    assert app.vista_previa_umbral.cget("text") == "Con 0.12: 3 en rojo · 1 en amarillo · 1 sin marcar por precio · 1 en blanco"
    poner(app.umbral_entry, "0.08")
    app.actualizar_vista_previa_umbral(en_segundo_plano=False)
    assert app.vista_previa_umbral.cget("text").startswith("Con 0.08: 2 en rojo")
    poner(app.umbral_entry, "0,08")
    app.actualizar_vista_previa_umbral(en_segundo_plano=False)
    assert "punto decimal" in app.vista_previa_umbral.cget("text")
    poner(app.excel_quiter_entry, carpeta / "no_existe.xlsx")
    app.actualizar_vista_previa_umbral(en_segundo_plano=False)
    assert app.vista_previa_umbral.cget("text").startswith("Elige el Excel de Quiter")


def test_vista_previa_avisa_si_aun_no_hay_disponibilidad(app, archivos, tmp_path):
    quiter, propuesta, _ = archivos
    wb = openpyxl.load_workbook(quiter)
    for fila in range(2, wb.active.max_row + 1):
        wb.active.cell(row=fila, column=18).value = None
    sin_resultados = tmp_path / "sin_resultados.xlsx"
    wb.save(sin_resultados)
    poner(app.excel_quiter_entry, sin_resultados)
    poner(app.excel_rot_entry, propuesta)
    app.actualizar_vista_previa_umbral(en_segundo_plano=False)
    assert app.vista_previa_umbral.cget("text").endswith("(el amarillo aparece tras la automatización)")


def test_copiar_lista_para_quiter(app, archivos, monkeypatch):
    quiter, propuesta, carpeta = archivos
    avisos = []
    monkeypatch.setattr(app_mod.messagebox, "showinfo", lambda *a, **k: avisos.append(a))
    poner(app.excel_quiter_entry, quiter)
    poner(app.folder_rot_entry, carpeta)
    assert app.copiar_lista_quiter() is None and avisos  # aún no hay Excel de rotación
    ejecutar_filtro(quiter, propuesta, carpeta)
    texto = app.copiar_lista_quiter()
    assert texto.startswith("rotación S39P: 4 de 6 líneas a poner a 0")
    assert app.clipboard_get() == texto
