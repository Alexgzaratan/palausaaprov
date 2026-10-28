import time
from datetime import date, datetime, timedelta

import pytest

import app_automatizacion_optimizada as app_mod
import configuracion
import tareas


def esperar(app, condicion, limite=5):
    """Ejecuta el bucle de eventos real (los hilos avisan a la ventana con after) hasta que se cumple la condición."""
    fin = time.time() + limite

    def revisar():
        if condicion() or time.time() > fin:
            app.quit()
        else:
            app.after(30, revisar)

    app.after(30, revisar)
    app.mainloop()
    return condicion()


def lanzar_en_el_bucle(app, funcion):
    """Lanza la acción con el bucle de eventos ya en marcha, como al pulsar el botón en la app."""
    resultado = []
    app.after(0, lambda: resultado.append(funcion()))
    return resultado


def poner(campo, valor):
    campo.delete(0, "end")
    campo.insert(0, str(valor))


# ---------- 1. propuesta automática al cambiar de pedido ----------

@pytest.fixture
def semana(app, tmp_path):
    rotacion = tmp_path / "ROTACION"
    rotacion.mkdir()
    (rotacion / "S40PR.xls").write_text("x")
    nuevo = tmp_path / "S40P.xlsx"
    nuevo.write_text("x")
    anterior = (tmp_path / "S39P.xlsx").as_posix()
    for campo, valor in ((app.excel_auto_entry, anterior), (app.excel_quiter_entry, anterior),
                         (app.excel_rot_entry, (rotacion / "S39PR.xls").as_posix()), (app.folder_rot_entry, rotacion.as_posix())):
        poner(campo, valor)
        campo._valor_anterior = valor
    return nuevo, rotacion


def test_elegir_el_excel_de_stock_rellena_quiter_y_propuesta(app, semana, monkeypatch):
    nuevo, rotacion = semana
    monkeypatch.setattr(app_mod.filedialog, "askopenfilename", lambda **k: nuevo.as_posix())
    app.seleccionar_archivo(app.excel_auto_entry, [("Excel files", "*.xlsx")])
    assert app.excel_quiter_entry.get() == nuevo.as_posix()
    assert app.excel_rot_entry.get() == (rotacion / "S40PR.xls").as_posix()
    assert app.avisos.historial[-1] == ("ok", "Pedido S40P: he puesto también el Excel de Quiter y su propuesta S40PR.xls.")


def test_soltar_el_excel_de_stock_tambien_rellena(app, semana):
    nuevo, rotacion = semana
    assert app.soltar_en_campo(app.excel_auto_entry, "{" + str(nuevo) + "}", "archivo", (".xlsx",))
    assert app.excel_rot_entry.get() == (rotacion / "S40PR.xls").as_posix()


def test_no_toca_un_excel_de_quiter_distinto(app, semana, tmp_path):
    nuevo, _ = semana
    otro = (tmp_path / "otro.xlsx").as_posix()
    poner(app.excel_quiter_entry, otro)
    poner(app.excel_auto_entry, nuevo.as_posix())
    app.campo_cambiado(app.excel_auto_entry)  # como al escribir la ruta y salir del campo
    assert app.excel_quiter_entry.get() == otro
    assert app.avisos.historial[-1] == ("ok", "Pedido S40P: he puesto también su propuesta S40PR.xls.")


def test_avisa_si_la_propuesta_aun_no_esta(app, semana, tmp_path):
    _, rotacion = semana
    (rotacion / "S40PR.xls").unlink()
    poner(app.excel_quiter_entry, "")
    sin = tmp_path / "S40P.xlsx"
    poner(app.excel_auto_entry, sin.as_posix())
    app.campo_cambiado(app.excel_auto_entry)
    assert app.excel_rot_entry.get().endswith("S39PR.xls")  # la anterior se queda hasta que exportes la nueva
    assert app.avisos.historial[-1] == ("info", "Pedido S40P: he puesto también el Excel de Quiter. Aún no está la propuesta S40PR.xls en ROTACION: elígela cuando la exportes.")
    poner(app.excel_quiter_entry, "C:/otro.xlsx")
    poner(app.excel_auto_entry, (tmp_path / "S41P.xlsx").as_posix())
    app.campo_cambiado(app.excel_auto_entry)
    assert app.avisos.historial[-1] == ("info", "Aún no está la propuesta S41PR.xls en ROTACION: elígela cuando la exportes.")


# ---------- 2. hoja de hoy en Avisos ----------

@pytest.fixture
def hoja_de_hoy(app, tmp_path):
    app.tabview.set("Pedidos")
    app.al_cambiar_pestana()
    poner(app.carpeta_hojas_entry, tmp_path)
    hoja = app_mod.ruta_hoja_pedidos_del_dia(str(tmp_path))
    hoja.parent.mkdir(parents=True)
    hoja.write_text("x")
    return hoja


def test_boton_hoy_pone_la_hoja_de_hoy(app, hoja_de_hoy):
    poner(app.hoja_avisos_entry, "C:/HOJAS/27-10-2026.xlsx")
    boton = next(h for h in app.hoja_avisos_entry.master.winfo_children() if type(h).__name__ == "CTkButton" and h.cget("text") == "Hoy")
    boton.invoke()
    assert app.hoja_avisos_entry.get() == hoja_de_hoy.as_posix()


def test_propone_la_hoja_de_hoy_sin_cambiarla_sola(app, hoja_de_hoy):
    poner(app.hoja_avisos_entry, "C:/HOJAS/27-10-2026.xlsx")
    aviso = app.sugerir_hoja_de_hoy_en_avisos()
    assert app.hoja_avisos_entry.get() == "C:/HOJAS/27-10-2026.xlsx"  # no la cambia sin preguntar
    aviso = app.avisos.activos[-1]
    assert "27-10-2026.xlsx" in aviso.texto and aviso.boton_accion.cget("text") == "Usar"
    aviso.boton_accion.invoke()
    assert app.hoja_avisos_entry.get() == hoja_de_hoy.as_posix()
    cuantos = len(app.avisos.historial)
    app.sugerir_hoja_de_hoy_en_avisos()  # ya es la de hoy: no vuelve a proponer
    assert len(app.avisos.historial) == cuantos


# ---------- 3. avisos con botón ----------

def test_aviso_con_accion(app):
    llamadas = []
    aviso = app.avisar("Rotación completada: S40P_ROTACION.xlsx", "ok", accion=("Abrir", lambda: llamadas.append(1)))
    assert aviso.boton_accion.cget("text") == "Abrir"
    aviso.boton_accion.invoke()
    assert llamadas == [1]
    assert esperar(app, lambda: not aviso.winfo_exists())


def test_rotacion_completada_ofrece_abrir_el_excel(app, tmp_path, monkeypatch):
    abiertos = []
    monkeypatch.setattr(app_mod, "abrir_ruta_en_sistema", lambda ruta: abiertos.append(str(ruta)) or True)
    salida = tmp_path / "S40P_ROTACION.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append([None, "REF1"])
    wb.save(salida)
    poner(app.excel_quiter_entry, tmp_path / "S40P.xlsx")
    poner(app.folder_rot_entry, tmp_path)
    app._inicio_rotacion = time.time() - 1
    app.finalizar_filtro_rotacion()
    aviso = app.avisos.activos[-1]
    aviso.boton_accion.invoke()
    assert abiertos == [str(salida)]
    assert app.etiquetas_ultima_vez["rotacion"].cget("text").endswith("S40P: 0 de 1 líneas")


# ---------- 4. recordatorio de tareas ----------

def test_recordatorio_de_tareas_al_abrir(app, tmp_path):
    app.lista_tareas = tareas.ListaTareas(str(tmp_path / "t.json"))
    assert app.recordar_tareas() is None  # sin nada para hoy, no molesta
    hoy = date.today()
    app.lista_tareas.anadir("Para hoy", hoy)
    app.lista_tareas.anadir("Vencida", hoy - timedelta(days=3))
    app.lista_tareas.anadir("Sin fecha")
    aviso = app.recordar_tareas()
    assert aviso.texto == "Tienes 1 tarea para hoy y 1 vencida." and app.avisos.historial[-1][0] == "aviso"
    aviso.boton_accion.invoke()
    assert app.tabview.get() == "Comparador"


# ---------- 6. última vez ----------

def test_texto_ultima_vez():
    ahora = datetime(2026, 10, 28, 12)
    assert app_mod.texto_ultima_vez(None, ahora) == ""
    assert app_mod.texto_ultima_vez({"cuando": "2026-10-28T09:42:10", "detalle": "S40P: 24 de 46 líneas"}, ahora) == "Última vez: hoy 09:42 · S40P: 24 de 46 líneas"
    assert app_mod.texto_ultima_vez({"cuando": "2026-10-27T17:05:00", "detalle": ""}, ahora) == "Última vez: ayer 17:05"
    assert app_mod.texto_ultima_vez({"cuando": "2026-10-20T08:00:00"}, ahora) == "Última vez: 20/10 08:00"


def test_ultima_vez_se_guarda_y_se_ve_en_la_tarjeta(app):
    app.registrar_ejecucion("auto", "S40P completa")
    assert configuracion.cargar_configuracion()["ultimas_ejecuciones"]["auto"]["detalle"] == "S40P completa"
    assert app.etiquetas_ultima_vez["auto"].cget("text").startswith("Última vez: hoy ")
    app.guardar_configuracion_ui_internal()  # guardar la configuración no la borra
    assert configuracion.cargar_configuracion()["ultimas_ejecuciones"]["auto"]["detalle"] == "S40P completa"


# ---------- 7. reenviar facturas de un día ----------

def test_interpretar_fecha_pasada():
    hoy = date(2026, 10, 28)
    assert app_mod.interpretar_fecha_pasada("ayer", hoy) == date(2026, 10, 27)
    assert app_mod.interpretar_fecha_pasada("06/10", hoy) == date(2026, 10, 6)
    assert app_mod.interpretar_fecha_pasada("20/12", hoy) == date(2025, 12, 20)  # aún no ha llegado: el del año pasado
    assert app_mod.interpretar_fecha_pasada("8-10-26", hoy) == date(2026, 10, 8)
    for malo in ("", "lunes", "31/02", "01/01/2030"):
        with pytest.raises(ValueError):
            app_mod.interpretar_fecha_pasada(malo, hoy)


def test_facturas_locales_del_dia(tmp_path):
    for nombre in ("Factura_1_2026-10-06.pdf", "Factura_2_2026-10-06.pdf", "Factura_3_2026-10-07.pdf", "otra_2026-10-06.pdf"):
        (tmp_path / nombre).write_text("x")
    assert [r.name for r in app_mod.facturas_locales_del_dia(tmp_path, date(2026, 10, 6))] == ["Factura_1_2026-10-06.pdf", "Factura_2_2026-10-06.pdf"]


def test_la_cabecera_no_cuenta_el_reenvio_de_otro_dia(tmp_path, monkeypatch):
    hoy = datetime(2026, 10, 28, 9)
    log = tmp_path / "log.txt"
    log.write_text("--- CORREO [2026-10-28 08:30:00] ---\nFecha objetivo: 2026-10-06\nEstado: enviado\n", encoding="utf-8")
    monkeypatch.setattr(app_mod, "LOG_FACTURAS_MADD", str(log))
    assert app_mod.resumen_del_dia({"usuario_madd": "X"}, hoy) == [("factura", "Facturas de hoy aún sin enviar", False)]
    with open(log, "a", encoding="utf-8") as f:
        f.write("--- CORREO [2026-10-28 08:31:00] ---\nFecha objetivo: 2026-10-27\nEstado: enviado\n")
    assert app_mod.resumen_del_dia({"usuario_madd": "X"}, hoy) == [("check", "Facturas enviadas a las 08:31", True)]


@pytest.fixture
def facturas(app, tmp_path, monkeypatch):
    app.tabview.set("Facturas")
    app.al_cambiar_pestana()
    poner(app.carpeta_facturas_madd_entry, tmp_path)
    poner(app.email_facturas_madd_entry, "a@x.es")
    monkeypatch.setattr(app_mod.messagebox, "askyesno", lambda *a, **k: True)
    return tmp_path


def test_reenviar_con_los_pdf_de_la_carpeta(app, facturas, monkeypatch):
    for nombre in ("Factura_1_2026-10-06.pdf", "Factura_2_2026-10-06.pdf"):
        (facturas / nombre).write_text("x")
    enviados = []
    monkeypatch.setattr(app_mod, "_madd_enviar_correo_resultado", lambda destino, fecha, rutas, log: enviados.append((destino, fecha, len(rutas))))
    poner(app.fecha_reenvio_madd_entry, "06/10")
    lanzado = lanzar_en_el_bucle(app, app.reenviar_facturas_dia)
    assert esperar(app, lambda: bool(enviados) and app.btn_reenviar_madd.cget("text") == "Reenviar")
    assert lanzado == [True]
    assert enviados == [("a@x.es", app_mod.interpretar_fecha_pasada("06/10"), 2)]
    assert not app_mod.madd_running
    assert app.etiquetas_ultima_vez["madd"].cget("text").endswith("reenvío del 06/10 (2 facturas)")


def test_reenviar_sin_pdf_descarga_de_madd_ese_dia(app, facturas, monkeypatch):
    llamadas = []

    def tarea_falsa(*args, **kwargs):
        llamadas.append(kwargs["fecha_objetivo"])
        app_mod.madd_running = False
        kwargs["finalizado_callback"]()

    monkeypatch.setattr(app_mod, "tarea_facturas_madd", tarea_falsa)
    poner(app.fecha_reenvio_madd_entry, "08/10")
    poner(app.usuario_madd_entry, "")
    assert not app.reenviar_facturas_dia() and app.mensajes  # sin credenciales no puede descargar
    poner(app.usuario_madd_entry, "U")
    poner(app.password_madd_entry, "P")
    lanzado = lanzar_en_el_bucle(app, app.reenviar_facturas_dia)
    assert esperar(app, lambda: bool(llamadas) and app.btn_reenviar_madd.cget("text") == "Reenviar")
    assert lanzado == [True]
    assert llamadas == [app_mod.interpretar_fecha_pasada("08/10")]


def test_tarea_de_facturas_de_un_dia_concreto(monkeypatch):
    dias = []
    monkeypatch.setattr(app_mod, "ejecutar_descarga_facturas_madd", lambda u, p, c, fechas, log: dias.append(fechas) or {"facturas": ["a.pdf"], "error": None})
    monkeypatch.setattr(app_mod, "_madd_enviar_correo_resultado", lambda *a: None)
    monkeypatch.setattr(app_mod, "registrar_log_facturas_madd", lambda *a: None)
    app_mod.tarea_facturas_madd("u", "p", "c", "a@x.es", "", fecha_objetivo=date(2026, 10, 8))
    assert dias == [{date(2026, 10, 8)}]
    assert configuracion.cargar_configuracion()["ultimas_ejecuciones"]["madd"]["detalle"] == "1 facturas del 08/10"
