import customtkinter as ctk
import threading
import os
import sys
import openpyxl
from openpyxl import load_workbook
import csv
import re
import ctypes
import pandas as pd
import time
import traceback
import shutil
import xlrd
import string
import subprocess
import tkinter
import tkinter.font
import queue
from datetime import datetime, timedelta
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from tkinter import filedialog, messagebox
from openpyxl.styles import PatternFill
from concurrent.futures import ThreadPoolExecutor
from reclamaciones_power_line import ModuloReclamacionesPowerLine
import estilo
import animacion_inicio
import efectos
import piezas_caras
import actualizaciones
import cambio_version
import instancia_unica
from version import VERSION
import tareas
from rutas_app import ruta_datos
from configuracion import LOG_DETALLADO, cargar_configuracion, guardar_configuracion, registrar_error_detallado

# --- CONFIGURACIÓN DE WINDOWS PARA ICONO EN BARRA DE TAREAS ---
try:
    myappid = 'Palausa.Aprovisionamiento'
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
except Exception:
    pass

# --- CONFIGURACIÓN GLOBAL ---
ICONO_NOMBRE = "ultraball.ico"
LOG_MAX_LINEAS = 2000
LOG_LINEAS_A_RECORTAR = 500
ERROR_FOLDER = ruta_datos("errores_capturas")

# Bloqueo para escritura segura en Excel desde múltiples hilos
excel_lock = threading.Lock()

def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

ICONO_PATH = resource_path(ICONO_NOMBRE)

os.makedirs(ERROR_FOLDER, exist_ok=True)

# Variables globales para control de ejecución
automation_running = False
pedidos_running = False
comparador_running = False
filtro_rotacion_running = False
mopar_running = False
stock_fiat_running = False
hoja_pedidos_running = False
avisos_running = False
madd_running = False
auto_completada_ok = False
auto_resultado = "error"  # "completa", "incompleta", "detenida", "credenciales" o "error"
stop_event_auto = threading.Event()
stop_event_pedidos = threading.Event()
stop_event_comparador = threading.Event()
stop_event_mopar = threading.Event()
stop_event_stock_fiat = threading.Event()

# (clave en config, atributo del widget, valor por defecto, conversor opcional)
CAMPOS_CONFIG = (
    ("usuario", "user_entry", ""),
    ("password", "pass_entry", ""),
    ("num_hilos", "hilos_option", "3", int),
    ("ruta_excel_auto", "excel_auto_entry", ""),
    ("ruta_excel_pedidos", "excel_pedidos_entry", ""),
    ("ruta_folder_pedidos", "folder_pedidos_entry", ""),
    ("ruta_plantilla_hojas", "plantilla_hoja_entry", ""),
    ("ruta_carpeta_hojas", "carpeta_hojas_entry", ""),
    ("ruta_hoja_avisos", "hoja_avisos_entry", ""),
    ("ruta_excel_clientes", "clientes_avisos_entry", ""),
    ("email_companeros_avisos", "email_companeros_avisos_entry", ""),
    ("precio_min", "precio_entry", "0"),
    ("ruta_excel_antiguo", "excel_old_entry", ""),
    ("ruta_excel_nuevo", "excel_new_entry", ""),
    ("ruta_folder_comparador", "folder_comp_entry", ""),
    ("ruta_excel_quiter", "excel_quiter_entry", ""),
    ("ruta_excel_rotacion", "excel_rot_entry", ""),
    ("ruta_folder_rotacion", "folder_rot_entry", ""),
    ("umbral_rotacion", "umbral_entry", "0.08"),
    ("limite_revisar_rotacion", "limite_revisar_entry", "1600"),
    ("ruta_excel_mopar", "excel_mopar_entry", ""),
    ("ruta_folder_mopar", "folder_mopar_entry", ""),
    ("ruta_excel_stock_fiat", "excel_stock_fiat_entry", ""),
    ("ruta_folder_stock_fiat", "folder_stock_fiat_entry", ""),
    ("usuario_madd", "usuario_madd_entry", ""),
    ("password_madd", "password_madd_entry", ""),
    ("ruta_carpeta_facturas_madd", "carpeta_facturas_madd_entry", ""),
    ("email_facturas_madd", "email_facturas_madd_entry", ""),
    ("email_aviso_fallo_madd", "email_aviso_fallo_madd_entry", ""),
    ("filtro_dia_facturas_madd", "filtro_dia_madd_option", "Ayer"),
    ("color_resaltado", "color_optionmenu", "Amarillo"),
    ("nombre_usuario", "nombre_entry", ""),
    ("ruta_carpeta_informes", "carpeta_informes_entry", ""),
    ("ruta_excel_pedidos_stock", "excel_pedidos_stock_entry", ""),
)

DIAS_ES = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")
MESES_LARGOS_ES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")

def texto_saludo(nombre="", ahora=None):
    hora = (ahora or datetime.now()).hour
    saludo = "Buenos días" if 6 <= hora < 14 else "Buenas tardes" if 14 <= hora < 21 else "Buenas noches"
    nombre = (nombre or "").strip()
    return f"{saludo}, {nombre}" if nombre else saludo

def fecha_larga(ahora=None):
    ahora = ahora or datetime.now()
    texto = f"{DIAS_ES[ahora.weekday()]}, {ahora.day} de {MESES_LARGOS_ES[ahora.month - 1]}"
    return texto[0].upper() + texto[1:]

def resumen_del_dia(config, ahora=None):
    """Fichas (icono, texto, correcto) con lo que ya está hecho hoy. Solo incluye lo que está configurado."""
    ahora = ahora or datetime.now()
    fichas = []
    carpeta_hojas = config.get("ruta_carpeta_hojas", "")
    if carpeta_hojas and os.path.isdir(carpeta_hojas):
        if ruta_hoja_pedidos_del_dia(carpeta_hojas, ahora).exists():
            fichas.append(("check", "Hoja de pedidos de hoy creada", True))
        else:
            fichas.append(("calendario", "Hoja de pedidos de hoy pendiente", False))
    if config.get("usuario_madd") and ahora.weekday() < 5 and os.path.exists(LOG_FACTURAS_MADD):
        hoy = ahora.strftime("%Y-%m-%d")
        # Las facturas que envía hoy la tarea diaria son las del día laborable anterior (o las de hoy, según la configuración)
        objetivo = (ahora.date() if config.get("filtro_dia_facturas_madd") == "Hoy" else _dia_laborable_anterior(ahora.date())).isoformat()
        estado = None
        try:
            with open(LOG_FACTURAS_MADD, encoding="utf-8") as f:
                bloque_hoy = None
                for linea in f:
                    if linea.startswith("--- CORREO [" + hoy):
                        bloque_hoy = linea.split(" ")[3][:5]  # "--- CORREO [AAAA-MM-DD HH:MM:SS] ---" -> "HH:MM"
                    elif bloque_hoy and linea.startswith("Fecha objetivo:") and linea.split(":", 1)[1].strip() != objetivo:
                        bloque_hoy = None  # correo de otro día (p. ej. un reenvío)
                    elif bloque_hoy and linea.startswith("Estado:"):
                        estado = (bloque_hoy, linea.split(":", 1)[1].strip())
                        bloque_hoy = None
        except OSError:
            estado = None
        if estado and estado[1].startswith("enviado"):
            fichas.append(("check", f"Facturas enviadas a las {estado[0]}", True))
        elif estado:
            fichas.append(("aviso", "Facturas pendientes en Outlook", False))
        else:
            fichas.append(("factura", "Facturas de hoy aún sin enviar", False))
    return fichas

def tarea_activa():
    """Nombre de la tarea en curso que trabaja con ficheros o envía correos, o None si no hay ninguna."""
    for activa, nombre in (
        (automation_running, "Automatizacion"),
        (filtro_rotacion_running, "Filtro de rotacion"),
        (pedidos_running, "Generacion de pedidos"),
        (hoja_pedidos_running, "Hoja de pedidos del dia"),
        (avisos_running, "Avisos de no suministro"),
        (comparador_running, "Comparador de Excel"),
        (mopar_running, "Fichero MOPAR"),
        (stock_fiat_running, "Stock FIAT"),
        (madd_running, "Facturas MADD"),
    ):
        if activa:
            return nombre
    return None

def correos_invalidos(texto):
    """Direcciones de una lista separada por ';' o ',' que no tienen forma de correo electrónico."""
    direcciones = [d.strip() for d in re.split(r"[;,]", texto or "") if d.strip()]
    return [d for d in direcciones if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", d)]

def archivo_bloqueado(ruta):
    """True si el fichero existe y otro programa (normalmente Excel) lo tiene abierto y no se puede escribir."""
    if not ruta or not os.path.exists(ruta):
        return False
    try:
        with open(ruta, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return False

def abrir_ruta_en_sistema(ruta):
    if not ruta or not os.path.exists(ruta):
        return False
    try:
        os.startfile(ruta)
        return True
    except Exception:
        return False

def crear_copia_seguridad(ruta_archivo, carpeta_destino=None):
    carpeta_destino = carpeta_destino or ruta_datos("backups_pedidos")
    Path(carpeta_destino).mkdir(parents=True, exist_ok=True)
    nombre_base = Path(ruta_archivo).stem
    extension = Path(ruta_archivo).suffix
    marca_tiempo = datetime.now().strftime("%Y%m%d_%H%M%S")
    ruta_backup = Path(carpeta_destino) / f"{nombre_base}_backup_{marca_tiempo}{extension}"
    shutil.copy2(ruta_archivo, ruta_backup)
    return str(ruta_backup)

def obtener_numero_mopar_del_dia(carpeta_salida, prefijo_clave="mopar"):
    config = cargar_configuracion()
    clave_fecha = f"{prefijo_clave}_fecha"
    clave_numero = f"{prefijo_clave}_numero"
    hoy = datetime.now().strftime("%Y-%m-%d")
    ultimo_dia = config.get(clave_fecha)
    ultimo_numero = config.get(clave_numero)

    if ultimo_dia == hoy and ultimo_numero is not None:
        return int(ultimo_numero)

    ultimo_f = 0
    for f in os.listdir(carpeta_salida):
        m = re.match(r"F(\d+)", f)
        if m:
            ultimo_f = max(ultimo_f, int(m.group(1)))

    proximo_f = ultimo_f + 1
    config[clave_fecha] = hoy
    config[clave_numero] = proximo_f
    guardar_configuracion(config)
    return proximo_f

# ---------------------------------------------------
# LÓGICA DE GESTIÓN DE CARPETAS POR RANGOS (PARA PEDIDOS)
# ---------------------------------------------------

def obtener_ultimo_contador_y_ruta(carpeta_base, extension=".csv"):
    if not os.path.exists(carpeta_base):
        os.makedirs(carpeta_base, exist_ok=True)
    ultimo_contador = 4999
    subcarpetas = [d for d in os.listdir(carpeta_base) if os.path.isdir(os.path.join(carpeta_base, d))]
    subcarpetas_rango = []
    for d in subcarpetas:
        match = re.match(r"(\d+)-(\d+)", d)
        if match:
            subcarpetas_rango.append((int(match.group(1)), d))
    if subcarpetas_rango:
        subcarpetas_rango.sort()
        ultimo_rango_inicio, nombre_subcarpeta = subcarpetas_rango[-1]
        ruta_subcarpeta = os.path.join(carpeta_base, nombre_subcarpeta)
        regex_archivo = r"(\d+)" + re.escape(extension)
        for f in os.listdir(ruta_subcarpeta):
            m = re.match(regex_archivo, f)
            if m:
                ultimo_contador = max(ultimo_contador, int(m.group(1)))
    else:
        regex_archivo = r"(\d+)" + re.escape(extension)
        for f in os.listdir(carpeta_base):
            m = re.match(regex_archivo, f)
            if m:
                ultimo_contador = max(ultimo_contador, int(m.group(1)))
    return ultimo_contador

def obtener_ruta_para_contador(carpeta_base, contador):
    rango_inicio = (contador // 100) * 100
    rango_fin = rango_inicio + 99
    nombre_rango = f"{rango_inicio}-{rango_fin}"
    ruta_rango = os.path.join(carpeta_base, nombre_rango)
    if not os.path.exists(ruta_rango):
        os.makedirs(ruta_rango, exist_ok=True)
    return ruta_rango

# ---------------------------------------------------
# 1. LÓGICA DE FILTRO DE ROTACIÓN
# ---------------------------------------------------

# ---------- regla del filtro de rotación (la usan el filtro, su vista previa y la lista para Quiter) ----------
RESULTADOS_DISPONIBLES = ("pieza disponible", "disponible")
COLOR_ROJO_ROTACION = "FF0000"
COLOR_AMARILLO_ROTACION = "FFFF00"
COLOR_REVISAR_ROTACION = "9DC3E6"  # azul claro: piezas caras que decide Álex


def precio_fila_rotacion(valor):
    """Precio de la columna F del Excel de Quiter, o None si no es un número."""
    try:
        return float(str(valor).replace(",", "."))
    except Exception:
        return None


def referencia_rotacion(valor):
    return str(valor).strip().replace('2/', '')


def resultado_disponibilidad(valor):
    """Resultado de la automatización (columna R) normalizado para compararlo."""
    return "" if valor is None else " ".join(str(valor).strip().lower().split())


def clasificar_fila_rotacion(precio, media, resultado, umbral_media, precio_minimo, precio_unitario=None, limite_revisar=0):
    """'precio' (se deja sin marcar), 'rojo' (baja rotación), 'amarillo' (no disponible), 'revisar' (pieza cara:
    el precio unitario supera el límite; la decide Álex) o 'blanco'. El rojo y el amarillo mandan sobre el azul.
    precio es el total de la línea (columna F) y precio_unitario el precio de compra (columna E)."""
    if precio is not None and precio < precio_minimo:
        return "precio"
    if media < umbral_media:
        return "rojo"
    if resultado and resultado not in RESULTADOS_DISPONIBLES:
        return "amarillo"
    if limite_revisar and precio_unitario is not None and precio_unitario > limite_revisar:
        return "revisar"
    return "blanco"


def leer_medias_propuesta(ruta_propuesta_xls):
    """Media de cada referencia en la hoja 'Informe' de la propuesta; None si faltan las columnas."""
    df_propuesta = pd.read_excel(ruta_propuesta_xls, sheet_name='Informe')
    columnas = [str(c).strip() for c in df_propuesta.columns]
    col_articulo = next((c for c in columnas if c.lower() in ["articulo", "artículo", "referencia", "ref", "código", "codigo"]), None)
    col_media = next((c for c in columnas if c.lower() in ["media", "consumo", "promedio"]), None)
    if not col_articulo or not col_media:
        return None
    df_propuesta.columns = columnas
    dict_medias = {}
    for _, fila in df_propuesta.iterrows():
        try:
            dict_medias[referencia_rotacion(fila[col_articulo])] = float(fila[col_media])
        except Exception:
            continue
    return dict_medias


def leer_filas_quiter(ruta_excel_quiter):
    """(total de la línea, referencia, resultado, precio unitario) de cada fila del Excel de Quiter, sin la cabecera."""
    wb = load_workbook(ruta_excel_quiter, read_only=True)
    try:
        filas = []
        for valores in wb.active.iter_rows(min_row=2, max_col=18, values_only=True):
            valores = tuple(valores) + (None,) * (18 - len(valores))
            filas.append((precio_fila_rotacion(valores[5]), referencia_rotacion(valores[1]), resultado_disponibilidad(valores[17]),
                          precio_fila_rotacion(valores[4])))
        return filas
    finally:
        wb.close()


def contar_rotacion(filas, dict_medias, umbral_media, precio_minimo, limite_revisar=0):
    conteo = {"rojo": 0, "amarillo": 0, "revisar": 0, "precio": 0, "blanco": 0}
    for precio, ref, resultado, *unitario in filas:
        conteo[clasificar_fila_rotacion(precio, dict_medias.get(ref, 0.0), resultado, umbral_media, precio_minimo,
                                        unitario[0] if unitario else None, limite_revisar)] += 1
    return conteo


def colores_rotacion(ruta_rotacion):
    """Filas del Excel _ROTACION.xlsx ya generado: [(línea, referencia, color, precio unitario)], con color
    'rojo', 'amarillo', 'revisar' o ''. La fila N del archivo es la línea N del pedido en Quiter (el filtro quita la cabecera)."""
    wb = load_workbook(ruta_rotacion)
    try:
        return colores_rotacion_hoja(wb.active)
    finally:
        wb.close()


def colores_rotacion_hoja(ws):
    filas = []
    for fila in range(1, ws.max_row + 1):
        celda = ws.cell(row=fila, column=2)
        if celda.value is None or not str(celda.value).strip():
            continue
        relleno = str(celda.fill.fgColor.rgb or "")[-6:].upper() if celda.fill.fill_type == "solid" else ""
        color = {COLOR_ROJO_ROTACION: "rojo", COLOR_AMARILLO_ROTACION: "amarillo", COLOR_REVISAR_ROTACION: "revisar"}.get(relleno, "")
        filas.append((fila, str(celda.value).strip(), color, precio_fila_rotacion(ws.cell(row=fila, column=5).value)))
    return filas


def lineas_marcadas_rotacion(ruta_rotacion):
    """Líneas en rojo o amarillo (las que se ponen a 0): ([(línea, referencia, color)], total de líneas)."""
    filas = colores_rotacion(ruta_rotacion)
    return [(linea, ref, color) for linea, ref, color, _ in filas if color in ("rojo", "amarillo")], len(filas)


def piezas_caras_rotacion(ruta_rotacion):
    """Piezas en azul del Excel _ROTACION.xlsx (aún sin decidir), con lo necesario para decidir:
    [{"linea", "referencia", "descripcion", "pedidas", "unitario", "total"}]."""
    wb = load_workbook(ruta_rotacion)
    try:
        ws = wb.active
        piezas = []
        for linea, ref, color, unitario in colores_rotacion_hoja(ws):
            if color == "revisar":
                piezas.append({"linea": linea, "referencia": ref, "descripcion": ws.cell(row=linea, column=3).value or "",
                               "pedidas": piezas_caras.numero(ws.cell(row=linea, column=4).value), "unitario": unitario,
                               "total": precio_fila_rotacion(ws.cell(row=linea, column=6).value)})
        return piezas
    finally:
        wb.close()


def decidir_piezas_caras_en_excel(ruta_rotacion, decisiones):
    """Aplica lo decidido a las filas azules: pedir -> sin relleno (blanco), poner a 0 -> rojo. Solo toca filas
    que sigan en azul. Devuelve (para pedir, a cero). Lanza PermissionError si el Excel está abierto."""
    wb = load_workbook(ruta_rotacion)
    try:
        ws = wb.active
        rojo = PatternFill(start_color=COLOR_ROJO_ROTACION, end_color=COLOR_ROJO_ROTACION, fill_type="solid")
        azules = {linea for linea, _, color, _ in colores_rotacion_hoja(ws) if color == "revisar"}
        pedir = cero = 0
        for linea, decision in decisiones.items():
            if linea not in azules:
                continue
            for celda in ws[linea]:
                celda.fill = rojo if decision == piezas_caras.A_CERO else PatternFill(fill_type=None)
            if decision == piezas_caras.A_CERO:
                cero += 1
            else:
                pedir += 1
        wb.save(ruta_rotacion)
        return pedir, cero
    finally:
        wb.close()


def texto_lista_quiter(pedido, marcadas, total):
    """Texto para pegar en Claude (skill quiter-lineas-a-cero)."""
    if not marcadas:
        return f"rotación {pedido}: ninguna línea marcada de {total}"
    lineas = [f"rotación {pedido}: {len(marcadas)} de {total} líneas a poner a 0"]
    lineas += [f"Lín. {linea} · {ref} ({color})" for linea, ref, color in marcadas]
    return "\n".join(lineas)


def filtrar_por_rotacion_task(ruta_excel_quiter, ruta_propuesta_xls, ruta_salida_folder, umbral_media, color_resaltado, precio_minimo=0.0, log_callback=None, progreso_callback=None, finalizado_callback=None, limite_revisar=0.0):
    global filtro_rotacion_running
    filtro_rotacion_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[ROTACIÓN] {texto}", tag)
    log(f"📊 Iniciando análisis de rotación...")
    try:
        log("📖 Leyendo archivo de rotación...")
        try:
            dict_medias = leer_medias_propuesta(ruta_propuesta_xls)
        except Exception as e:
            log(f"❌ Error al leer el archivo Excel: {e}", "error")
            registrar_error_detallado(str(e), traceback.format_exc())
            return
        if dict_medias is None:
            log(f"❌ Error: No se encuentran las columnas necesarias.", "error")
            return
        log(f"✅ Cargadas {len(dict_medias)} referencias.", "success")
        wb = load_workbook(ruta_excel_quiter)
        ws = wb.active
        fill_color = PatternFill(start_color=color_resaltado, end_color=color_resaltado, fill_type="solid")
        fill_no_disponible = PatternFill(start_color=COLOR_AMARILLO_ROTACION, end_color=COLOR_AMARILLO_ROTACION, fill_type="solid")
        fill_revisar = PatternFill(start_color=COLOR_REVISAR_ROTACION, end_color=COLOR_REVISAR_ROTACION, fill_type="solid")
        total_filas = max(ws.max_row - 1, 1)
        filas_baja_rotacion = 0
        filas_no_disponibles = 0
        filas_revisar = 0
        filas_bajo_precio = 0
        for i in range(2, ws.max_row + 1):
            ref = referencia_rotacion(ws.cell(row=i, column=2).value)
            clase = clasificar_fila_rotacion(precio_fila_rotacion(ws.cell(row=i, column=6).value), dict_medias.get(ref, 0.0),
                                             resultado_disponibilidad(ws.cell(row=i, column=18).value), umbral_media, precio_minimo,
                                             precio_fila_rotacion(ws.cell(row=i, column=5).value), limite_revisar)
            if clase == "precio":
                filas_bajo_precio += 1
            elif clase == "rojo":
                for cell in ws[i]: cell.fill = fill_color
                filas_baja_rotacion += 1
            elif clase == "amarillo":
                for cell in ws[i]: cell.fill = fill_no_disponible
                filas_no_disponibles += 1
            elif clase == "revisar":
                for cell in ws[i]: cell.fill = fill_revisar
                filas_revisar += 1
            if progreso_callback: progreso_callback((i - 1) / total_filas)
        
        # AJUSTE v13: Nombre original + rotacion
        nombre_base = os.path.splitext(os.path.basename(ruta_excel_quiter))[0]
        nombre_archivo = f"{nombre_base}_ROTACION.xlsx"
        ruta_final = os.path.join(ruta_salida_folder, nombre_archivo)

        ws.delete_rows(1)
        ws.delete_cols(8, 10)
        
        wb.save(ruta_final)
        wb.close()
        if limite_revisar:
            log(f"Piezas en azul para revisar (más de {limite_revisar:g} EUR la unidad): {filas_revisar}.", "success")
        log(f"Resumen rotacion: rojas por umbral {filas_baja_rotacion}, amarillas por disponibilidad {filas_no_disponibles}, sin marcar por precio inferior a {precio_minimo:g} EUR {filas_bajo_precio}.", "success")
        log(f"✅ Análisis completado. Guardado en: {ruta_final}", "success")
        if progreso_callback: progreso_callback(1)
    except Exception as e:
        log(f"❌ Error: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc())
    finally:
        filtro_rotacion_running = False
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 2. LÓGICA DE COMPARADOR DE EXCEL
# ---------------------------------------------------

def comparar_excel_task(ruta_excel_antiguo, ruta_excel_nuevo, ruta_salida_excel, color_resaltado, log_callback=None, progreso_callback=None, finalizado_callback=None, stop_event=None):
    global comparador_running
    def log(texto, tag=None):
        if log_callback: log_callback(f"[COMPARACIÓN] {texto}", tag)
    log(f"🔍 Iniciando comparación...")
    comparador_running = True
    try:
        wb_antiguo = load_workbook(ruta_excel_antiguo, data_only=True)
        ws_antiguo = wb_antiguo.active
        lineas_antiguo = set(tuple(fila) for fila in ws_antiguo.iter_rows(values_only=True) if any(fila))
        wb_nuevo = load_workbook(ruta_excel_nuevo, data_only=True)
        ws_nuevo = wb_nuevo.active
        wb_salida = openpyxl.Workbook()
        ws_salida = wb_salida.active
        fill_color = PatternFill(start_color=color_resaltado, end_color=color_resaltado, fill_type="solid")
        total_lineas = ws_nuevo.max_row
        for i, row in enumerate(ws_nuevo.iter_rows(values_only=True)):
            if stop_event and stop_event.is_set(): break
            ws_salida.append(row)
            if tuple(row) in lineas_antiguo:
                for cell in ws_salida[ws_salida.max_row]: cell.fill = fill_color
            if progreso_callback: progreso_callback((i + 1) / total_lineas)
        if not (stop_event and stop_event.is_set()):
            Path(os.path.dirname(ruta_salida_excel)).mkdir(parents=True, exist_ok=True)
            wb_salida.save(ruta_salida_excel)
            log(f"✅ Resultado guardado en: {ruta_salida_excel}", "success")
    except Exception as e:
        log(f"❌ Error: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc())
    finally:
        comparador_running = False
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 3. LÓGICA DE AUTOMATIZACIÓN MULTI-HILO (WORKER)
# ---------------------------------------------------

def guardar_resultado_excel(ruta_excel, fila_idx, resultado, intentos=4, espera=3):
    """Escribe el resultado de una referencia en el Excel; si está abierto en Excel reintenta unos segundos antes de rendirse."""
    for intento in range(intentos):
        try:
            with excel_lock:
                wb = openpyxl.load_workbook(ruta_excel)
                try:
                    wb.active.cell(row=fila_idx, column=18).value = resultado
                    wb.save(ruta_excel)
                finally:
                    wb.close()
            return
        except PermissionError:
            if intento == intentos - 1:
                raise
            time.sleep(espera)

URL_PORTAL = "https://parts.citroen.com/zzlogpl0"
INTENTOS_ACCESO = 3           # entrar al portal: 1 intento + 2 reintentos
PAUSA_REINTENTO_ACCESO = 5.0  # segundos entre intentos de acceso
ERRORES_SEGUIDOS_RECARGA = 2  # tras tantos errores seguidos se vuelve a cargar el portal
ESCALA_ESPERAS = 1.0          # multiplica los tiempos de espera del portal (las pruebas lo reducen)
credenciales_incorrectas = threading.Event()  # el portal respondió 401: paran todos los hilos


class Detenido(Exception):
    """Se pulsó Detener (o el portal rechazó el usuario): el hilo sale en cuanto lo ve."""


class CredencialesIncorrectas(Exception):
    pass


class ColaReferencias:
    """Referencias pendientes compartidas por todos los hilos: cada uno coge la siguiente libre, así que si un
    hilo no consigue entrar al portal los demás hacen su parte. Se recorre con un for normal."""

    def __init__(self, filas):
        self._cola = queue.SimpleQueue()
        for fila in filas:
            self._cola.put(fila)

    def __iter__(self):
        return self

    def __next__(self):
        try:
            return self._cola.get_nowait()
        except queue.Empty:
            raise StopIteration


def worker_playwright(worker_id, sub_filas, ruta_excel, usuario, password, log_callback, update_global_progress, stop_event):
    def log(texto, tag=None):
        log_callback(f"[HILO {worker_id}] {texto}", tag)

    def debe_parar():
        return stop_event.is_set() or credenciales_incorrectas.is_set()

    def pausa(segundos):
        fin = time.monotonic() + segundos
        while time.monotonic() < fin:
            if debe_parar():
                raise Detenido()
            time.sleep(min(0.2, max(0.0, fin - time.monotonic())))
        if debe_parar():
            raise Detenido()

    def esperar(locator, state="visible", timeout=15000):
        """wait_for en tramos de 1 s para que Detener se note enseguida."""
        fin = time.monotonic() + timeout * ESCALA_ESPERAS / 1000
        while True:
            if debe_parar():
                raise Detenido()
            restante_ms = (fin - time.monotonic()) * 1000
            try:
                locator.wait_for(state=state, timeout=max(1, min(1000, restante_ms)))
                return
            except PlaywrightTimeoutError:
                if time.monotonic() >= fin:
                    raise

    def esperar_y_hacer_click(locator, timeout=15000):
        esperar(locator, timeout=timeout)
        locator.click()

    def entrar(page):
        """Abre el portal hasta el menú de trabajo y devuelve su marco, o None si no se pudo tras varios intentos."""
        for intento in range(1, INTENTOS_ACCESO + 1):
            try:
                respuesta = page.goto(URL_PORTAL, timeout=45000, wait_until="commit")
                if respuesta is not None and respuesta.status == 401:
                    raise CredencialesIncorrectas()
                if respuesta is not None and respuesta.status >= 500:
                    raise RuntimeError(f"el portal respondió con el error {respuesta.status}")
                idioma = page.locator("text=/Espa.ol/").first
                esperar(idioma, timeout=20000)
                idioma.click()
                esperar(page.locator("text=/Referencia/").first, timeout=25000)
                for f in page.frames:
                    try:
                        if f.locator("text=/Referencia/").count() > 0:
                            return f
                    except Exception:
                        continue
                raise RuntimeError("no se encontró el marco de trabajo")
            except (Detenido, CredencialesIncorrectas):
                raise
            except Exception as e:
                if intento == INTENTOS_ACCESO:
                    log(f"Error de acceso/login tras {INTENTOS_ACCESO} intentos: {e}. Los demás hilos harán sus referencias.", "error")
                    registrar_error_detallado(str(e), traceback.format_exc())
                    return None
                log(f"No se pudo entrar al portal (intento {intento} de {INTENTOS_ACCESO}): {str(e).splitlines()[0]}. Se reintenta...", "error")
                pausa(PAUSA_REINTENTO_ACCESO)
        return None

    def campo_referencia(frame):
        return frame.locator('input[name="mara-matnr[1]"]').first

    def preparar_formulario(frame, formulario_listo):
        campo_ref = campo_referencia(frame)
        # Tras «Otro Articulo» el portal ya muestra el formulario vacío: no hace falta volver a pasar por los menús
        if not (formulario_listo and campo_ref.is_visible()):
            esperar_y_hacer_click(frame.locator("text=Referencia").first, timeout=20000)
            esperar_y_hacer_click(frame.locator("text=Mostrar disponibilidad pieza").first, timeout=20000)
            esperar(campo_ref, timeout=20000)
        campo_ref.click()
        campo_ref.fill("")
        return campo_ref

    def pasar_a_otro_articulo(frame):
        """Vuelve al formulario para la siguiente referencia; True si quedó listo."""
        try:
            esperar_y_hacer_click(frame.locator("text=Otro Articulo").first, timeout=10000)
            esperar(campo_referencia(frame), timeout=5000)
            return True
        except Detenido:
            raise
        except Exception:
            return False

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--disable-gpu", "--disable-dev-shm-usage"])
            try:
                context = browser.new_context(http_credentials={"username": usuario, "password": password})
                context.route("**/*.{png,jpg,jpeg,gif,svg,woff,woff2}", lambda route: route.abort())
                page = context.new_page()
                target_frame = entrar(page)
                if target_frame is None:
                    return
                formulario_listo = False
                errores_seguidos = 0

                for fila_idx, referencia in sub_filas:
                    if debe_parar():
                        break
                    log(f"Consultando: {referencia}")
                    try:
                        ultimo_error = None
                        for intento in range(2):
                            try:
                                campo_ref = preparar_formulario(target_frame, formulario_listo)
                                campo_ref.fill(str(referencia))
                                radio = target_frame.locator('input[type="radio"]').nth(1)
                                esperar(radio, state="attached", timeout=10000)
                                radio.click(force=True)
                                esperar_y_hacer_click(target_frame.locator("text=Ejecutar programa").first, timeout=15000)
                                break
                            except Detenido:
                                raise
                            except Exception as intento_error:
                                ultimo_error = intento_error
                                formulario_listo = False
                                if intento == 0:
                                    log(f"Reintentando {referencia}: el formulario no estaba listo.", "error")
                                    try:
                                        btn_otro = target_frame.locator("text=Otro Articulo").first
                                        if btn_otro.is_visible(timeout=3000):
                                            btn_otro.click()
                                    except Exception:
                                        pass
                                    pausa(1.2)
                                    continue
                                raise ultimo_error
                        selector = "tr.moduloTableauBlanc td"
                        esperar(target_frame.locator(selector).first, timeout=15000)
                        resultado = target_frame.locator(selector).first.inner_text().strip()
                        errores_seguidos = 0
                        try:
                            guardar_resultado_excel(ruta_excel, fila_idx, resultado)
                            log(f"OK {referencia}: {resultado}")
                        except PermissionError:
                            log(f"No se pudo guardar {referencia} porque el Excel está abierto. Ciérralo y repite esa referencia. Resultado obtenido: {resultado}", "error")
                        formulario_listo = pasar_a_otro_articulo(target_frame)
                    except Detenido:
                        raise
                    except Exception as e:
                        log(f"Error en {referencia}: {e}", "error")
                        formulario_listo = False
                        errores_seguidos += 1
                    update_global_progress()
                    if errores_seguidos >= ERRORES_SEGUIDOS_RECARGA and not debe_parar():
                        log(f"{errores_seguidos} errores seguidos: se vuelve a cargar el portal.", "error")
                        errores_seguidos = 0
                        target_frame = entrar(page)
                        if target_frame is None:
                            return
            finally:
                browser.close()
    except Detenido:
        pass
    except CredencialesIncorrectas:
        credenciales_incorrectas.set()
    except Exception as e:
        log(f"Error critico en hilo: {e}", "error")

def filas_a_consultar(filas_trabajo, precio_minimo):
    """De (fila, referencia, precio) deja solo las filas con precio válido y no inferior al mínimo, como (fila, referencia)."""
    resultado = []
    for fila_idx, referencia, valor_precio in filas_trabajo:
        try:
            if float(str(valor_precio).replace(",", ".")) >= precio_minimo:
                resultado.append((fila_idx, referencia))
        except Exception:
            continue
    return resultado

def ejecutar_automatizacion_task(ruta_excel, usuario, password, precio_minimo, num_hilos, log_callback=None, progreso_callback=None, finalizado_callback=None, stop_event=None):
    global automation_running, auto_completada_ok, auto_resultado
    automation_running = True
    auto_completada_ok = False
    auto_resultado = "error"
    cierre_notificado = False
    if stop_event:
        stop_event.clear()
    credenciales_incorrectas.clear()
    try:
        wb = openpyxl.load_workbook(ruta_excel)
        sheet = wb.active
        filas_trabajo = []
        for r in range(2, sheet.max_row + 1):
            ref = sheet.cell(row=r, column=2).value
            pre = sheet.cell(row=r, column=6).value
            if ref:
                filas_trabajo.append((r, ref, pre))
        wb.close()
        if not filas_trabajo:
            log_callback("No hay referencias en el Excel.", "error")
            return
        filas = filas_a_consultar(filas_trabajo, precio_minimo)
        total_tareas = len(filas)
        if total_tareas == 0:
            auto_completada_ok = True
            auto_resultado = "completa"
            log_callback(f"Ninguna referencia alcanza el precio mínimo ({precio_minimo:g} EUR): no hay nada que consultar.", "success")
            return

        procesadas = 0
        progreso_lock = threading.Lock()

        def update_progress():
            nonlocal procesadas, cierre_notificado
            global automation_running
            with progreso_lock:
                procesadas += 1
                progreso = procesadas / total_tareas
                deberia_cerrar = procesadas >= total_tareas and not cierre_notificado
                if deberia_cerrar:
                    cierre_notificado = True
            if progreso_callback:
                progreso_callback(progreso)
            if deberia_cerrar:
                automation_running = False
                if finalizado_callback:
                    finalizado_callback()

        hilos = max(1, min(num_hilos, total_tareas))
        pendientes = ColaReferencias(filas)
        log_callback(f"[SISTEMA] Consultando {total_tareas} de {len(filas_trabajo)} referencias (precio mínimo {precio_minimo:g} EUR) con {hilos} hilos paralelos...")
        with ThreadPoolExecutor(max_workers=hilos) as executor:
            for i in range(hilos):
                executor.submit(worker_playwright, i + 1, pendientes, ruta_excel, usuario, password, log_callback, update_progress, stop_event)
        if credenciales_incorrectas.is_set():
            auto_resultado = "credenciales"
            log_callback("❌ El portal ha rechazado el usuario o la contraseña: revísalos en la pestaña Configuración.", "error")
            return
        if stop_event.is_set():
            auto_resultado = "detenida"
            return
        if procesadas >= total_tareas:
            auto_completada_ok = True
            auto_resultado = "completa"
            log_callback("Automatizacion MULTI-HILO completada con exito.", "success")
        else:
            auto_resultado = "incompleta"
            log_callback(f"Automatización terminada con incidencias: se han procesado {procesadas} de {total_tareas} referencias (revisa los errores del registro).", "error")
    except Exception as e:
        log_callback(f"Error general: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc(), ref="ejecutar_automatizacion_task")
    finally:
        automation_running = False
        if finalizado_callback and not cierre_notificado:
            finalizado_callback()

# ---------------------------------------------------
# 4. LÓGICA DE GENERACIÓN DE PEDIDOS (CON CARPETAS POR RANGO) - v13
# ---------------------------------------------------

def generar_pedidos_task(ruta_excel, carpeta_pedidos, log_callback=None, progreso_callback=None, finalizado_callback=None, stop_event=None, numero_callback=None):
    global pedidos_running
    pedidos_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[PEDIDOS] {texto}", tag) # AJUSTE v13: Prefijo [PEDIDOS]
    try:
        ruta_backup = crear_copia_seguridad(ruta_excel)
        log(f"Backup creado: {ruta_backup}", "success")
        # 1. LIMPIEZA DE REFERENCIAS (2/) Y GUARDADO INTERMEDIO
        wb = load_workbook(ruta_excel)
        ws = wb.active
        for fila in range(2, ws.max_row + 1):
            valor = ws[f"A{fila}"].value
            if valor:
                valor_str = str(valor)
                if "2/" in valor_str:
                    ws[f"A{fila}"] = valor_str.replace("2/", "")
        wb.save(ruta_excel)
        
        # 2. LECTURA DE DATOS FILTRADOS (SOLO COLUMNAS A, C, F)
        wb_data = load_workbook(ruta_excel, data_only=True)
        ws_data = wb_data.active
        
        datos = []
        for row in ws_data.iter_rows(min_row=2, values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                ref = row[0] if len(row) > 0 else ""
                cant = row[2] if len(row) > 2 else ""
                cli = row[5] if len(row) > 5 else ""
                datos.append([ref, cant, cli])
        
        if not datos:
            log("⚠️ El Excel está vacío o no tiene datos válidos.", "error")
            return

        # 3. GESTIÓN DE CARPETAS Y GENERACIÓN DE CSV
        contador = obtener_ultimo_contador_y_ruta(carpeta_pedidos, ".csv")
        
        for i in range(0, len(datos), 100):
            if stop_event and stop_event.is_set(): break
            bloque = datos[i:i+100]
            
            contador += 1
            ruta_rango = obtener_ruta_para_contador(carpeta_pedidos, contador)
            ruta_csv = os.path.join(ruta_rango, f"{contador}.csv")
            
            with open(ruta_csv, "w", newline="", encoding="utf-8") as f:
                csv.writer(f, delimiter=";").writerows(bloque)
            
            log(f"✅ Pedido generado: {contador}.csv (en {os.path.basename(ruta_rango)})", "success")
            if numero_callback: numero_callback(str(contador))
            if progreso_callback: progreso_callback((i + len(bloque)) / len(datos))
        
        # 4. BORRADO FÍSICO FINAL DEL EXCEL
        if not (stop_event and stop_event.is_set()):
            wb_clean = load_workbook(ruta_excel)
            ws_clean = wb_clean.active
            if ws_clean.max_row > 0:
                ws_clean.delete_rows(1, ws_clean.max_row)
            wb_clean.save(ruta_excel)
            log("✅ Pedidos generados (Ref, Cant, Cliente) y Excel limpiado.", "success")
            
    except Exception as e:
        log(f"❌ Error: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc())
    finally:
        pedidos_running = False
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 4bis. LÓGICA DE HOJA DE PEDIDOS DIARIA (COPIA EN LIMPIO DE LA PLANTILLA)
# ---------------------------------------------------

MESES_ES = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
    7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE",
}

def ruta_hoja_pedidos_del_dia(carpeta_hojas, fecha=None):
    """Ruta AÑO/MES/DD-MM-AAAA.xlsx que le corresponde a una fecha dentro de la carpeta de hojas."""
    fecha = fecha or datetime.now()
    carpeta_mes = Path(carpeta_hojas) / str(fecha.year) / MESES_ES[fecha.month]
    nombre = fecha.strftime("%d-%m-%Y") + ".xlsx"
    return carpeta_mes / nombre

EXTENSIONES_INFORME = (".xlsx", ".xlsm", ".xls", ".csv")


def fecha_en_nombre(nombre):
    """Fecha dd-mm-aaaa (o dd-mm-aa, con - _ o .) del nombre de un archivo, o None."""
    encontrada = re.search(r"(\d{1,2})[-_.](\d{1,2})[-_.](\d{4}|\d{2})(?!\d)", nombre)
    if not encontrada:
        return None
    dia, mes, anio = (int(n) for n in encontrada.groups())
    try:
        return datetime(anio + 2000 if anio < 100 else anio, mes, dia).date()
    except ValueError:
        return None


def informe_diario_anterior(carpeta, hoy=None):
    """El informe más reciente anterior a hoy de la carpeta (por la fecha del nombre; si ningún nombre
    tiene fecha, por la fecha de modificación). None si no hay ninguno."""
    hoy = hoy or datetime.now().date()
    try:
        archivos = [a for a in Path(carpeta).iterdir()
                    if a.is_file() and a.suffix.lower() in EXTENSIONES_INFORME and not a.name.startswith("~$")]
    except OSError:
        return None
    con_fecha = [(fecha_en_nombre(a.name), a) for a in archivos]
    con_fecha = [(f, a) for f, a in con_fecha if f is not None and f < hoy]
    if con_fecha:
        return max(con_fecha, key=lambda par: (par[0], par[1].stat().st_mtime))[1]
    inicio_de_hoy = datetime.combine(hoy, datetime.min.time()).timestamp()
    anteriores = [a for a in archivos if fecha_en_nombre(a.name) is None and a.stat().st_mtime < inicio_de_hoy]
    return max(anteriores, key=lambda a: a.stat().st_mtime) if anteriores else None


def generar_hoja_pedidos_task(ruta_plantilla, carpeta_hojas, log_callback=None, progreso_callback=None, finalizado_callback=None, forzar=False, fecha=None):
    global hoja_pedidos_running
    hoja_pedidos_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[HOJA PEDIDOS] {texto}", tag)
    log("Generando hoja de pedidos del día a partir de la plantilla...")
    try:
        if not ruta_plantilla or not os.path.exists(ruta_plantilla):
            log("❌ No se encuentra la plantilla de hoja de pedidos.", "error")
            return
        ruta_final = ruta_hoja_pedidos_del_dia(carpeta_hojas, fecha)
        if progreso_callback: progreso_callback(0.2)
        if ruta_final.exists() and not forzar:
            log(f"La hoja de hoy ya existía, no se ha sobrescrito: {ruta_final}", "success")
            if progreso_callback: progreso_callback(1)
            return
        ruta_final.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ruta_plantilla, ruta_final)
        if progreso_callback: progreso_callback(1)
        log(f"✅ Hoja de pedidos generada: {ruta_final}", "success")
    except Exception as e:
        log(f"❌ Error: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_hoja_pedidos_task")
    finally:
        hoja_pedidos_running = False
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 4bis-2. AVISOS A CLIENTES DE REFERENCIAS QUE NO SUMINISTRAN
# ---------------------------------------------------

def _normalizar_texto_estado(valor):
    return " ".join(str(valor).strip().upper().split())

def _normalizar_cliente_dms(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()

def _columnas_clientes(cabecera, columnas):
    """Índices (nombre, email) de clientes.xlsx. Se buscan por la cabecera (CORREO/EMAIL, NOMBRE/TALLER); si no se
    reconoce: A=num_cliente, B=email con dos columnas, o A=num_cliente, B=nombre, C=email con tres."""
    titulos = [str(valor or "").strip().upper() for valor in cabecera]
    email = next((i for i, t in enumerate(titulos) if i > 0 and ("CORREO" in t or "MAIL" in t)), None)
    nombre = next((i for i, t in enumerate(titulos) if i > 0 and ("NOMBRE" in t or "TALLER" in t)), None)
    if email is None:
        email = 1 if columnas <= 2 else 2
        if nombre is None and columnas > 2:
            nombre = 1
    return nombre, email

def cargar_clientes_excel(ruta_clientes):
    """Lee clientes.xlsx -> {num_cliente: {"nombre", "email"}}. Vale con A=num_cliente, B=email (sin nombre del taller)
    o con A=num_cliente, B=nombre, C=email; varios emails separados por ';'."""
    clientes = {}
    wb = openpyxl.load_workbook(ruta_clientes, data_only=True)
    ws = wb.active
    cabecera = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    col_nombre, col_email = _columnas_clientes(cabecera, ws.max_column)

    def celda(fila, indice):
        if indice is None or len(fila) <= indice or fila[indice] in (None, ""):
            return ""
        return str(fila[indice]).strip()

    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or fila[0] is None:
            continue
        num = _normalizar_cliente_dms(fila[0])
        if not num:
            continue
        clientes[num] = {"nombre": celda(fila, col_nombre), "email": celda(fila, col_email)}
    wb.close()
    return clientes

def nombre_cliente_aviso(entrada):
    """'TALLER (8049)' si el Excel de clientes trae el nombre; si no, 'Cliente 8049'."""
    return f"{entrada['nombre']} ({entrada['cliente']})" if entrada.get("nombre") else f"Cliente {entrada['cliente']}"

def extraer_avisos(ruta_hoja, log=lambda *a, **k: None):
    """Lee la hoja de pedido en streaming (una sola pasada) y devuelve (no_suministra, sustituciones), agrupados por cliente DMS.
    no_suministra: {cliente: [referencia, ...]} para ESTADO 'NO SUMINISTRA'.
    sustituciones: {cliente: [(referencia_original, referencia_sustituta), ...]} con la casilla SUSTITUCIÓN rellena y sin NO SUMINISTRA.
    Las hojas suelen declarar más de un millón de filas, por eso no se usa ws.max_row ni acceso celda a celda."""
    wb = openpyxl.load_workbook(ruta_hoja, data_only=True, read_only=True)
    no_suministra = {}
    sustituciones = {}
    try:
        for ws in wb.worksheets:
            pedido = None
            encontradas_no = 0
            encontradas_sus = 0
            cabecera_correcta = False
            for numero_fila, fila in enumerate(ws.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
                if numero_fila == 2:
                    pedido = fila[4] if len(fila) > 4 else None
                    continue
                if numero_fila == 4:
                    cabecera = fila[1] if len(fila) > 1 else None
                    cabecera_correcta = bool(cabecera) and "REFERENCIA" in str(cabecera).upper()
                if numero_fila < 5:
                    continue
                if not fila:
                    continue
                if len(fila) < 8:
                    fila = tuple(fila) + (None,) * (8 - len(fila))
                referencia, estado, sustitucion, cliente = fila[1], fila[4], fila[5], fila[7]
                if not referencia:
                    continue
                if estado and "NO SUMINISTRA" in _normalizar_texto_estado(estado):
                    cliente_norm = _normalizar_cliente_dms(cliente)
                    if cliente_norm:
                        no_suministra.setdefault(cliente_norm, []).append(str(referencia).strip())
                        encontradas_no += 1
                elif sustitucion not in (None, ""):
                    cliente_norm = _normalizar_cliente_dms(cliente)
                    if cliente_norm:
                        sustituciones.setdefault(cliente_norm, []).append((str(referencia).strip(), str(sustitucion).strip()))
                        encontradas_sus += 1
            if not cabecera_correcta:
                log(f"⚠️ La pestaña '{ws.title}' no tiene el formato esperado (falta la cabecera «REFERENCIAS A PEDIR» en B4); puede que no se detecte nada.", "error")
            if encontradas_no:
                log(f"Pestaña '{ws.title}' (pedido {pedido}): {encontradas_no} referencia(s) NO SUMINISTRA.")
            if encontradas_sus:
                log(f"Pestaña '{ws.title}' (pedido {pedido}): {encontradas_sus} referencia(s) sustituida(s).")
    finally:
        wb.close()
    return no_suministra, sustituciones

def construir_cuerpo_aviso(referencias, fecha_str):
    if len(referencias) == 1:
        return (
            "Buenas tardes.\n\n"
            f"La referencia {referencias[0]} correspondiente al pedido del día {fecha_str} no suministra.\n\n"
            "Gracias por su atención,"
        )
    lineas = "\n".join(f"- {ref}" for ref in referencias)
    return (
        "Buenas tardes.\n\n"
        f"Las siguientes referencias correspondientes al pedido del día {fecha_str} no suministran:\n\n"
        f"{lineas}\n\n"
        "Gracias por su atención,"
    )

def construir_cuerpo_sustitucion(sustituciones, fecha_str):
    if len(sustituciones) == 1:
        original, final = sustituciones[0]
        return (
            "Buenas tardes.\n\n"
            f"La referencia {original} correspondiente al pedido del día {fecha_str} ha sido sustituida por la referencia {final}.\n\n"
            "Gracias por su atención,"
        )
    lineas = "\n".join(f"- {original} → {final}" for original, final in sustituciones)
    return (
        "Buenas tardes.\n\n"
        f"Las siguientes referencias correspondientes al pedido del día {fecha_str} han sido sustituidas:\n\n"
        f"{lineas}\n\n"
        "Gracias por su atención,"
    )

def crear_borrador_outlook(outlook, destinatarios, asunto, cuerpo_texto):
    """Crea un borrador en Outlook de escritorio (no lo envía) respetando la firma por defecto del usuario.
    'outlook' es la aplicación Outlook ya conectada (COM ya inicializado en el hilo llamante)."""
    mail = outlook.CreateItem(0)  # olMailItem
    mail.Subject = asunto
    if destinatarios:
        mail.To = destinatarios
    mail.GetInspector  # fuerza a Outlook a cargar la firma por defecto de "mensaje nuevo"
    firma_html = mail.HTMLBody or ""
    cuerpo_html = cuerpo_texto.replace("\n", "<br>")
    mail.HTMLBody = f"<div style=\"font-family:Calibri,Arial,sans-serif;font-size:11pt\">{cuerpo_html}</div><br>{firma_html}"
    mail.Save()

def enviar_correo_outlook_aviso(outlook, destinatarios, asunto, cuerpo_texto):
    """Igual que crear_borrador_outlook pero envía el correo directamente (olMailItem.Send)."""
    mail = outlook.CreateItem(0)  # olMailItem
    mail.Subject = asunto
    if destinatarios:
        mail.To = destinatarios
    mail.GetInspector  # fuerza a Outlook a cargar la firma por defecto de "mensaje nuevo"
    firma_html = mail.HTMLBody or ""
    cuerpo_html = cuerpo_texto.replace("\n", "<br>")
    mail.HTMLBody = f"<div style=\"font-family:Calibri,Arial,sans-serif;font-size:11pt\">{cuerpo_html}</div><br>{firma_html}"
    mail.Send()

def generar_avisos_no_suministro_task(ruta_hoja, ruta_clientes, enviar_automaticamente=False, emails_companeros="", log_callback=None, progreso_callback=None, finalizado_callback=None, resultado_callback=None):
    global avisos_running
    avisos_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[AVISOS] {texto}", tag)
    log("Analizando la hoja de pedido en busca de referencias que no suministran o han sido sustituidas...")
    resultado = {
        "no_suministra": {"enviados": [], "borradores": [], "sin_correo": []},
        "sustitucion": {"enviados": [], "borradores": [], "sin_correo": []},
        "hoja_companeros_enviada": False,
        "error": None,
    }
    try:
        match = re.search(r"\d{2}-\d{2}-\d{4}", Path(ruta_hoja).stem)
        fecha_str = match.group(0) if match else Path(ruta_hoja).stem
        try:
            clientes = cargar_clientes_excel(ruta_clientes)
        except Exception as e:
            log(f"❌ No se pudo leer el Excel de clientes: {e}", "error")
            registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_avisos_no_suministro_task_clientes")
            resultado["error"] = str(e)
            return
        try:
            agrupado_no_suministra, agrupado_sustituciones = extraer_avisos(ruta_hoja, log)
        except Exception as e:
            log(f"❌ No se pudo leer la hoja de pedido: {e}", "error")
            registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_avisos_no_suministro_task_hoja")
            resultado["error"] = str(e)
            return
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        try:
            try:
                outlook = win32com.client.Dispatch("Outlook.Application")
            except Exception as e:
                log(f"❌ No se pudo conectar con Outlook: {e}. Abre Outlook de escritorio, comprueba que tiene una cuenta configurada y vuelve a intentarlo.", "error")
                registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_avisos_no_suministro_task_outlook")
                resultado["error"] = str(e)
                return
            if emails_companeros:
                try:
                    cuerpo_companeros = "Buenas tardes.\n\nAdjunto hoja de pedido.\n\n"
                    enviar_correo_outlook_madd(outlook, emails_companeros, f"Hoja de pedido del día {fecha_str}", cuerpo_companeros, adjuntos=[ruta_hoja])
                    resultado["hoja_companeros_enviada"] = True
                    log("✅ Hoja de pedido enviada a los compañeros.", "success")
                except Exception as e:
                    log(f"❌ Error enviando la hoja de pedido a los compañeros: {e}", "error")
                    registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_avisos_no_suministro_task_companeros")
            if not agrupado_no_suministra and not agrupado_sustituciones:
                log("No se han encontrado referencias con estado NO SUMINISTRA ni sustituciones.", "success")
                return
            total = len(agrupado_no_suministra) + len(agrupado_sustituciones)
            hecho = 0

            def procesar_cliente(cliente_num, email, nombre, asunto, cuerpo, detalle, bucket, tipo_label, ref_sufijo):
                nonlocal hecho
                enviar_directo = enviar_automaticamente and bool(email)
                try:
                    if enviar_directo:
                        enviar_correo_outlook_aviso(outlook, email, asunto, cuerpo)
                    else:
                        crear_borrador_outlook(outlook, email, asunto, cuerpo)
                except Exception as e:
                    accion = "enviando el correo" if enviar_directo else "creando el borrador"
                    log(f"❌ Error {accion} ({tipo_label}) para el cliente {cliente_num}: {e}", "error")
                    registrar_error_detallado(str(e), traceback.format_exc(), ref=f"generar_avisos_no_suministro_task_{ref_sufijo}_{cliente_num}")
                    hecho += 1
                    if progreso_callback: progreso_callback(hecho / total)
                    return
                entrada = {"cliente": cliente_num, "nombre": nombre, "email": email, "detalle": detalle}
                if enviar_directo:
                    bucket["enviados"].append(entrada)
                    log(f"✅ Correo enviado ({tipo_label}) a {nombre or f'cliente {cliente_num}'} ({len(detalle)} ref.)", "success")
                elif email:
                    bucket["borradores"].append(entrada)
                    log(f"✅ Borrador creado ({tipo_label}) para {nombre or f'cliente {cliente_num}'} ({len(detalle)} ref.)", "success")
                else:
                    bucket["sin_correo"].append(entrada)
                    log(f"⚠️ Cliente {cliente_num} sin email en el Excel de clientes ({tipo_label}); borrador creado sin destinatario.", "error")
                hecho += 1
                if progreso_callback: progreso_callback(hecho / total)

            for cliente_num, referencias in agrupado_no_suministra.items():
                info = clientes.get(cliente_num)
                nombre = info["nombre"] if info else ""
                email = info["email"] if info else ""
                asunto = f"Aviso referencia no suministrada {fecha_str}"
                cuerpo = construir_cuerpo_aviso(referencias, fecha_str)
                procesar_cliente(cliente_num, email, nombre, asunto, cuerpo, referencias, resultado["no_suministra"], "no suministra", "nosum")

            for cliente_num, sustituciones in agrupado_sustituciones.items():
                info = clientes.get(cliente_num)
                nombre = info["nombre"] if info else ""
                email = info["email"] if info else ""
                asunto = f"Aviso referencia sustituida {fecha_str}"
                cuerpo = construir_cuerpo_sustitucion(sustituciones, fecha_str)
                detalle = [f"{original} → {final}" for original, final in sustituciones]
                procesar_cliente(cliente_num, email, nombre, asunto, cuerpo, detalle, resultado["sustitucion"], "sustitución", "sust")

            for etiqueta, bucket in (("no suministra", resultado["no_suministra"]), ("sustitución", resultado["sustitucion"])):
                partes = []
                if bucket["enviados"]:
                    partes.append(f"{len(bucket['enviados'])} correo(s) enviado(s)")
                if bucket["borradores"]:
                    partes.append(f"{len(bucket['borradores'])} borrador(es) con destinatario")
                if bucket["sin_correo"]:
                    partes.append(f"{len(bucket['sin_correo'])} sin email (borrador sin destinatario)")
                if partes:
                    log(f"✨ {etiqueta}: {', '.join(partes)}.", "success")
        finally:
            pythoncom.CoUninitialize()
    except Exception as e:
        resultado["error"] = str(e)
        log(f"❌ Error: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc(), ref="generar_avisos_no_suministro_task")
    finally:
        avisos_running = False
        if resultado_callback: resultado_callback(resultado)
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 5. LÓGICA DE FICHERO MOPAR (CON NOMBRES FXXXX + LETRAS) - v13
# ---------------------------------------------------

def generar_fichero_mopar_task(ruta_excel, carpeta_salida, log_callback=None, progreso_callback=None, finalizado_callback=None, stop_event=None, numero_callback=None):
    global mopar_running
    mopar_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[MOPAR] {texto}", tag)
    log(f"🚀 Iniciando generación de ficheros MOPAR...")
    try:
        log("📖 Leyendo archivo Excel...")
        try:
            wb = xlrd.open_workbook(ruta_excel)
            ws = wb.sheet_by_index(0)
        except Exception as e:
            log(f"❌ Error al abrir el Excel: {e}", "error")
            return
        headers = [str(ws.cell_value(0, i)).strip() for i in range(ws.ncols)]
        try:
            col_ref = headers.index("Referencia")
            col_cant = headers.index("Cant")
            col_cliente = headers.index("Cliente")
        except ValueError as e:
            col_faltante = str(e).split()[-1]
            log(f"❌ Error: No se encuentra la columna {col_faltante} en el Excel.", "error")
            return
        total_filas = ws.nrows - 1
        if total_filas <= 0:
            log("⚠️ El Excel no tiene datos de artículos.", "error")
            return
        lineas_fichero = []
        log(f"🔍 Procesando {total_filas} filas...")
        for r in range(1, ws.nrows):
            if stop_event and stop_event.is_set(): break
            try:
                ref_raw = str(ws.cell_value(r, col_ref)).strip()
                ref_clean = ref_raw.replace("2/", "")
                referencia = ref_clean.zfill(13)[:13]
                try:
                    cant_val = int(float(ws.cell_value(r, col_cant)))
                except Exception:
                    cant_val = 0
                    log(f"⚠️ Cantidad no válida en fila {r+1}, se usará 0.", "error")
                cantidad = str(cant_val).zfill(7)[:7]
                cliente_raw = str(ws.cell_value(r, col_cliente)).strip()
                cliente_nums = "".join(filter(str.isdigit, cliente_raw))
                ubicacion = cliente_nums.zfill(8)[:8]
                linea = f"{referencia}{cantidad}{ubicacion}"
                lineas_fichero.append(linea)
            except Exception as e:
                log(f"⚠️ Error en fila {r+1}: {e}", "error")
            if progreso_callback: progreso_callback(r / total_filas)
        
        if not (stop_event and stop_event.is_set()):
            # AJUSTE v13: Lógica FXXXX + Letras
            if not os.path.exists(carpeta_salida):
                os.makedirs(carpeta_salida, exist_ok=True)

            proximo_f = obtener_numero_mopar_del_dia(carpeta_salida)
            letras = string.ascii_uppercase # A, B, C...
            archivos_creados = 0
            
            for i in range(0, len(lineas_fichero), 10):
                grupo = lineas_fichero[i:i+10]
                # Si solo hay un archivo, no lleva letra. Si hay varios, llevan A, B, C...
                if len(lineas_fichero) <= 10:
                    nombre_fichero = f"F{proximo_f}.txt"
                else:
                    indice_letra = i // 10
                    letra = letras[indice_letra % len(letras)]
                    nombre_fichero = f"F{proximo_f}{letra}.txt"
                
                ruta_final = os.path.join(carpeta_salida, nombre_fichero)
                with open(ruta_final, "w", encoding="utf-8") as f:
                    f.write("\n".join(grupo))
                
                log(f"✅ Generado MOPAR: {nombre_fichero}", "success")
                if numero_callback: numero_callback(nombre_fichero)
                archivos_creados += 1

            log(f"✨ Proceso completado. Se han generado {archivos_creados} archivos con prefijo F{proximo_f}.", "success")
    except Exception as e:
        log(f"❌ Error crítico: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc())
    finally:
        mopar_running = False
        if finalizado_callback: finalizado_callback()

# ---------------------------------------------------
# 5bis. LÓGICA DE FACTURAS STELLANTIS MADD (DESCARGA DIARIA)
# ---------------------------------------------------
# Estructura de la web confirmada por el usuario el 2026-09-21 (login SSO en
# idfed.mpsa.com, selector de perfil ng-select, tabla de documentos "eDoc-Partners").
# Solo lectura y descarga: nunca se pulsan los botones de estado (leído/impreso/etc).

MADD_URL_DOCUMENTOS = "https://www.madd.stellantis.com/documents"
MADD_PERFIL_TEXTO = "Piezas de recambio"
LOG_FACTURAS_MADD = ruta_datos("log_facturas_madd.txt", migrar=False)
MADD_PAGINAS_MAX = 10  # 100 docs/página -> hasta 1000 documentos más recientes revisados

def _dia_laborable_anterior(fecha):
    """Día laborable anterior a 'fecha', saltando sábados y domingos (los lunes devuelve el viernes)."""
    anterior = fecha - timedelta(days=1)
    while anterior.weekday() >= 5:  # 5=sábado, 6=domingo
        anterior -= timedelta(days=1)
    return anterior

def registrar_log_facturas_madd(fecha_objetivo, resultado):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(LOG_FACTURAS_MADD, "a", encoding="utf-8") as f:
        f.write(f"--- EJECUCIÓN [{timestamp}] ---\n")
        f.write(f"Fecha objetivo: {fecha_objetivo.isoformat()}\n")
        f.write(f"Nº facturas descargadas: {len(resultado.get('facturas', []))}\n")
        if resultado.get("error"):
            f.write(f"Error: {resultado['error']} (paso: {resultado.get('paso_fallido') or 'desconocido'})\n")
        else:
            f.write("Sin errores.\n")
        f.write("-" * 30 + "\n\n")

def _madd_login(page, usuario, password):
    page.goto(MADD_URL_DOCUMENTOS, timeout=45000)
    page.wait_for_selector("#username", timeout=30000)
    page.fill("#username", usuario)
    page.fill("#password", password)
    page.click("#signOnButton")
    page.wait_for_load_state("networkidle", timeout=30000)

def _madd_seleccionar_perfil(page):
    selector_perfil = page.locator('ng-select[formcontrolname="profile"]')
    # Si ya hay un perfil activo de una sesión previa, esta pantalla no aparece.
    if selector_perfil.count() == 0 or not selector_perfil.first.is_visible():
        return
    selector_perfil.first.click()
    opcion = page.locator(".ng-dropdown-panel .ng-option", has_text=MADD_PERFIL_TEXTO).first
    opcion.wait_for(state="visible", timeout=10000)
    opcion.click()
    boton_confirmar = page.locator('button:has(i.bi-check-circle):not([disabled])').first
    boton_confirmar.wait_for(state="visible", timeout=10000)
    boton_confirmar.click()
    page.wait_for_selector("tr.border-bottom", timeout=20000)

def _madd_procesar_facturas_del_dia(page, fechas_objetivo, carpeta_destino, log):
    """Recorre las páginas de resultados, filtra solo por los títulos reales 'Factura ...' (ignora
    cualquier otro elemento/fila duplicada del DOM) y descarga cada factura en el momento en que se
    encuentra, para no depender de índices que quedarían obsoletos tras cambiar de página."""
    try:
        page.select_option('select[name="pageSize"]', label="100 elementos por página")
        page.wait_for_timeout(1000)
    except Exception:
        log("No se pudo ampliar el tamaño de página, se sigue con el valor por defecto.", "error")
    patron = re.compile(r"Factura\s+(\S+)\s+(\d{2}/\d{2}/\d{4})")
    facturas_descargadas = []
    for pagina in range(1, MADD_PAGINAS_MAX + 1):
        titulos = page.locator("a").filter(has_text=re.compile(r"^Factura\s"))
        total_titulos = titulos.count()
        log(f"Página {pagina}/{MADD_PAGINAS_MAX}: {total_titulos} factura(s) en la página, revisando...")
        coincidencias_pagina = 0
        for i in range(total_titulos):
            enlace = titulos.nth(i)
            try:
                texto_titulo = enlace.inner_text(timeout=3000).strip()
            except Exception:
                continue
            match = patron.search(texto_titulo)
            if not match:
                continue
            numero, fecha_str = match.group(1), match.group(2)
            dia, mes, anio = fecha_str.split("/")
            fecha_factura = datetime(int(anio), int(mes), int(dia)).date()
            if fecha_factura not in fechas_objetivo:
                continue
            coincidencias_pagina += 1
            nombre_archivo = f"Factura_{numero}_{fecha_factura.isoformat()}.pdf"
            ruta_destino = Path(carpeta_destino) / nombre_archivo
            if ruta_destino.exists():
                facturas_descargadas.append(str(ruta_destino))
                log(f"Ya existía {nombre_archivo}, no se ha vuelto a descargar.", "success")
                continue
            try:
                fila = enlace.locator("xpath=ancestor::tr[1]")
                boton_descarga = fila.locator("td.text-center button").first
                with page.expect_download(timeout=30000) as info_descarga:
                    boton_descarga.click()
                descarga = info_descarga.value
                descarga.save_as(str(ruta_destino))
                facturas_descargadas.append(str(ruta_destino))
                log(f"✅ Descargada: {nombre_archivo}", "success")
            except Exception as e:
                log(f"❌ Error descargando {nombre_archivo}: {e}", "error")
                raise
        log(f"Página {pagina}/{MADD_PAGINAS_MAX} revisada. {coincidencias_pagina} coincidencia(s) en esta página.")
        if pagina >= MADD_PAGINAS_MAX:
            break
        boton_siguiente = page.locator('a[aria-label="Next"]')
        if boton_siguiente.count() == 0:
            log("No hay botón 'Next', se detiene la paginación.")
            break
        li_padre = boton_siguiente.first.locator("xpath=..")
        if "disabled" in (li_padre.get_attribute("class") or ""):
            log("Botón 'Next' deshabilitado (última página), se detiene la paginación.")
            break
        log(f"Pasando a la página {pagina + 1}...")
        boton_siguiente.first.click()
        page.wait_for_timeout(1500)
    return facturas_descargadas

def ejecutar_descarga_facturas_madd(usuario, password, carpeta_destino, fechas_objetivo, log):
    Path(carpeta_destino).mkdir(parents=True, exist_ok=True)
    resultado = {"facturas": [], "error": None, "paso_fallido": None}
    ultimo_error = None
    for intento in range(3):
        resultado["paso_fallido"] = None
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, args=["--disable-gpu", "--disable-dev-shm-usage"])
                try:
                    context = browser.new_context(accept_downloads=True)
                    page = context.new_page()
                    try:
                        _madd_login(page, usuario, password)
                    except Exception:
                        resultado["paso_fallido"] = "acceso/login"
                        raise
                    try:
                        _madd_seleccionar_perfil(page)
                    except Exception:
                        resultado["paso_fallido"] = "cambio de perfil"
                        raise
                    try:
                        resultado["facturas"] = _madd_procesar_facturas_del_dia(page, fechas_objetivo, carpeta_destino, log)
                    except Exception:
                        resultado["paso_fallido"] = resultado["paso_fallido"] or "listado/descarga de documentos"
                        raise
                    return resultado
                finally:
                    browser.close()
        except Exception as e:
            ultimo_error = e
            log(f"⚠️ Intento {intento + 1}/3 fallido en '{resultado['paso_fallido']}': {e}", "error")
            registrar_error_detallado(str(e), traceback.format_exc(), ref="facturas_madd")
            if intento < 2:
                time.sleep(5)
    resultado["error"] = str(ultimo_error)
    return resultado

def enviar_correo_outlook_madd(outlook, destinatarios, asunto, cuerpo_texto, adjuntos=None):
    """Crea y ENVÍA (Send, no Save) un correo desde Outlook de escritorio, respetando la firma por defecto."""
    mail = outlook.CreateItem(0)  # olMailItem
    mail.Subject = asunto
    if destinatarios:
        mail.To = destinatarios
    mail.GetInspector  # fuerza a Outlook a cargar la firma por defecto de "mensaje nuevo"
    firma_html = mail.HTMLBody or ""
    cuerpo_html = cuerpo_texto.replace("\n", "<br>")
    mail.HTMLBody = f"<div style=\"font-family:Calibri,Arial,sans-serif;font-size:11pt\">{cuerpo_html}</div><br>{firma_html}"
    if adjuntos:
        for ruta in adjuntos:
            mail.Attachments.Add(str(ruta))
    mail.Send()

def outlook_en_ejecucion():
    """True si hay un Outlook de escritorio abierto en esta sesión."""
    try:
        salida = subprocess.run(["tasklist", "/FI", "IMAGENAME eq OUTLOOK.EXE", "/NH"], capture_output=True, timeout=30,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        return b"OUTLOOK.EXE" in salida.stdout.upper()
    except Exception:
        return False

def asegurar_outlook_abierto(log, espera_max=120):
    """Si Outlook está cerrado lo abre (como haría el usuario, para que siga enviando y sincronizando)
    y espera a que responda. Devuelve el objeto Outlook.Application. Se deja abierto al terminar."""
    import win32com.client
    if not outlook_en_ejecucion():
        log("Outlook estaba cerrado: abriéndolo para poder enviar el correo...")
        try:
            os.startfile("outlook.exe")
        except Exception as e:
            log(f"⚠️ No se pudo abrir Outlook de forma directa ({e}); se intentará arrancarlo por COM.", "error")
    limite = time.time() + espera_max
    ultimo_error = None
    while time.time() < limite:
        try:
            outlook = win32com.client.Dispatch("Outlook.Application")
            outlook.GetNamespace("MAPI").GetDefaultFolder(4)  # bandeja de salida: confirma que el perfil ya está cargado
            return outlook
        except Exception as e:
            ultimo_error = e
            time.sleep(3)
    raise RuntimeError(f"Outlook no respondió en {espera_max} s: {ultimo_error}")

def esperar_salida_de_correo(outlook, asunto, desde, espera_max=180, intervalo=3):
    """Espera a que el correo con ese asunto, creado desde 'desde', salga de la bandeja de salida.
    Devuelve True si salió y False si sigue pendiente al agotar el tiempo."""
    ns = outlook.GetNamespace("MAPI")
    limite = time.time() + espera_max
    empujado = False
    inicio = time.time()
    while time.time() < limite:
        pendiente = False
        try:
            for item in ns.GetDefaultFolder(4).Items:
                try:
                    if (item.Subject or "") == asunto and item.CreationTime.timestamp() >= desde - 5:
                        pendiente = True
                        break
                except Exception:
                    continue
        except Exception:
            pendiente = True
        if not pendiente:
            return True
        if not empujado and time.time() - inicio > 15:
            empujado = True
            try:
                ns.SyncObjects.Item(1).Start()  # fuerza un ciclo de enviar/recibir
            except Exception:
                pass
        time.sleep(intervalo)
    return False

def registrar_estado_correo_madd(fecha_objetivo, estado, detalle=""):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    try:
        with open(LOG_FACTURAS_MADD, "a", encoding="utf-8") as f:
            f.write(f"--- CORREO [{timestamp}] ---\n")
            f.write(f"Fecha objetivo: {fecha_objetivo.isoformat()}\n")
            f.write(f"Estado: {estado}\n")
            if detalle:
                f.write(f"{detalle}\n")
            f.write("-" * 30 + "\n\n")
    except OSError:
        pass

def enviar_con_outlook_abierto(destinatarios, asunto, cuerpo, adjuntos, log):
    """Abre Outlook si hace falta, envía el correo y espera a que salga de la bandeja de salida.
    Devuelve 'enviado' o 'pendiente' (sigue en la bandeja de salida; saldrá cuando Outlook pueda enviarlo)."""
    import pythoncom
    pythoncom.CoInitialize()
    try:
        outlook = asegurar_outlook_abierto(log)
        creado = time.time()
        enviar_correo_outlook_madd(outlook, destinatarios, asunto, cuerpo, adjuntos=adjuntos)
        if esperar_salida_de_correo(outlook, asunto, creado):
            return "enviado"
        return "pendiente"
    finally:
        pythoncom.CoUninitialize()

def _madd_enviar_correo_resultado(email_destino, fecha_objetivo, rutas_facturas, log):
    fecha_str = fecha_objetivo.strftime("%Y-%m-%d")
    asunto = f"Facturas Stellantis del {fecha_str}"
    if rutas_facturas:
        cuerpo = f"Adjunto las facturas del día {fecha_str}. Total: {len(rutas_facturas)} facturas."
    else:
        cuerpo = f"No hay facturas del día {fecha_str}"
    estado = enviar_con_outlook_abierto(email_destino, asunto, cuerpo, rutas_facturas or None, log)
    if estado == "enviado":
        log(f"✅ Correo enviado a {email_destino} ({len(rutas_facturas)} factura(s) adjuntas).", "success")
        registrar_estado_correo_madd(fecha_objetivo, "enviado", f"Destinatarios: {email_destino} | adjuntos: {len(rutas_facturas)}")
    else:
        log(f"⚠️ El correo para {email_destino} sigue en la bandeja de salida de Outlook; saldrá en cuanto Outlook pueda enviarlo.", "error")
        registrar_estado_correo_madd(fecha_objetivo, "PENDIENTE en la bandeja de salida de Outlook", f"Destinatarios: {email_destino} | adjuntos: {len(rutas_facturas)}")

def _madd_enviar_aviso_fallo(email_aviso, fecha_objetivo, resultado, log):
    if not email_aviso:
        log("No hay correo de aviso de fallo configurado; no se envía ningún aviso.", "error")
        return
    fecha_str = fecha_objetivo.strftime("%Y-%m-%d")
    paso = resultado.get("paso_fallido") or "desconocido"
    error = resultado.get("error") or "error desconocido"
    asunto = f"[AVISO] Fallo en descarga de facturas Stellantis del {fecha_str}"
    cuerpo = f"La descarga automática de facturas Stellantis del día {fecha_str} ha fallado tras 3 intentos.\n\nPaso que falló: {paso}\nError: {error}"
    estado = enviar_con_outlook_abierto(email_aviso, asunto, cuerpo, None, log)
    if estado == "enviado":
        log(f"✅ Aviso de fallo enviado a {email_aviso}.", "error")
    else:
        log(f"⚠️ El aviso de fallo para {email_aviso} sigue en la bandeja de salida de Outlook.", "error")
    registrar_estado_correo_madd(fecha_objetivo, f"aviso de fallo: {estado}", f"Destinatarios: {email_aviso}")

def tarea_facturas_madd(usuario, password, carpeta_destino, email_destino, email_aviso_fallo, filtro_dia="Ayer", log_callback=None, finalizado_callback=None, resultado_callback=None, fecha_objetivo=None):
    global madd_running
    madd_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[FACTURAS MADD] {texto}", tag)
    log("Iniciando descarga de facturas Stellantis MADD...")
    hoy = datetime.now().date()
    fecha_objetivo = fecha_objetivo or (hoy if filtro_dia == "Hoy" else _dia_laborable_anterior(hoy))
    resultado = {"facturas": [], "error": None, "paso_fallido": None}
    try:
        resultado = ejecutar_descarga_facturas_madd(usuario, password, carpeta_destino, {fecha_objetivo}, log)
        registrar_log_facturas_madd(fecha_objetivo, resultado)
        if resultado.get("error"):
            log(f"❌ Fallo tras 3 intentos en el paso '{resultado.get('paso_fallido')}': {resultado['error']}", "error")
            _madd_enviar_aviso_fallo(email_aviso_fallo, fecha_objetivo, resultado, log)
            return
        _madd_enviar_correo_resultado(email_destino, fecha_objetivo, resultado["facturas"], log)
        log(f"✨ Proceso completado: {len(resultado['facturas'])} factura(s) del {fecha_objetivo.isoformat()}.", "success")
    except Exception as e:
        resultado["error"] = str(e)
        resultado["paso_fallido"] = resultado.get("paso_fallido") or "envío de correo"
        log(f"❌ Error crítico: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc(), ref="tarea_facturas_madd")
        registrar_log_facturas_madd(fecha_objetivo, resultado)
        try:
            _madd_enviar_aviso_fallo(email_aviso_fallo, fecha_objetivo, resultado, log)
        except Exception as e2:
            log(f"❌ No se pudo enviar el aviso de fallo: {e2}", "error")
    finally:
        madd_running = False
        detalle = "falló" if resultado.get("error") else f"{len(resultado.get('facturas') or [])} facturas del {fecha_objetivo:%d/%m}"
        registrar_ultima_ejecucion("madd", detalle)
        if resultado_callback: resultado_callback(resultado)
        if finalizado_callback: finalizado_callback()

def ejecutar_tarea_madd_standalone():
    """Punto de entrada para la tarea programada de Windows (--tarea-madd), sin GUI. Devuelve el código de
    salida: 0 si terminó bien y 1 si falló o falta configuración, para que el Programador de tareas lo muestre."""
    config = cargar_configuracion()
    usuario = config.get("usuario_madd", "")
    password = config.get("password_madd", "")
    carpeta = config.get("ruta_carpeta_facturas_madd", "")
    email_destino = config.get("email_facturas_madd", "")
    email_aviso = config.get("email_aviso_fallo_madd", "")
    filtro_dia = config.get("filtro_dia_facturas_madd", "Ayer")
    if not usuario or not password or not carpeta or not email_destino:
        registrar_error_detallado("Configuración incompleta para la tarea MADD (usuario/password/carpeta/email).", "", ref="madd_standalone")
        return 1
    resultados = []
    tarea_facturas_madd(usuario, password, carpeta, email_destino, email_aviso, filtro_dia=filtro_dia, log_callback=None, resultado_callback=resultados.append)
    return 1 if not resultados or resultados[-1].get("error") else 0

# ---------- «Última vez» de cada tarea (se guarda en la configuración, clave ultimas_ejecuciones) ----------

def registrar_ultima_ejecucion(clave, detalle="", ahora=None):
    """Apunta cuándo terminó una tarea. Lo usan la app y la tarea programada de facturas."""
    try:
        completa = cargar_configuracion()
        ultimas = dict(completa.get("ultimas_ejecuciones") or {})
        ultimas[clave] = {"cuando": (ahora or datetime.now()).isoformat(timespec="seconds"), "detalle": detalle}
        completa["ultimas_ejecuciones"] = ultimas
        guardar_configuracion(completa)
        return ultimas[clave]
    except Exception as e:
        registrar_error_detallado(f"No se pudo guardar la última ejecución de {clave}: {e}", traceback.format_exc())
        return None


def texto_ultima_vez(info, ahora=None):
    """'Última vez: hoy 09:42 · 24 de 46 líneas' (ayer / dd/mm para días anteriores); '' si nunca."""
    if not info or not info.get("cuando"):
        return ""
    try:
        cuando = datetime.fromisoformat(info["cuando"])
    except ValueError:
        return ""
    hoy = (ahora or datetime.now()).date()
    if cuando.date() == hoy:
        dia = "hoy"
    elif cuando.date() == hoy - timedelta(days=1):
        dia = "ayer"
    else:
        dia = cuando.strftime("%d/%m")
    detalle = info.get("detalle") or ""
    return f"Última vez: {dia} {cuando:%H:%M}" + (f" · {detalle}" if detalle else "")


# ---------- reenviar las facturas de un día ----------

def interpretar_fecha_pasada(texto, hoy=None):
    """'hoy', 'ayer', 'dd/mm' o 'dd/mm/aa(aa)' -> date de hoy o anterior (un dd/mm aún por llegar es del año pasado).
    Lanza ValueError si no se entiende o si es futura."""
    hoy = hoy or datetime.now().date()
    texto = (texto or "").strip().lower()
    if texto == "hoy":
        return hoy
    if texto == "ayer":
        return hoy - timedelta(days=1)
    partes = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2}|\d{4}))?", texto)
    if not partes:
        raise ValueError("Escribe el día como dd/mm (por ejemplo 06/10) o «ayer».")
    dia, mes, anio = int(partes.group(1)), int(partes.group(2)), partes.group(3)
    try:
        if anio:
            fecha = datetime(int(anio) + (2000 if len(anio) == 2 else 0), mes, dia).date()
        else:
            fecha = datetime(hoy.year, mes, dia).date()
            if fecha > hoy:
                fecha = datetime(hoy.year - 1, mes, dia).date()
    except ValueError:
        raise ValueError(f"El día {texto} no existe.")
    if fecha > hoy:
        raise ValueError("No se pueden enviar facturas de un día que aún no ha llegado.")
    return fecha


def facturas_locales_del_dia(carpeta, fecha):
    """PDF ya descargados de ese día en la carpeta de facturas (Factura_<número>_AAAA-MM-DD.pdf)."""
    try:
        return sorted(Path(carpeta).glob(f"Factura_*_{fecha.isoformat()}.pdf"))
    except OSError:
        return []


# ---------------------------------------------------
# 6. INTERFAZ GRÁFICA
# ---------------------------------------------------

def generar_fichero_stock_fiat_task(ruta_excel, carpeta_salida, log_callback=None, progreso_callback=None, finalizado_callback=None, stop_event=None):
    global stock_fiat_running
    stock_fiat_running = True
    def log(texto, tag=None):
        if log_callback: log_callback(f"[STOCK FIAT] {texto}", tag)
    log("Iniciando generacion de ficheros Stock FIAT...")
    try:
        log("Leyendo archivo Excel...")
        try:
            wb = xlrd.open_workbook(ruta_excel)
            ws = wb.sheet_by_index(0)
        except Exception as e:
            log(f"Error al abrir el Excel: {e}", "error")
            return
        if ws.ncols < 4:
            log("Error: El Excel debe tener al menos columnas hasta la D.", "error")
            return

        col_ref = 1
        col_cant = 3
        total_filas = ws.nrows
        if total_filas <= 0:
            log("El Excel no tiene datos de articulos.", "error")
            return

        def valor_limpio(valor):
            if isinstance(valor, float) and valor.is_integer():
                return str(int(valor))
            return str(valor).strip()

        lineas_fichero = []
        log(f"Procesando {total_filas} filas...")
        for fila_idx in range(ws.nrows):
            if stop_event and stop_event.is_set(): break
            try:
                ref_raw = valor_limpio(ws.cell_value(fila_idx, col_ref))
                if not ref_raw:
                    continue
                ref_clean = ref_raw.replace("2/", "")
                referencia = ref_clean.zfill(13)[:13]
                try:
                    cant_val = int(float(ws.cell_value(fila_idx, col_cant)))
                except Exception:
                    cant_val = 0
                    log(f"Cantidad no válida en fila {fila_idx+1}, se usará 0.", "error")
                cantidad = str(cant_val).zfill(7)[:7]
                ubicacion = "00000000"
                linea = f"{referencia}{cantidad}{ubicacion}"
                lineas_fichero.append(linea)
            except Exception as e:
                log(f"Error en fila {fila_idx+1}: {e}", "error")
            if progreso_callback: progreso_callback((fila_idx + 1) / total_filas)

        if not (stop_event and stop_event.is_set()):
            if not os.path.exists(carpeta_salida):
                os.makedirs(carpeta_salida, exist_ok=True)

            proximo_f = obtener_numero_mopar_del_dia(carpeta_salida, prefijo_clave="stock_fiat")
            letras = string.ascii_uppercase
            archivos_creados = 0

            for i in range(0, len(lineas_fichero), 10):
                grupo = lineas_fichero[i:i+10]
                if len(lineas_fichero) <= 10:
                    nombre_fichero = f"F{proximo_f}.txt"
                else:
                    indice_letra = i // 10
                    letra = letras[indice_letra % len(letras)]
                    nombre_fichero = f"F{proximo_f}{letra}.txt"

                ruta_final = os.path.join(carpeta_salida, nombre_fichero)
                with open(ruta_final, "w", encoding="utf-8") as f:
                    f.write("\n".join(grupo))

                log(f"Generado Stock FIAT: {nombre_fichero}", "success")
                archivos_creados += 1

            log(f"Proceso completado. Se han generado {archivos_creados} archivos con prefijo F{proximo_f}.", "success")
    except Exception as e:
        log(f"Error critico: {e}", "error")
        registrar_error_detallado(str(e), traceback.format_exc())
    finally:
        stock_fiat_running = False
        if finalizado_callback: finalizado_callback()

class AppIntegrada(ctk.CTk):
    def __init__(self):
        super().__init__()
        estilo.configurar_fuentes(tkinter.font.families(self))
        self.config = cargar_configuracion()
        self._ultimo_porcentaje = -1
        self.aplicar_tema_inicial()
        self.title("Palausa Aprovisionamiento")
        self.ajustar_tamano_ventana()
        self.resizable(True, True)
        if os.path.exists(ICONO_PATH):
            try: self.iconbitmap(ICONO_PATH)
            except Exception: pass
        self.avisos = efectos.GestorAvisos(self)
        self.lista_tareas = tareas.ListaTareas(ruta_datos(tareas.ARCHIVO))
        self.arrastre_activo = efectos.preparar_arrastre(self)
        self.setup_ui()
        self.configurar_atajos()
        self.comprobar_hoja_pedidos_al_inicio()

    def ajustar_tamano_ventana(self):
        """Tamaño inicial proporcional a la pantalla (hasta 1500x1000), centrado, y mínimo que mantiene todo legible."""
        try:
            escala = ctk.ScalingTracker.get_window_scaling(self)
        except Exception:
            escala = 1.0
        ancho_pantalla = self.winfo_screenwidth() / escala
        alto_pantalla = self.winfo_screenheight() / escala
        ancho = int(min(1500, ancho_pantalla * 0.85))
        alto = int(min(1000, alto_pantalla * 0.88))
        x = int((ancho_pantalla - ancho) / 2)
        y = max(0, int((alto_pantalla - alto) / 2) - 20)
        self.geometry(f"{ancho}x{alto}+{x}+{y}")
        self.minsize(900, 680)

    def aplicar_tema_inicial(self):
        modo = self.config.get("modo_apariencia", "Light")
        ctk.set_appearance_mode("Dark" if modo == "Dark" else "Light")
        ctk.set_default_color_theme("blue")

    def aplicar_estilos_ui(self):
        # Paleta del tema de color elegido, en su variante clara u oscura; definida en estilo.py
        for nombre, valor in estilo.paleta(self.config.get("modo_apariencia", "Light") == "Dark", self.tema_actual()).items():
            setattr(self, f"color_{nombre}", valor)

    def tema_actual(self):
        tema = self.config.get("tema_color")
        return tema if tema in estilo.TEMAS else estilo.TEMA_POR_DEFECTO

    # ---------- componentes reutilizables (todos usan estilo.py) ----------

    def crear_etiqueta(self, parent, texto, tam=estilo.TAM_TEXTO, negrita=False, secundario=False, mono=False, **kw):
        return ctk.CTkLabel(parent, text=texto, font=estilo.fuente(tam, negrita, mono), text_color=self.color_texto_sec if secundario else self.color_texto, **kw)

    def crear_entrada(self, parent, **kw):
        kw.setdefault("font", estilo.fuente())
        entrada = ctk.CTkEntry(parent, height=estilo.ALTO_CONTROL, corner_radius=estilo.RADIO, border_width=1, border_color=self.color_borde_campo,
                               fg_color=self.color_campo, text_color=self.color_texto, placeholder_text_color=self.color_texto_placeholder, **kw)
        # Foco visible: el campo activo se resalta con el color principal
        entrada.bind("<FocusIn>", lambda evento, campo=entrada: campo.configure(border_color=self.color_principal, border_width=2), add="+")
        entrada.bind("<FocusOut>", lambda evento, campo=entrada: campo.configure(border_color=self.color_borde_campo, border_width=1), add="+")
        return entrada

    def crear_boton_secundario(self, parent, texto, comando, alto=estilo.ALTO_CONTROL, **kw):
        kw.setdefault("font", estilo.fuente())
        return ctk.CTkButton(parent, text=texto, height=alto, corner_radius=estilo.RADIO, fg_color=self.color_panel_medio, text_color=self.color_texto, hover_color=self.color_panel_medio_hover, border_width=1, border_color=self.color_borde, command=comando, **kw)

    def crear_boton_principal(self, parent, text, command, alto=estilo.ALTO_BOTON):
        return ctk.CTkButton(parent, text=text, height=alto, corner_radius=estilo.RADIO, font=estilo.fuente(estilo.TAM_BOTON, True), fg_color=self.color_principal, text_color="#ffffff", hover_color=self.color_hover, command=command)

    def crear_menu(self, parent, valores, comando=None, **kw):
        kw.setdefault("font", estilo.fuente())
        kw.setdefault("dropdown_font", estilo.fuente())
        return ctk.CTkOptionMenu(parent, values=valores, height=estilo.ALTO_CONTROL, corner_radius=estilo.RADIO, fg_color=self.color_principal, button_color=self.color_principal, button_hover_color=self.color_hover, command=comando, **kw)

    def crear_tarjeta_tab(self, tab, titulo, subtitulo, en_rejilla=False, icono=None, ultima_vez=None):
        card = ctk.CTkFrame(tab, fg_color=self.color_superficie, corner_radius=estilo.RADIO_TARJETA, border_width=1, border_color=self.color_borde)
        if en_rejilla:
            self.registrar_tarjeta(tab, card)
        else:
            card.pack(padx=20, pady=estilo.PAD_TARJETA * 2, fill="both", expand=True)
        fila_titulo = ctk.CTkFrame(card, fg_color="transparent")
        fila_titulo.pack(anchor="w", fill="x", padx=estilo.PAD_X, pady=(14, 2))
        if icono and estilo.FUENTE_ICONOS:
            ctk.CTkLabel(fila_titulo, text=estilo.ICONOS[icono], font=estilo.fuente_icono(), text_color=self.color_principal).pack(side="left", padx=(0, 10))
        ctk.CTkLabel(fila_titulo, text=titulo, font=estilo.fuente_titulo(), text_color=self.color_texto).pack(side="left")
        if ultima_vez:  # cuándo se hizo por última vez, a la derecha del título (sin ocupar más alto)
            etiqueta = self.crear_etiqueta(fila_titulo, texto_ultima_vez((self.config.get("ultimas_ejecuciones") or {}).get(ultima_vez)), estilo.TAM_PEQUENO, secundario=True)
            etiqueta.pack(side="right", padx=(12, 0))
            if not hasattr(self, "etiquetas_ultima_vez"):
                self.etiquetas_ultima_vez = {}
            self.etiquetas_ultima_vez[ultima_vez] = etiqueta
        etiqueta_subtitulo = self.crear_etiqueta(card, subtitulo, secundario=True, wraplength=470, justify="left")
        etiqueta_subtitulo.pack(anchor="w", padx=estilo.PAD_X, pady=(0, 8))
        card.bind("<Configure>", lambda evento, etiqueta=etiqueta_subtitulo: self.ajustar_subtitulo(etiqueta, evento.width), add="+")
        efectos.activar_resalte(self, card)
        ctk.CTkFrame(card, fg_color=self.color_borde, height=1).pack(fill="x", padx=estilo.PAD_X, pady=(0, estilo.PAD_Y + 2))
        return card

    def ajustar_subtitulo(self, etiqueta, ancho_tarjeta):
        """El subtítulo ocupa todo el ancho de la tarjeta (menos líneas en tarjetas anchas, sin cortes en las estrechas)."""
        try:
            escala = ctk.ScalingTracker.get_widget_scaling(etiqueta)
        except Exception:
            escala = 1.0
        ancho = max(200, int(ancho_tarjeta / escala) - 2 * estilo.PAD_X - 8)
        if getattr(etiqueta, "_ancho_ajuste", None) != ancho:
            etiqueta._ancho_ajuste = ancho
            etiqueta.configure(wraplength=ancho)

    def resaltar_pestana_activa(self, tabview):
        """La pestaña activa lleva texto blanco sobre el azul (contraste 4.5:1); las demás, el texto normal."""
        try:
            activa = tabview.get()
            for nombre, boton in tabview._segmented_button._buttons_dict.items():
                boton.configure(text_color="#ffffff" if nombre == activa else self.color_texto)
        except Exception:
            pass

    def registrar_tarjeta(self, contenedor, card):
        """Las tarjetas de una pestaña se reparten en una o dos columnas según el ancho disponible."""
        if not hasattr(contenedor, "_tarjetas"):
            contenedor._tarjetas = []
            contenedor._distribucion = None
            contenedor.bind("<Configure>", lambda evento, c=contenedor: self.repartir_tarjetas(c, evento.width), add="+")
        contenedor._tarjetas.append(card)
        self.repartir_tarjetas(contenedor, contenedor.winfo_width())

    def repartir_tarjetas(self, contenedor, ancho):
        tarjetas = contenedor._tarjetas
        columnas = 2 if (ancho >= estilo.ANCHO_DOS_COLUMNAS and len(tarjetas) > 1) else 1
        distribucion = (columnas, len(tarjetas))
        if distribucion == contenedor._distribucion:
            return
        contenedor._distribucion = distribucion
        for columna in range(2):
            contenedor.grid_columnconfigure(columna, weight=1 if columna < columnas else 0, uniform="tarjetas" if columna < columnas else "")
        for posicion, card in enumerate(tarjetas):
            fila, columna = divmod(posicion, columnas)
            ocupa_todo = columnas == 2 and posicion == len(tarjetas) - 1 and len(tarjetas) % 2 == 1
            card.grid(row=fila, column=0 if ocupa_todo else columna, columnspan=2 if ocupa_todo else 1, sticky="nsew", padx=estilo.PAD_TARJETA, pady=estilo.PAD_TARJETA)

    def crear_fila_archivo(self, parent, label_text, entry_attr, config_key, placeholder, button_text, command, extensiones=(".xlsx",)):
        """Fila con campo de ruta y botón Examinar/Carpeta; también admite soltar encima un archivo o carpeta
        (de carpeta si el botón es «Carpeta»; si no, un archivo con alguna de las extensiones indicadas)."""
        self.crear_etiqueta(parent, label_text, negrita=True).pack(anchor="w", padx=estilo.PAD_X, pady=(0, 4))
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        entry = self.crear_entrada(row, placeholder_text=placeholder)
        entry.insert(0, self.config.get(config_key, ""))
        entry.pack(side="left", fill="x", expand=True, padx=(0, 8))
        entry._valor_anterior = entry.get().strip()
        entry.bind("<FocusOut>", lambda evento, campo=entry: self.campo_cambiado(campo), add="+")
        setattr(self, entry_attr, entry)
        self.crear_boton_secundario(row, button_text, command, width=92).pack(side="right")
        if self.arrastre_activo:
            tipo = "carpeta" if button_text == "Carpeta" else "archivo"
            try:
                efectos.activar_soltar(self, entry, lambda datos, campo=entry: self.soltar_en_campo(campo, datos, tipo, extensiones))
            except Exception:
                pass  # sin arrastrar y soltar sigue funcionando «Examinar»
        return entry

    def soltar_en_campo(self, entrada, datos, tipo, extensiones=(".xlsx",)):
        """Rellena el campo con lo soltado si es del tipo que espera; si no, avisa y lo deja como estaba."""
        rutas = self.tk.splitlist(datos) if datos else ()
        if not rutas:
            return False
        ruta = rutas[0]
        if not os.path.exists(ruta) and os.path.exists(datos):
            ruta = datos  # una sola ruta sin formato de lista Tcl (con barras invertidas sin escapar)
        if tipo == "carpeta":
            if not os.path.isdir(ruta):
                self.avisar("Aquí va una carpeta: suelta una carpeta, no un archivo.", "aviso")
                return False
        elif not os.path.isfile(ruta) or Path(ruta).suffix.lower() not in extensiones:
            self.avisar(f"Aquí va un Excel ({' o '.join(extensiones)}).", "aviso")
            return False
        entrada.delete(0, "end")
        entrada.insert(0, ruta.replace("\\", "/"))  # mismo formato que «Examinar»
        self.campo_cambiado(entrada)
        return True

    def avisar(self, texto, tipo="info", duracion=None, accion=None):
        """Aviso que no interrumpe (abajo a la derecha, se va solo). tipo: 'info', 'ok' o 'aviso'.
        accion: (texto del botón, función), p. ej. ("Abrir", ...)."""
        try:
            return self.avisos.mostrar(texto, tipo, duracion, accion)
        except Exception:
            self.agregar_log(texto, "error" if tipo == "aviso" else ("success" if tipo == "ok" else None))
            return None

    def crear_fila_ajuste(self, card, texto):
        """Fila con etiqueta a la izquierda; el control se coloca a la derecha con pack(side='right', padx=estilo.PAD_FILA, pady=6)."""
        fila = ctk.CTkFrame(card, fg_color=self.color_superficie_alt, corner_radius=10)
        fila.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        self.crear_etiqueta(fila, texto, negrita=True).pack(side="left", padx=(estilo.PAD_FILA, 12), pady=estilo.PAD_FILA - 4)
        return fila

    def crear_campo_en_fila(self, card, texto, atributo, clave, defecto="", **kw):
        """Fila con etiqueta arriba y campo de texto a todo el ancho (por ejemplo, listas de correos)."""
        fila = ctk.CTkFrame(card, fg_color=self.color_superficie_alt, corner_radius=10)
        fila.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        self.crear_etiqueta(fila, texto, negrita=True).pack(anchor="w", padx=estilo.PAD_FILA, pady=(estilo.PAD_FILA - 4, 4))
        entrada = self.crear_entrada(fila, **kw)
        entrada.insert(0, self.config.get(clave, defecto))
        entrada.pack(padx=estilo.PAD_FILA, pady=(0, estilo.PAD_FILA - 2), fill="x")
        setattr(self, atributo, entrada)
        return entrada

    def crear_bloque_credenciales(self, card, filas):
        """filas: lista de (etiqueta, atributo, clave, defecto, ocultar)."""
        marco = ctk.CTkFrame(card, fg_color=self.color_superficie_alt, corner_radius=10)
        marco.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        for numero, (texto, atributo, clave, defecto, ocultar) in enumerate(filas):
            self.crear_etiqueta(marco, texto, negrita=True).grid(row=numero, column=0, padx=estilo.PAD_FILA, pady=6, sticky="w")
            entrada = self.crear_entrada(marco, width=220, show="*" if ocultar else "")
            entrada.insert(0, self.config.get(clave, defecto))
            entrada.grid(row=numero, column=1, padx=estilo.PAD_FILA, pady=6, sticky="ew")
            setattr(self, atributo, entrada)
        marco.grid_columnconfigure(1, weight=1)

    def crear_caja_log(self, parent, alto=110):
        caja = ctk.CTkTextbox(parent, height=alto, font=estilo.fuente(estilo.TAM_TEXTO, mono=True), border_width=1, border_color=self.color_borde, corner_radius=estilo.RADIO, fg_color=self.color_campo, text_color=self.color_texto)
        caja.tag_config("success", foreground=self.color_ok)
        caja.tag_config("error", foreground=self.color_texto_error)
        caja.configure(state="disabled")
        return caja

    def crear_registro(self, card, titulo, atributo, alto, pady_final=estilo.PAD_Y):
        self.crear_etiqueta(card, titulo, negrita=True).pack(anchor="w", padx=estilo.PAD_X, pady=(0, 4))
        caja = self.crear_caja_log(card, alto=alto)
        caja.pack(padx=estilo.PAD_X, pady=(0, pady_final), fill="x")
        setattr(self, atributo, caja)
        return caja

    # ---------- ventana ----------

    def setup_ui(self):
        self._ultimo_porcentaje = -1
        self.aplicar_estilos_ui()
        self.ultima_ruta_generada = ""
        self.configure(fg_color=self.color_panel_claro)

        self.construir_cabecera()

        self.construir_panel_actividad()

        self.tabview = ctk.CTkTabview(
            self,
            corner_radius=estilo.RADIO_TARJETA,
            segmented_button_selected_color=self.color_principal,
            segmented_button_selected_hover_color=self.color_hover,
            segmented_button_fg_color=self.color_panel_medio,
            segmented_button_unselected_color=self.color_panel_medio,
            text_color=self.color_texto,
            fg_color=self.color_panel_claro,
            border_width=0,
        )
        self.tabview.pack(pady=(10, 6), padx=16, fill="both", expand=True)
        # Las pestañas se construyen la primera vez que se abren (arranque más rápido)
        self.constructores_pestanas = {
            "Automatización": ("tab_auto", self.build_tab_auto),
            "Pedidos": ("tab_pedidos", self.build_tab_pedidos),
            "Comparador": ("tab_comparador", self.build_tab_comparador),
            "Fichero Mopar": ("tab_mopar", self.build_tab_mopar),
            "Stock FIAT": ("tab_stock_fiat", self.build_tab_stock_fiat),
            "Facturas": ("tab_madd", self.build_tab_madd),
            "Reclamaciones": ("tab_reclamaciones_power_line", self.build_tab_reclamaciones),
            "Configuración": ("tab_config", self.build_tab_config),
        }
        self.pestanas_construidas = set()
        for nombre, (atributo, _) in self.constructores_pestanas.items():
            setattr(self, atributo, self.tabview.add(nombre))
        self.reclamaciones_power_line = ModuloReclamacionesPowerLine(self, self.config)
        self.construir_pestana(self.tabview.get())
        self.tabview.configure(command=self.al_cambiar_pestana)
        self.resaltar_pestana_activa(self.tabview)

    def construir_cabecera(self):
        """Cabecera de bienvenida: logo, saludo según la hora, fecha y un resumen de cómo va el día."""
        cabecera = ctk.CTkFrame(self, fg_color=self.color_superficie, corner_radius=0, border_width=0)
        cabecera.pack(fill="x")
        cabecera_int = ctk.CTkFrame(cabecera, fg_color="transparent")
        cabecera_int.pack(fill="x", padx=20, pady=(12, 10))
        logo = self.cargar_logo(38)
        if logo is not None:
            ctk.CTkLabel(cabecera_int, image=logo, text="").pack(side="left", padx=(0, 14))
        bloque = ctk.CTkFrame(cabecera_int, fg_color="transparent")
        bloque.pack(side="left")
        ctk.CTkLabel(bloque, text=texto_saludo(self.config.get("nombre_usuario", "")), font=estilo.fuente_titulo(estilo.TAM_APP), text_color=self.color_texto).pack(anchor="w")
        self.crear_etiqueta(bloque, f"{fecha_larga()}  ·  Palausa Aprovisionamiento {VERSION}", secundario=True).pack(anchor="w")
        derecha = ctk.CTkFrame(cabecera_int, fg_color="transparent")
        derecha.pack(side="right")
        # Tamaño mínimo explícito: vacío (sin fichas, p. ej. en una instalación nueva) un CTkFrame ocupa 200x200
        self.frame_resumen = ctk.CTkFrame(derecha, fg_color="transparent", width=1, height=1)
        self.frame_resumen.pack(anchor="e")
        self.crear_etiqueta(derecha, "Ctrl+1…8 pestañas  ·  Ctrl+Intro ejecutar  ·  Esc detener", estilo.TAM_PEQUENO, secundario=True).pack(anchor="e", pady=(6, 0))
        self.actualizar_resumen()
        ctk.CTkFrame(self, fg_color=self.color_borde, height=1).pack(fill="x")

    def imagen_logo(self):
        try:
            from PIL import Image
            imagen = Image.open(ICONO_PATH)
            if hasattr(imagen, "ico"):
                imagen = imagen.ico.getimage(max(imagen.ico.sizes()))  # la resolución más grande del .ico
            return imagen
        except Exception:
            return None

    def cargar_logo(self, tam):
        imagen = self.imagen_logo()
        if imagen is None:
            return None
        return ctk.CTkImage(light_image=imagen, dark_image=imagen, size=(tam, tam))

    def mostrar_animacion_inicio(self):
        """La ultraball se abre y deja paso a la app (se desactiva en Configuración; un clic la salta)."""
        if not self.config.get("animacion_inicio", True):
            return None
        imagen = self.imagen_logo()
        if imagen is None:
            return None
        oscuro = self.config.get("modo_apariencia", "Light") == "Dark"
        # Luz dorada, como el amarillo de la ultraball: más clara sobre fondo oscuro, más intensa sobre fondo claro
        color_luz = animacion_inicio.mezclar(animacion_inicio.AMARILLO_BOLA, "#ffffff", 0.5 if oscuro else 0.0)
        return animacion_inicio.reproducir(self, imagen, fondo=self.color_panel_claro, color_texto=self.color_texto_sec,
                                           color_luz=color_luz, fuente=estilo.fuente_titulo(15))

    def actualizar_resumen(self):
        """Fichas con el estado del día (hoja de pedidos, facturas MADD y tareas pendientes)."""
        marco = getattr(self, "frame_resumen", None)
        if marco is None or not marco.winfo_exists():
            return
        for hijo in marco.winfo_children():
            hijo.destroy()
        for icono, texto, correcto in resumen_del_dia(self.config):
            self.crear_ficha_resumen(marco, icono, texto, self.color_ok if correcto else self.color_texto_sec)
        lista = getattr(self, "lista_tareas", None)
        pendientes = lista.pendientes() if lista is not None else 0
        if pendientes:
            vencidas = lista.vencidas(datetime.now().date())
            texto = f"{pendientes} tarea{'s' if pendientes != 1 else ''} pendiente{'s' if pendientes != 1 else ''}"
            if vencidas:
                texto += f" · {vencidas} vencida{'s' if vencidas != 1 else ''}"
            self.ficha_tareas = self.crear_ficha_resumen(marco, "tareas", texto, self.color_texto_error if vencidas else self.color_principal, self.ir_a_tareas)
        texto_boton, comando = self.estado_boton_actualizacion()
        if texto_boton:
            self.btn_actualizacion = ctk.CTkButton(marco, text=texto_boton, command=comando, height=28, corner_radius=12, font=estilo.fuente(estilo.TAM_PEQUENO, True),
                                                   fg_color=self.color_principal, hover_color=self.color_hover, text_color="#ffffff",
                                                   text_color_disabled="#ffffff", state="normal" if comando else "disabled")
            self.btn_actualizacion.pack(side="left", padx=(8, 0))

    def crear_ficha_resumen(self, marco, icono, texto, color_icono, al_pulsar=None):
        ficha = ctk.CTkFrame(marco, fg_color=self.color_superficie_alt, corner_radius=12)
        ficha.pack(side="left", padx=(8, 0))
        cursor = "hand2" if al_pulsar else ""
        if estilo.FUENTE_ICONOS:
            ctk.CTkLabel(ficha, text=estilo.ICONOS[icono], font=estilo.fuente_icono(13), text_color=color_icono, cursor=cursor).pack(side="left", padx=(10, 6), pady=4)
        self.crear_etiqueta(ficha, texto, estilo.TAM_PEQUENO, cursor=cursor).pack(side="left", padx=(0 if estilo.FUENTE_ICONOS else 10, 10), pady=4)
        if al_pulsar:
            for widget in [ficha] + ficha.winfo_children():
                widget.bind("<Button-1>", al_pulsar)
        return ficha

    # ---------- actualizaciones (GitHub Releases; ver actualizaciones.py y cambio_version.py) ----------

    def estado_boton_actualizacion(self):
        """Texto y acción del botón de la cabecera: ('', None) si no hay nada que actualizar."""
        if getattr(self, "descargando_actualizacion", False):
            return "Descargando actualización...", None
        preparada = getattr(self, "version_preparada", None)
        if preparada:
            return f"Reiniciar para actualizar a {preparada}", self.reiniciar_para_actualizar
        info = getattr(self, "actualizacion_disponible", None)
        if info:
            return f"Actualizar a {info['version']}", self.iniciar_actualizacion
        return "", None

    def mostrar_captura(self, marcadas, total, pedido):
        """«¡Capturado!» al terminar bien el flujo completo (se desactiva con el interruptor de animaciones)."""
        if not self.config.get("animacion_inicio", True):
            return None
        imagen = self.imagen_logo()
        if imagen is None:
            return None
        oscuro = self.config.get("modo_apariencia", "Light") == "Dark"
        color_luz = animacion_inicio.mezclar(animacion_inicio.AMARILLO_BOLA, "#ffffff", 0.5 if oscuro else 0.0)
        subtitulo = f"{marcadas} de {total} líneas a poner a 0  ·  {pedido}" if total else pedido
        return animacion_inicio.reproducir_captura(self, imagen, fondo=self.color_panel_claro, color_texto=self.color_texto, color_luz=color_luz,
                                                   titulo="¡Capturado!", subtitulo=subtitulo, fuente=estilo.fuente_titulo(24),
                                                   fuente_subtitulo=estilo.fuente(14))

    def comprobar_actualizaciones_al_inicio(self):
        """Al abrir: avisa si hay una versión publicada más nueva (en segundo plano; sin internet no pasa nada)."""
        carpeta = actualizaciones.carpeta_instalacion()
        if carpeta is None:
            return
        self.version_preparada = cambio_version.version_preparada(carpeta)
        if self.version_preparada:
            self.actualizar_resumen()
            return

        def buscar():
            try:
                info = actualizaciones.buscar_actualizacion()
            except Exception:
                return  # sin conexión o GitHub no responde: se volverá a mirar en el próximo arranque
            if info:
                self.after(0, self.mostrar_aviso_actualizacion, info)

        threading.Thread(target=buscar, daemon=True).start()

    def mostrar_aviso_actualizacion(self, info):
        self.actualizacion_disponible = info
        self.agregar_log(f"[ACTUALIZACIÓN] Hay una versión nueva: {info['version']} (tienes la {VERSION}). Pulsa «Actualizar a {info['version']}» arriba a la derecha.")
        self.actualizar_resumen()

    def iniciar_actualizacion(self):
        info = getattr(self, "actualizacion_disponible", None)
        carpeta = actualizaciones.carpeta_instalacion()
        if not info or carpeta is None or getattr(self, "descargando_actualizacion", False):
            return
        notas = info["notas"] or "Sin notas de la versión."
        if len(notas) > 1200:
            notas = notas[:1200] + "..."
        tamano = f" ({info['tamano'] / 1048576:.0f} MB)" if info.get("tamano") else ""
        if not messagebox.askyesno("Actualización disponible", f"Hay una versión nueva: {info['version']} (tienes la {VERSION}).\n\nNovedades:\n{notas}\n\n"
                                   f"Se descarga ahora{tamano} y se aplica al reiniciar la app. Tu configuración y tus datos no se tocan. ¿Descargar?"):
            return
        self.descargando_actualizacion = True
        self.actualizar_resumen()
        self.actualizar_estado(f"descargando la versión {info['version']}")
        self._ultimo_porcentaje = -1

        def trabajo():
            try:
                actualizaciones.descargar_y_preparar(info, carpeta, progreso=lambda v: self.after(0, self.actualizar_progreso, v))
                self.after(0, self.actualizacion_descargada, info["version"], None)
            except Exception as e:
                registrar_error_detallado(f"Actualización a {info['version']}: {e}", traceback.format_exc())
                self.after(0, self.actualizacion_descargada, info["version"], str(e))

        threading.Thread(target=trabajo, daemon=True).start()

    def actualizacion_descargada(self, version, error):
        self.descargando_actualizacion = False
        if error:
            self.actualizar_estado("la actualización no se pudo descargar")
            self.agregar_log(f"[ACTUALIZACIÓN] No se pudo descargar la versión {version}: {error}", "error")
            messagebox.showwarning("Actualización", f"No se pudo descargar la versión {version}:\n{error}\n\nPuedes volver a intentarlo desde el botón de arriba.")
            self.actualizar_resumen()
            return
        self.version_preparada = version
        self.actualizar_resumen()
        self.actualizar_estado(f"versión {version} lista para instalar")
        self.agregar_log(f"[ACTUALIZACIÓN] Versión {version} descargada. Se instala al reiniciar la app.", "success")
        if messagebox.askyesno("Actualización", f"La versión {version} está lista.\n\n¿Reiniciar la app ahora para instalarla?\n(Si no, se instalará la próxima vez que la abras.)"):
            self.reiniciar_para_actualizar()

    def reiniciar_para_actualizar(self):
        carpeta = actualizaciones.carpeta_instalacion()
        if carpeta is None or self.hay_tarea_activa():
            return
        try:
            self.guardar_configuracion_ui_internal()
        except Exception:
            pass
        lanzador = carpeta / cambio_version.NOMBRE_EJECUTABLE
        try:
            subprocess.Popen([str(lanzador)], cwd=str(carpeta))
        except OSError as e:
            messagebox.showwarning("Actualización", f"No se pudo reiniciar la app: {e}\n\nCiérrala y vuelve a abrirla para instalar la versión nueva.")
            return
        self.destroy()  # al cerrarse, el lanzador cambia la carpeta del programa y abre la versión nueva

    def construir_panel_actividad(self):
        """Un único panel inferior: estado y progreso en una fila, acciones en otra y registro plegable."""
        self.panel_actividad = ctk.CTkFrame(self, fg_color=self.color_superficie, corner_radius=10, border_width=1, border_color=self.color_borde)
        self.panel_actividad.pack(side="bottom", pady=(0, 12), padx=16, fill="x")

        self.frame_progreso = ctk.CTkFrame(self.panel_actividad, fg_color="transparent")
        self.frame_progreso.pack(padx=estilo.PAD_X, pady=(10, 2), fill="x")
        self.estado_label = self.crear_etiqueta(self.frame_progreso, "Estado: listo", negrita=True)
        self.estado_label.configure(text_color=self.color_enfasis)
        self.estado_label.pack(side="left")
        self.porcentaje_label = self.crear_etiqueta(self.frame_progreso, "0%", estilo.TAM_BOTON, negrita=True)
        self.porcentaje_label.configure(text_color=self.color_enfasis)
        self.porcentaje_label.pack(side="right")
        self.progreso_var = ctk.DoubleVar()
        self.barra_prog = ctk.CTkProgressBar(self.frame_progreso, variable=self.progreso_var, height=10, corner_radius=999, progress_color=self.color_principal, fg_color=self.color_superficie_alt)
        self.barra_prog.pack(side="left", fill="x", expand=True, padx=14)

        self.frame_estado = ctk.CTkFrame(self.panel_actividad, fg_color="transparent")
        self.frame_estado.pack(padx=estilo.PAD_X, pady=(2, 6), fill="x")
        self.frame_acciones = ctk.CTkFrame(self.frame_estado, fg_color="transparent")
        self.frame_acciones.pack(side="right")
        for texto, comando in (
            ("Abrir última salida", self.abrir_ultima_salida),
            ("Abrir errores", lambda: self.abrir_directorio(ERROR_FOLDER)),
            ("Abrir log", self.abrir_log_detallado),
            ("Limpiar log", self.limpiar_log),
        ):
            self.crear_boton_secundario(self.frame_acciones, texto, comando, alto=estilo.ALTO_BOTON_COMPACTO).pack(side="left", padx=(6, 0))
        self.btn_registro = self.crear_boton_secundario(self.frame_acciones, "", self.alternar_registro, alto=estilo.ALTO_BOTON_COMPACTO, width=104)
        self.btn_registro.pack(side="left", padx=(6, 0))
        self.ruta_label = self.crear_etiqueta(self.frame_estado, "Última salida: -", estilo.TAM_PEQUENO, secundario=True, mono=True, anchor="w", justify="left")
        self.ruta_label.pack(side="left", fill="x", expand=True)

        self.frame_log = ctk.CTkFrame(self.panel_actividad, fg_color="transparent")
        self.log_box = ctk.CTkTextbox(self.frame_log, height=130, font=estilo.fuente(estilo.TAM_TEXTO, mono=True), border_width=1, border_color=self.color_borde, corner_radius=estilo.RADIO, fg_color=self.color_campo, text_color=self.color_texto)
        self.log_box.pack(padx=0, pady=(0, 0), fill="both", expand=True)
        self.log_box.tag_config("success", foreground=self.color_ok)
        self.log_box.tag_config("error", foreground=self.color_texto_error)
        self.mostrar_registro(self.config.get("registro_visible", True))

    def mostrar_registro(self, visible):
        self.registro_visible = bool(visible)
        if self.registro_visible:
            self.frame_log.pack(padx=estilo.PAD_X, pady=(0, 10), fill="x")
            self.btn_registro.configure(text="Registro ▾")
        else:
            self.frame_log.pack_forget()
            self.btn_registro.configure(text="Registro ▸")

    def alternar_registro(self):
        self.mostrar_registro(not self.registro_visible)
        completa = cargar_configuracion()
        completa["registro_visible"] = self.registro_visible
        guardar_configuracion(completa)
        self.config["registro_visible"] = self.registro_visible

    def build_tab_reclamaciones(self, acc_color):
        self.reclamaciones_power_line.construir(self.tab_reclamaciones_power_line)

    def construir_pestana(self, nombre):
        if nombre in self.pestanas_construidas:
            return
        self.pestanas_construidas.add(nombre)
        self.constructores_pestanas[nombre][1](self.color_principal)

    def al_cambiar_pestana(self):
        self.construir_pestana(self.tabview.get())
        self.actualizar_paneles_por_pestana()

    def actualizar_paneles_por_pestana(self):
        """La pestaña Reclamaciones usa todo el alto de la ventana; en el resto se muestra el panel de actividad."""
        self.resaltar_pestana_activa(self.tabview)
        if self.tabview.get() == "Reclamaciones":
            self.panel_actividad.pack_forget()
            self.tabview.pack_configure(pady=(10, 12))
            return
        if not self.panel_actividad.winfo_manager():
            self.panel_actividad.pack(side="bottom", pady=(0, 12), padx=16, fill="x", before=self.tabview)
        self.tabview.pack_configure(pady=(10, 6))

    # ---------- atajos de teclado ----------

    def configurar_atajos(self):
        for numero in range(1, 9):
            self.bind_all(f"<Control-Key-{numero}>", lambda evento, n=numero - 1: self.ir_a_pestana(n))
        self.bind_all("<Control-Return>", lambda evento: self.accion_principal())
        self.bind_all("<Escape>", lambda evento: self.detener_tarea_actual())

    def ir_a_pestana(self, indice):
        nombres = list(self.constructores_pestanas)
        if 0 <= indice < len(nombres):
            self.tabview.set(nombres[indice])
            self.al_cambiar_pestana()

    def accion_principal(self):
        """Ejecuta la acción principal de la pestaña visible (nunca detiene una tarea en curso)."""
        pestana = self.tabview.get()
        acciones = {
            "Automatización": (automation_running, self.iniciar_auto),
            "Pedidos": (pedidos_running, self.iniciar_pedidos),
            "Comparador": (comparador_running, self.iniciar_comparador),
            "Fichero Mopar": (mopar_running, self.iniciar_mopar),
            "Stock FIAT": (stock_fiat_running, self.iniciar_stock_fiat),
            "Facturas": (madd_running, self.iniciar_facturas_madd),
            "Reclamaciones": (False, self.reclamaciones_power_line.generar_documento),
            "Configuración": (False, self.guardar_config_ui),
        }
        en_curso, accion = acciones[pestana]
        if not en_curso:
            accion()

    def detener_tarea_actual(self):
        tareas = (
            (automation_running, "la automatización", self.toggle_auto),
            (pedidos_running, "la generación de pedidos", self.toggle_pedidos),
            (comparador_running, "el comparador", self.toggle_comparador),
            (mopar_running, "el fichero Mopar", self.toggle_mopar),
            (stock_fiat_running, "Stock FIAT", self.toggle_stock_fiat),
        )
        for en_curso, descripcion, detener in tareas:
            if en_curso:
                if messagebox.askyesno("Detener", f"¿Detener {descripcion}? Lo ya procesado se conserva."):
                    detener()
                return

    # ---------- pestañas ----------

    def build_tab_auto(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_auto, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Consulta de disponibilidad", "Carga tu Excel y ejecuta la consulta paralela contra el portal.", en_rejilla=True, ultima_vez="auto", icono="buscar")
        self.crear_fila_archivo(card, "Excel de stock", "excel_auto_entry", "ruta_excel_auto", "Excel de stock...", "Examinar", lambda: self.seleccionar_archivo(self.excel_auto_entry, [("Excel files", "*.xlsx")]))
        self.excel_auto_entry._al_cambiar = self.al_cambiar_excel_stock
        fila = self.crear_fila_ajuste(card, "Precio mínimo (EUR)")
        self.precio_entry = self.crear_entrada(fila, width=120, justify="center")
        self.precio_entry.insert(0, self.config.get("precio_min", "0"))
        self.precio_entry.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        self.btn_auto = self.crear_boton_principal(card, "Iniciar automatización", self.toggle_auto)
        self.btn_auto.pack(padx=estilo.PAD_X, pady=(2, 6), fill="x")
        self.btn_flujo = self.crear_boton_secundario(card, "Flujo completo: automatización + filtro de rotación", self.iniciar_flujo_completo, alto=estilo.ALTO_BOTON, font=estilo.fuente(estilo.TAM_BOTON, True))
        self.btn_flujo.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        self.crear_registro(card, "Referencias consultadas", "auto_log_box", 140)
        botones_auto = ctk.CTkFrame(card, fg_color="transparent")
        botones_auto.pack(padx=estilo.PAD_X, pady=(0, 14), fill="x")
        botones_auto.grid_columnconfigure((0, 1), weight=1, uniform="botones_auto")
        self.crear_boton_secundario(botones_auto, "Abrir carpeta del Excel", self.abrir_carpeta_auto).grid(row=0, column=0, sticky="ew", padx=(0, 4))
        self.btn_pedidos_stock = self.crear_boton_secundario(botones_auto, "Abrir pedidos stock", self.abrir_pedidos_stock)
        self.btn_pedidos_stock.grid(row=0, column=1, sticky="ew", padx=(4, 0))
        efectos.poner_pista(self, self.btn_pedidos_stock, self.pista_pedidos_stock)

        card_rotacion = self.crear_tarjeta_tab(contenedor, "Filtro de rotación", "Marca referencias de baja rotación según el umbral y la propuesta de pedido. Las filas con precio inferior al precio mínimo de Automatización se dejan sin marcar, y las piezas caras salen en azul para que decidas tú.", en_rejilla=True, ultima_vez="rotacion", icono="filtro")
        self.crear_fila_archivo(card_rotacion, "Excel de Quiter", "excel_quiter_entry", "ruta_excel_quiter", "Excel de Quiter (DMS)...", "Examinar", lambda: self.seleccionar_archivo(self.excel_quiter_entry, [("Excel files", "*.xlsx")]))
        self.crear_fila_archivo(card_rotacion, "Archivo de propuesta", "excel_rot_entry", "ruta_excel_rotacion", "Archivo PROPUESTAPEDIDO.xls...", "Examinar", lambda: self.seleccionar_archivo(self.excel_rot_entry, [("Excel files", "*.xls *.xlsx")]), extensiones=(".xls", ".xlsx"))
        self.crear_fila_archivo(card_rotacion, "Carpeta destino", "folder_rot_entry", "ruta_folder_rotacion", "Carpeta destino...", "Carpeta", lambda: self.seleccionar_carpeta(self.folder_rot_entry))
        fila = self.crear_fila_ajuste(card_rotacion, "Umbral mínimo")
        self.umbral_entry = self.crear_entrada(fila, width=120, justify="center")
        self.umbral_entry.insert(0, self.config.get("umbral_rotacion", "0.08"))
        self.umbral_entry.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        fila.pack_configure(pady=(0, 4))
        # Ayuda breve pegada a la fila: qué hace el umbral (comparación en filtrar_por_rotacion_task)
        self.ayuda_umbral = self.crear_etiqueta(card_rotacion, "Se marcan en rojo las referencias cuya media en la propuesta es menor que este valor "
                                                "(si no aparecen en la propuesta, cuentan como 0). Subirlo endurece el filtro: más filas en rojo. "
                                                "Bajarlo lo suaviza.", estilo.TAM_PEQUENO, secundario=True, wraplength=470, justify="left")
        self.ayuda_umbral.pack(anchor="w", padx=estilo.PAD_X + 4, pady=(0, estilo.PAD_Y))
        fila = self.crear_fila_ajuste(card_rotacion, "Revisar si la unidad supera (EUR)")
        self.limite_revisar_entry = self.crear_entrada(fila, width=120, justify="center")
        self.limite_revisar_entry.insert(0, self.config.get("limite_revisar_rotacion", "1600"))
        self.limite_revisar_entry.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        efectos.poner_pista(self, self.limite_revisar_entry, "Las piezas cuyo precio unitario (P. compra) supere este importe salen en azul para que decidas tú; 0 = no marcar")
        fila.pack_configure(pady=(0, 2))
        # Vista previa: cuántas filas marcaría el filtro con el umbral escrito (se recalcula sola)
        self.vista_previa_umbral = self.crear_etiqueta(card_rotacion, "", estilo.TAM_PEQUENO, negrita=True, wraplength=470, justify="left")
        self.vista_previa_umbral.pack(anchor="w", padx=estilo.PAD_X + 4, pady=(0, estilo.PAD_Y))
        for etiqueta in (self.ayuda_umbral, self.vista_previa_umbral):
            card_rotacion.bind("<Configure>", lambda evento, etiqueta=etiqueta: self.ajustar_subtitulo(etiqueta, evento.width), add="+")
        self.btn_rotacion = self.crear_boton_principal(card_rotacion, "Resaltar por rotación", self.iniciar_filtro_rotacion)
        self.btn_rotacion.pack(padx=estilo.PAD_X, pady=(2, estilo.PAD_Y), fill="x")
        self.crear_registro(card_rotacion, "Registro de rotación", "rotacion_log_box", 110)
        botones = ctk.CTkFrame(card_rotacion, fg_color="transparent")
        botones.pack(padx=estilo.PAD_X, pady=(0, 14), fill="x")
        botones.grid_columnconfigure((0, 1), weight=1, uniform="botones_rotacion")
        self.crear_boton_secundario(botones, "Abrir carpeta destino", self.abrir_carpeta_rotacion).grid(row=0, column=0, sticky="ew", padx=(0, 4))
        self.btn_lista_quiter = self.crear_boton_secundario(botones, "Copiar lista para Quiter", self.copiar_lista_quiter)
        self.btn_lista_quiter.grid(row=0, column=1, sticky="ew", padx=(4, 0))
        self.iniciar_vista_previa_umbral()

    def build_tab_pedidos(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_pedidos, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Generación de pedidos", "Convierte el Excel urgente en lotes CSV y guarda copia de seguridad antes de limpiar.", en_rejilla=True, ultima_vez="pedidos", icono="carrito")
        self.crear_fila_archivo(card, "Excel urgente", "excel_pedidos_entry", "ruta_excel_pedidos", "Excel URGENTE.xlsx...", "Examinar", lambda: self.seleccionar_archivo(self.excel_pedidos_entry, [("Excel files", "*.xlsx")]))
        self.crear_fila_archivo(card, "Carpeta destino", "folder_pedidos_entry", "ruta_folder_pedidos", "Carpeta de destino...", "Carpeta", lambda: self.seleccionar_carpeta(self.folder_pedidos_entry))
        self.btn_pedidos = self.crear_boton_principal(card, "Generar pedidos", self.toggle_pedidos)
        self.btn_pedidos.pack(padx=estilo.PAD_X, pady=(2, estilo.PAD_Y), fill="x")
        self.crear_registro(card, "Números de pedido generados", "pedidos_generados_box", 90, pady_final=14)

        card_hoja = self.crear_tarjeta_tab(contenedor, "Hoja de pedidos del día", "Crea una copia en blanco de la plantilla, guardada con la fecha de hoy dentro de HOJAS DE PEDIDO.", en_rejilla=True, ultima_vez="hoja", icono="calendario")
        self.crear_fila_archivo(card_hoja, "Plantilla de hoja", "plantilla_hoja_entry", "ruta_plantilla_hojas", "HOJA DE PEDIDOS.xlsx...", "Examinar", lambda: self.seleccionar_archivo(self.plantilla_hoja_entry, [("Excel files", "*.xlsx")]))
        self.crear_fila_archivo(card_hoja, "Carpeta HOJAS DE PEDIDO", "carpeta_hojas_entry", "ruta_carpeta_hojas", "Carpeta HOJAS DE PEDIDO...", "Carpeta", lambda: self.seleccionar_carpeta(self.carpeta_hojas_entry))
        fila_hoja = ctk.CTkFrame(card_hoja, fg_color="transparent")
        fila_hoja.pack(padx=estilo.PAD_X, pady=(2, 14), fill="x")
        self.btn_hoja_pedidos = self.crear_boton_principal(fila_hoja, "Generar hoja de hoy", self.iniciar_hoja_pedidos)
        self.btn_hoja_pedidos.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.btn_abrir_hoja = self.crear_boton_secundario(fila_hoja, "Abrir hoja de hoy", self.abrir_hoja_de_hoy, alto=estilo.ALTO_BOTON, width=150)
        self.btn_abrir_hoja.pack(side="right")
        efectos.poner_pista(self, self.btn_abrir_hoja, lambda: f"Abrir {self.ruta_hoja_de_hoy().name}" if self.ruta_hoja_de_hoy() else "")

        card_avisos = self.crear_tarjeta_tab(contenedor, "Avisos de no suministro y sustituciones", "Lee la hoja de pedido ya rellenada y gestiona en Outlook un aviso por cliente con las referencias en NO SUMINISTRA y otro con las referencias SUSTITUIDAS.", en_rejilla=True, ultima_vez="avisos", icono="correo")
        self.crear_fila_archivo(card_avisos, "Hoja de pedido a procesar", "hoja_avisos_entry", "ruta_hoja_avisos", "Hoja dd-mm-aaaa.xlsx ya rellenada...", "Examinar", lambda: self.seleccionar_archivo(self.hoja_avisos_entry, [("Excel files", "*.xlsx")]))
        boton_hoy = self.crear_boton_secundario(self.hoja_avisos_entry.master, "Hoy", self.usar_hoja_de_hoy_en_avisos, width=56)
        boton_hoy.pack(side="right", padx=(0, 8))
        efectos.poner_pista(self, boton_hoy, lambda: f"Usar la hoja de hoy ({self.ruta_hoja_de_hoy().name})" if self.ruta_hoja_de_hoy() else "")
        self.crear_fila_archivo(card_avisos, "Excel de clientes", "clientes_avisos_entry", "ruta_excel_clientes", "clientes.xlsx...", "Examinar", lambda: self.seleccionar_archivo(self.clientes_avisos_entry, [("Excel files", "*.xlsx")]))
        self.crear_campo_en_fila(card_avisos, "Correo(s) de compañeros (se les envía la hoja de pedido completa)", "email_companeros_avisos_entry", "email_companeros_avisos", "", placeholder_text="correo1@palausa.es; correo2@palausa.es")
        self.avisos_auto_var = ctk.BooleanVar(value=self.config.get("avisos_envio_automatico", True))
        ctk.CTkCheckBox(card_avisos, text="Enviar automáticamente si el cliente tiene email (si no, se guarda como borrador)", variable=self.avisos_auto_var, fg_color=self.color_principal, hover_color=self.color_hover, text_color=self.color_texto, font=estilo.fuente()).pack(anchor="w", padx=estilo.PAD_X, pady=(0, estilo.PAD_Y))
        self.btn_avisos = self.crear_boton_principal(card_avisos, "Generar avisos de no suministro", self.iniciar_avisos_no_suministro)
        self.btn_avisos.pack(padx=estilo.PAD_X, pady=(2, estilo.PAD_Y), fill="x")
        self.crear_registro(card_avisos, "Registro de avisos", "avisos_log_box", 110, pady_final=14)
        self.after(400, self.sugerir_hoja_de_hoy_en_avisos)

    def build_tab_comparador(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_comparador, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Comparador de Excel", "Resalta las líneas repetidas y genera un archivo comparado en la carpeta elegida.", en_rejilla=True, ultima_vez="comparador", icono="comparar")
        self.crear_fila_archivo(card, "Excel antiguo", "excel_old_entry", "ruta_excel_antiguo", "Excel antiguo...", "Examinar", lambda: self.seleccionar_archivo(self.excel_old_entry, [("Excel files", "*.xlsx")]))
        self.crear_fila_archivo(card, "Excel nuevo", "excel_new_entry", "ruta_excel_nuevo", "Excel nuevo...", "Examinar", lambda: self.seleccionar_archivo(self.excel_new_entry, [("Excel files", "*.xlsx")]))
        self.crear_fila_archivo(card, "Carpeta destino", "folder_comp_entry", "ruta_folder_comparador", "Carpeta destino...", "Carpeta", lambda: self.seleccionar_carpeta(self.folder_comp_entry))
        fila = self.crear_fila_ajuste(card, "Color de resaltado")
        self.color_optionmenu = self.crear_menu(fila, ["Amarillo", "Rojo", "Verde", "Azul", "Naranja"])
        self.color_optionmenu.set(self.config.get("color_resaltado", "Amarillo"))
        self.color_optionmenu.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        fila_comparar = ctk.CTkFrame(card, fg_color="transparent")
        fila_comparar.pack(padx=estilo.PAD_X, pady=(2, 14), fill="x")
        self.btn_comparador = self.crear_boton_principal(fila_comparar, "Comparar Excels", self.toggle_comparador)
        self.btn_comparador.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.btn_informe_anterior = self.crear_boton_secundario(fila_comparar, "Abrir informe anterior", self.abrir_informe_anterior, alto=estilo.ALTO_BOTON, width=170)
        self.btn_informe_anterior.pack(side="right")
        efectos.poner_pista(self, self.btn_informe_anterior, self.pista_informe_anterior)
        self.construir_tarjeta_tareas(contenedor)

    def build_tab_mopar(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_mopar, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Generador Mopar", "Transforma el Excel en ficheros TXT Mopar con numeración diaria persistente.", en_rejilla=True, ultima_vez="mopar", icono="documento")
        self.crear_fila_archivo(card, "Excel de entrada", "excel_mopar_entry", "ruta_excel_mopar", "Excel de entrada (.xls)...", "Examinar", lambda: self.seleccionar_archivo(self.excel_mopar_entry, [("Excel files (.xls)", "*.xls")]), extensiones=(".xls",))
        self.crear_fila_archivo(card, "Carpeta destino", "folder_mopar_entry", "ruta_folder_mopar", "Carpeta de destino...", "Carpeta", lambda: self.seleccionar_carpeta(self.folder_mopar_entry))
        self.btn_mopar = self.crear_boton_principal(card, "Generar fichero Mopar", self.toggle_mopar)
        self.btn_mopar.pack(padx=estilo.PAD_X, pady=(2, estilo.PAD_Y), fill="x")
        self.crear_registro(card, "Ficheros Mopar generados", "mopar_generados_box", 90, pady_final=14)

    def build_tab_stock_fiat(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_stock_fiat, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Stock FIAT", "Transforma el Excel en ficheros TXT con el mismo formato final de Mopar.", en_rejilla=True, ultima_vez="stock_fiat", icono="coche")
        self.crear_fila_archivo(card, "Excel de entrada", "excel_stock_fiat_entry", "ruta_excel_stock_fiat", "Excel de entrada (.xls)...", "Examinar", lambda: self.seleccionar_archivo(self.excel_stock_fiat_entry, [("Excel files (.xls)", "*.xls")]), extensiones=(".xls",))
        self.crear_fila_archivo(card, "Carpeta destino", "folder_stock_fiat_entry", "ruta_folder_stock_fiat", "Carpeta de destino...", "Carpeta", lambda: self.seleccionar_carpeta(self.folder_stock_fiat_entry))
        self.btn_stock_fiat = self.crear_boton_principal(card, "Generar Stock FIAT", self.toggle_stock_fiat)
        self.btn_stock_fiat.pack(padx=estilo.PAD_X, pady=(2, 14), fill="x")

    def build_tab_madd(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_madd, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Facturas Stellantis MADD", "Descarga diaria de facturas desde el portal MADD (www.madd.stellantis.com) y envía un correo con los PDF adjuntos.", en_rejilla=True, ultima_vez="madd", icono="factura")
        self.crear_bloque_credenciales(card, [
            ("Usuario MADD", "usuario_madd_entry", "usuario_madd", "", False),
            ("Contraseña MADD", "password_madd_entry", "password_madd", "", True),
        ])
        self.crear_fila_archivo(card, "Carpeta de facturas", "carpeta_facturas_madd_entry", "ruta_carpeta_facturas_madd", "Carpeta donde guardar las facturas...", "Carpeta", lambda: self.seleccionar_carpeta(self.carpeta_facturas_madd_entry))
        self.crear_campo_en_fila(card, "Correo diario de facturas (varios separados por ;)", "email_facturas_madd_entry", "email_facturas_madd", "")
        self.crear_campo_en_fila(card, "Correo de aviso si falla", "email_aviso_fallo_madd_entry", "email_aviso_fallo_madd", "")
        fila = self.crear_fila_ajuste(card, "Facturas de")
        self.filtro_dia_madd_option = self.crear_menu(fila, ["Ayer", "Hoy"])
        self.filtro_dia_madd_option.set(self.config.get("filtro_dia_facturas_madd", "Ayer"))
        self.filtro_dia_madd_option.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        self.btn_madd = self.crear_boton_principal(card, "Ejecutar ahora (prueba manual)", self.iniciar_facturas_madd)
        self.btn_madd.pack(padx=estilo.PAD_X, pady=(2, estilo.PAD_Y), fill="x")
        fila = self.crear_fila_ajuste(card, "Reenviar las facturas de un día")
        self.btn_reenviar_madd = self.crear_boton_secundario(fila, "Reenviar", self.reenviar_facturas_dia, width=92)
        self.btn_reenviar_madd.pack(side="right", padx=(0, estilo.PAD_FILA), pady=6)
        self.fecha_reenvio_madd_entry = self.crear_entrada(fila, width=130, placeholder_text="dd/mm o «ayer»", justify="center")
        self.fecha_reenvio_madd_entry.pack(side="right", padx=(0, 8), pady=6)
        self.fecha_reenvio_madd_entry.bind("<Return>", lambda evento: (self.reenviar_facturas_dia(), "break")[1])
        efectos.poner_pista(self, self.btn_reenviar_madd, "Usa los PDF de ese día que ya estén en la carpeta; si no hay, los descarga de MADD")
        self.crear_registro(card, "Registro de la ejecución", "madd_log_box", 140, pady_final=14)

        card_prog = self.crear_tarjeta_tab(contenedor, "Programación diaria", "Cuando compruebes que la ejecución manual funciona bien, activa la tarea de Windows para que corra sola todos los días laborables a las 07:00 (despierta el PC si está en suspensión).", en_rejilla=True, icono="reloj")
        fila_prog = ctk.CTkFrame(card_prog, fg_color="transparent")
        fila_prog.pack(padx=estilo.PAD_X, pady=(0, 14), fill="x")
        self.btn_activar_tarea_madd = self.crear_boton_principal(fila_prog, "Activar tarea programada (07:00 L-V)", self.activar_tarea_programada_madd)
        self.btn_activar_tarea_madd.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.btn_desactivar_tarea_madd = self.crear_boton_secundario(fila_prog, "Desactivar", self.desactivar_tarea_programada_madd, alto=estilo.ALTO_BOTON)
        self.btn_desactivar_tarea_madd.pack(side="right")

    def build_tab_config(self, acc_color):
        contenedor = ctk.CTkScrollableFrame(self.tab_config, fg_color="transparent")
        contenedor.pack(fill="both", expand=True)
        card = self.crear_tarjeta_tab(contenedor, "Configuración general", "Ajusta credenciales, rendimiento y apariencia visual de la herramienta.", en_rejilla=True, icono="ajustes")
        self.crear_bloque_credenciales(card, [
            ("Usuario", "user_entry", "usuario", "", False),
            ("Contraseña", "pass_entry", "password", "", True),
        ])
        fila = self.crear_fila_ajuste(card, "Hilos paralelos")
        self.hilos_option = self.crear_menu(fila, ["1", "2", "3", "4", "5"])
        self.hilos_option.set(str(self.config.get("num_hilos", "3")))
        self.hilos_option.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        fila = self.crear_fila_ajuste(card, "Apariencia")
        self.modo_option = self.crear_menu(fila, ["Claro", "Oscuro"], comando=self.cambiar_modo_apariencia)
        self.modo_option.set("Oscuro" if self.config.get("modo_apariencia", "Light") == "Dark" else "Claro")
        self.modo_option.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        fila = self.crear_fila_ajuste(card, "Color")
        self.tema_option = self.crear_menu(fila, list(estilo.TEMAS), comando=self.cambiar_modo_apariencia)
        self.tema_option.set(self.tema_actual())
        self.tema_option.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        fila = self.crear_fila_ajuste(card, "Tu nombre (para el saludo)")
        self.nombre_entry = self.crear_entrada(fila, width=220, placeholder_text="Opcional")
        self.nombre_entry.insert(0, self.config.get("nombre_usuario", ""))
        self.nombre_entry.pack(side="right", padx=estilo.PAD_FILA, pady=6)
        self.crear_fila_archivo(card, "Excel de pedidos stock (botón «Abrir pedidos stock» de Automatización)", "excel_pedidos_stock_entry", "ruta_excel_pedidos_stock",
                                "PEDIDOS STOCK.xlsx...", "Examinar", lambda: self.seleccionar_archivo(self.excel_pedidos_stock_entry, [("Excel files", "*.xlsx *.xlsm *.xls")]),
                                extensiones=(".xlsx", ".xlsm", ".xls"))
        self.crear_fila_archivo(card, "Carpeta de informes diarios (botón «Abrir informe anterior» del Comparador)", "carpeta_informes_entry", "ruta_carpeta_informes",
                                "Carpeta de informes diarios...", "Carpeta", lambda: self.seleccionar_carpeta(self.carpeta_informes_entry))
        fila = self.crear_fila_ajuste(card, "Animaciones (apertura y «¡Capturado!»)")
        self.animacion_var = ctk.BooleanVar(value=self.config.get("animacion_inicio", True))
        ctk.CTkSwitch(fila, text="", variable=self.animacion_var, width=46, progress_color=self.color_principal).pack(side="right", padx=estilo.PAD_FILA, pady=6)
        self.reclamaciones_power_line.construir_ajustes(card)
        self.crear_boton_principal(card, "Guardar configuración", self.guardar_config_ui).pack(padx=estilo.PAD_X, pady=(2, 14), fill="x")

    # ---------- tareas pendientes (pestaña Comparador; datos y fechas en tareas.py) ----------

    def construir_tarjeta_tareas(self, contenedor):
        card = self.crear_tarjeta_tab(contenedor, "Tareas pendientes", "Apunta lo que tengas pendiente; se guarda solo. Doble clic sobre una tarea o su fecha para cambiarla.",
                                      en_rejilla=True, icono="tareas")
        fila = ctk.CTkFrame(card, fg_color="transparent")
        fila.pack(padx=estilo.PAD_X, pady=(0, estilo.PAD_Y), fill="x")
        self.nueva_tarea_entry = self.crear_entrada(fila, placeholder_text="Nueva tarea...")
        self.nueva_tarea_entry.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.nueva_fecha_entry = self.crear_entrada(fila, width=118, placeholder_text="Fecha (dd/mm)", justify="center")
        self.nueva_fecha_entry.pack(side="left", padx=(0, 8))
        for campo in (self.nueva_tarea_entry, self.nueva_fecha_entry):
            campo.bind("<Return>", lambda evento: self.anadir_tarea())
        self._nueva_importante = False
        self.btn_nueva_importante = self.crear_boton_estrella(fila, False, self.alternar_nueva_importante, estilo.ALTO_CONTROL)
        self.btn_nueva_importante.pack(side="left", padx=(0, 8))
        boton = self.crear_boton_principal(fila, "Añadir", self.anadir_tarea, alto=estilo.ALTO_CONTROL)
        boton.configure(width=92)
        boton.pack(side="left")
        self.lista_tareas_frame = ctk.CTkFrame(card, fg_color="transparent")
        self.lista_tareas_frame.pack(padx=estilo.PAD_X, fill="x")
        self.pie_tareas = ctk.CTkFrame(card, fg_color="transparent", height=1)
        self.pie_tareas.pack(padx=estilo.PAD_X, pady=(2, 14), fill="x")
        self.pintar_tareas()

    def crear_boton_estrella(self, padre, importante, comando, lado=28):
        """Estrella para marcar una tarea como importante (dorada si lo es)."""
        if estilo.FUENTE_ICONOS:
            texto, fuente = estilo.ICONOS["estrella_llena" if importante else "estrella"], estilo.fuente_icono(15)
        else:
            texto, fuente = "★" if importante else "☆", estilo.fuente(15)
        color = "#d99a00" if importante else self.color_texto_sec
        return ctk.CTkButton(padre, text=texto, font=fuente, width=lado, height=lado, corner_radius=estilo.RADIO, command=comando,
                             fg_color="transparent", hover_color=self.color_panel_medio_hover, text_color=color, border_width=0)

    def alternar_nueva_importante(self):
        self._nueva_importante = not self._nueva_importante
        estrella = self.crear_boton_estrella(self.btn_nueva_importante.master, self._nueva_importante, self.alternar_nueva_importante, estilo.ALTO_CONTROL)
        estrella.pack(side="left", padx=(0, 8), after=self.btn_nueva_importante)
        self.btn_nueva_importante.destroy()
        self.btn_nueva_importante = estrella

    def anadir_tarea(self):
        texto = self.nueva_tarea_entry.get().strip()
        if not texto:
            self.avisar("Escribe la tarea antes de añadirla.")
            return "break"
        try:
            fecha = tareas.interpretar_fecha(self.nueva_fecha_entry.get(), datetime.now().date())
        except ValueError as e:
            self.avisar(str(e), "aviso")
            return "break"
        if not self.guardar_tareas(lambda: self.lista_tareas.anadir(texto, fecha, self._nueva_importante)):
            return "break"
        self.nueva_tarea_entry.delete(0, "end")
        self.nueva_fecha_entry.delete(0, "end")
        if self._nueva_importante:
            self.alternar_nueva_importante()
        self.pintar_tareas()
        self.nueva_tarea_entry.focus_set()
        return "break"

    def guardar_tareas(self, cambio):
        """Aplica un cambio a la lista; si no se puede guardar el archivo, avisa en vez de fallar."""
        try:
            cambio()
            return True
        except OSError as e:
            registrar_error_detallado(f"Tareas pendientes: {e}", traceback.format_exc())
            self.avisar(f"No se pudieron guardar las tareas: {e}", "aviso")
            return False

    def pintar_tareas(self):
        marco = getattr(self, "lista_tareas_frame", None)
        if marco is None or not marco.winfo_exists():
            self.actualizar_resumen()
            return
        for hijo in marco.winfo_children():
            hijo.destroy()
        hoy = datetime.now().date()
        lista = self.lista_tareas.ordenadas()
        if not lista:
            self.crear_etiqueta(marco, "No tienes tareas pendientes. Escribe una arriba y pulsa Intro.", secundario=True).pack(anchor="w", pady=(2, 6))
        for tarea in lista:
            self.crear_fila_tarea(marco, tarea, hoy)
        for hijo in self.pie_tareas.winfo_children():
            hijo.destroy()
        hechas = sum(1 for t in lista if t["hecha"])
        if lista:
            pendientes = len(lista) - hechas
            self.crear_etiqueta(self.pie_tareas, f"{pendientes} pendiente{'s' if pendientes != 1 else ''}" + (f" · {hechas} hecha{'s' if hechas != 1 else ''}" if hechas else ""),
                                estilo.TAM_PEQUENO, secundario=True).pack(side="left", pady=(4, 0))
        if hechas:
            self.crear_boton_secundario(self.pie_tareas, f"Borrar completadas ({hechas})", self.borrar_tareas_completadas, alto=estilo.ALTO_BOTON_COMPACTO).pack(side="right", pady=(4, 0))
        self.actualizar_resumen()

    def crear_fila_tarea(self, marco, tarea, hoy):
        identificador = tarea["id"]
        fila = ctk.CTkFrame(marco, fg_color=self.color_superficie_alt, corner_radius=10)
        fila.pack(fill="x", pady=(0, 6))
        hecha = ctk.BooleanVar(value=tarea["hecha"])
        ctk.CTkCheckBox(fila, text="", variable=hecha, width=24, checkbox_width=20, checkbox_height=20, corner_radius=6, border_width=2,
                        fg_color=self.color_principal, hover_color=self.color_hover, border_color=self.color_borde_campo,
                        command=lambda: self.marcar_tarea(identificador, hecha.get())).pack(side="left", padx=(12, 2), pady=8)
        self.crear_boton_estrella(fila, tarea["importante"], lambda: self.alternar_importante_tarea(identificador)).pack(side="left", padx=(0, 4))
        borrar = ctk.CTkButton(fila, text=estilo.ICONOS["cerrar"] if estilo.FUENTE_ICONOS else "✕", font=estilo.fuente_icono(11) if estilo.FUENTE_ICONOS else estilo.fuente(),
                               width=28, height=28, corner_radius=estilo.RADIO, fg_color="transparent", hover_color=self.color_panel_medio_hover,
                               text_color=self.color_texto_sec, command=lambda: self.borrar_tarea(identificador))
        borrar.pack(side="right", padx=(0, 8))
        if not tarea["hecha"]:
            estado = tareas.estado_fecha(tarea["fecha"], hoy)
            color = {"vencida": self.color_texto_error, "hoy": self.color_principal}.get(estado, self.color_texto_sec if tarea["fecha"] else self.color_texto_placeholder)
            etiqueta_fecha = ctk.CTkLabel(fila, text=tareas.texto_fecha(tarea["fecha"], hoy) or "sin fecha", text_color=color, cursor="hand2",
                                          font=estilo.fuente(estilo.TAM_PEQUENO, estado in ("vencida", "hoy")))
            etiqueta_fecha.pack(side="right", padx=(6, 8))
            etiqueta_fecha.bind("<Double-Button-1>", lambda evento: self.editar_tarea_en_linea(etiqueta_fecha, tarea, "fecha"))
        if tarea["hecha"]:
            fuente = ctk.CTkFont(family=estilo.FUENTE, size=estilo.TAM_TEXTO, overstrike=True)
        else:
            fuente = estilo.fuente(estilo.TAM_TEXTO, tarea["importante"])
        etiqueta = ctk.CTkLabel(fila, text=tarea["texto"], font=fuente, anchor="w", justify="left", wraplength=420,
                                text_color=self.color_texto_sec if tarea["hecha"] else self.color_texto)
        etiqueta.pack(side="left", fill="x", expand=True, padx=(2, 6), pady=6)
        etiqueta.bind("<Double-Button-1>", lambda evento: self.editar_tarea_en_linea(etiqueta, tarea, "texto"))

    def editar_tarea_en_linea(self, etiqueta, tarea, campo):
        """Doble clic: un campo encima del texto (o de la fecha). Intro o salir guarda; Esc cancela."""
        if campo == "texto":
            valor = tarea["texto"]
        else:
            valor = datetime.fromisoformat(tarea["fecha"]).strftime("%d/%m/%Y") if tarea["fecha"] else ""
        entrada = self.crear_entrada(etiqueta.master, width=118 if campo == "fecha" else 200, justify="center" if campo == "fecha" else "left",
                                     placeholder_text="dd/mm" if campo == "fecha" else "")
        if campo == "fecha":
            entrada.place(in_=etiqueta, relx=1.0, rely=0.5, anchor="e")
        else:
            entrada.place(in_=etiqueta, relx=0, rely=0.5, anchor="w", relwidth=1.0)
        if valor:
            entrada.insert(0, valor)
        entrada.focus_set()
        entrada._entry.select_range(0, "end")
        terminado = {"si": False}

        def guardar(evento=None):
            if terminado["si"]:
                return "break"
            terminado["si"] = True
            nuevo = entrada.get()
            if campo == "texto":
                if nuevo.strip() and nuevo.strip() != tarea["texto"]:
                    self.guardar_tareas(lambda: self.lista_tareas.editar(tarea["id"], nuevo))
            else:
                try:
                    fecha = tareas.interpretar_fecha(nuevo, datetime.now().date())
                    self.guardar_tareas(lambda: self.lista_tareas.cambiar_fecha(tarea["id"], fecha))
                except ValueError as e:
                    self.avisar(str(e), "aviso")
            self.after_idle(self.pintar_tareas)
            return "break"

        def cancelar(evento=None):
            terminado["si"] = True
            entrada.destroy()
            return "break"  # que Esc no llegue al atajo global de detener la tarea en curso

        entrada.bind("<Return>", guardar)
        entrada.bind("<FocusOut>", guardar)
        entrada.bind("<Escape>", cancelar)
        entrada.guardar_edicion, entrada.cancelar_edicion = guardar, cancelar
        return entrada

    def marcar_tarea(self, identificador, hecha):
        self.guardar_tareas(lambda: self.lista_tareas.marcar(identificador, hecha))
        self.after_idle(self.pintar_tareas)

    def alternar_importante_tarea(self, identificador):
        self.guardar_tareas(lambda: self.lista_tareas.cambiar_importante(identificador))
        self.after_idle(self.pintar_tareas)

    def borrar_tarea(self, identificador):
        self.guardar_tareas(lambda: self.lista_tareas.borrar(identificador))
        self.after_idle(self.pintar_tareas)

    def borrar_tareas_completadas(self):
        borradas = []
        if self.guardar_tareas(lambda: borradas.append(self.lista_tareas.borrar_completadas())):
            self.avisar(f"Borradas {borradas[0]} tarea{'s' if borradas[0] != 1 else ''} completada{'s' if borradas[0] != 1 else ''}.", "ok")
        self.pintar_tareas()

    def ir_a_tareas(self, evento=None):
        self.tabview.set("Comparador")
        self.al_cambiar_pestana()
        campo = getattr(self, "nueva_tarea_entry", None)
        if campo is not None and campo.winfo_exists():
            campo.focus_set()

    # ---------- abrir los Excel de un clic ----------

    def abrir_archivo(self, ruta, sin_ruta="Elige primero el archivo."):
        """Abre un archivo con su programa (Excel); si falta, avisa sin ventanas."""
        ruta = str(ruta or "").strip()
        if not ruta:
            self.avisar(sin_ruta)
            return False
        if not os.path.isfile(ruta):
            self.avisar(f"No se encuentra {Path(ruta).name}: puede que se haya movido o borrado.", "aviso")
            return False
        if not abrir_ruta_en_sistema(ruta):
            self.avisar(f"No se pudo abrir {Path(ruta).name}.", "aviso")
            return False
        return True

    def ruta_hoja_de_hoy(self):
        carpeta = self.valor_widget_o_config("carpeta_hojas_entry", "ruta_carpeta_hojas", "")
        return ruta_hoja_pedidos_del_dia(carpeta) if carpeta else None

    def abrir_hoja_de_hoy(self):
        ruta = self.ruta_hoja_de_hoy()
        if ruta is None:
            self.avisar("Elige primero la carpeta de hojas de pedido.")
            return False
        if not ruta.exists():
            self.avisar("Todavía no hay hoja de hoy: pulsa «Generar hoja de hoy».")
            return False
        return self.abrir_archivo(ruta)

    def carpeta_informes(self, preguntar=False):
        """Carpeta de los informes diarios (Configuración). Si no está puesta y se pide, se elige una vez y se guarda."""
        carpeta = str(self.valor_widget_o_config("carpeta_informes_entry", "ruta_carpeta_informes", "") or "").strip()
        if carpeta or not preguntar:
            return carpeta
        carpeta = filedialog.askdirectory(title="Elige la carpeta de los informes diarios")
        if carpeta:
            self.guardar_ruta_elegida("ruta_carpeta_informes", "carpeta_informes_entry", carpeta, "Carpeta de informes diarios guardada")
        return carpeta

    def guardar_ruta_elegida(self, clave, atributo, ruta, mensaje):
        """Guarda una ruta elegida desde un botón y la refleja en su campo de Configuración si está abierto."""
        completa = cargar_configuracion()
        completa[clave] = ruta
        guardar_configuracion(completa)
        self.config[clave] = ruta
        campo = getattr(self, atributo, None)
        if campo is not None and campo.winfo_exists():
            campo.delete(0, "end")
            campo.insert(0, ruta)
        self.avisar(f"{mensaje}. Puedes cambiarla en Configuración.", "ok")

    def excel_pedidos_stock(self, preguntar=False):
        """Excel de pedidos stock (Configuración). Si no está puesto y se pide, se elige una vez y se guarda."""
        ruta = str(self.valor_widget_o_config("excel_pedidos_stock_entry", "ruta_excel_pedidos_stock", "") or "").strip()
        if ruta or not preguntar:
            return ruta
        ruta = filedialog.askopenfilename(title="Elige el Excel de pedidos stock", filetypes=[("Excel files", "*.xlsx *.xlsm *.xls")])
        if ruta:
            self.guardar_ruta_elegida("ruta_excel_pedidos_stock", "excel_pedidos_stock_entry", ruta, "Excel de pedidos stock guardado")
        return ruta

    def abrir_pedidos_stock(self):
        ruta = self.excel_pedidos_stock(preguntar=True)
        return self.abrir_archivo(ruta) if ruta else False

    def pista_pedidos_stock(self):
        ruta = self.excel_pedidos_stock()
        return f"Abrir {Path(ruta).name}" if ruta else "Elige el Excel de pedidos stock y ábrelo"

    def abrir_informe_anterior(self):
        carpeta = self.carpeta_informes(preguntar=True)
        if not carpeta:
            return False
        if not os.path.isdir(carpeta):
            self.avisar(f"No se encuentra la carpeta de informes diarios ({Path(carpeta).name}). Cámbiala en Configuración.", "aviso")
            return False
        informe = informe_diario_anterior(carpeta)
        if informe is None:
            self.avisar(f"No hay ningún informe anterior a hoy en {Path(carpeta).name}.")
            return False
        return self.abrir_archivo(informe)

    def pista_informe_anterior(self):
        carpeta = self.carpeta_informes()
        if not carpeta:
            return "Elige la carpeta de informes diarios y abre el anterior a hoy"
        informe = informe_diario_anterior(carpeta) if os.path.isdir(carpeta) else None
        return f"Abrir {informe.name}" if informe else "No hay ningún informe anterior a hoy"

    # ---------- «Última vez» en el título de cada tarjeta ----------

    def registrar_ejecucion(self, clave, detalle=""):
        info = registrar_ultima_ejecucion(clave, detalle)
        if info is not None:
            self.config.setdefault("ultimas_ejecuciones", {})[clave] = info
        self.refrescar_ultima_vez(clave, desde_disco=False)

    def refrescar_ultima_vez(self, clave, desde_disco=True):
        if desde_disco:  # la tarea programada de facturas escribe en el archivo desde otro proceso
            self.config["ultimas_ejecuciones"] = cargar_configuracion().get("ultimas_ejecuciones") or {}
        etiqueta = getattr(self, "etiquetas_ultima_vez", {}).get(clave)
        if etiqueta is not None and etiqueta.winfo_exists():
            etiqueta.configure(text=texto_ultima_vez((self.config.get("ultimas_ejecuciones") or {}).get(clave)))

    # ---------- rellenar solos los campos del pedido de la semana ----------

    def campo_cambiado(self, entrada):
        """Avisa al campo de que su valor ha cambiado (Examinar, soltar un archivo o escribir y salir)."""
        actual = entrada.get().strip()
        anterior = getattr(entrada, "_valor_anterior", "")
        if actual == anterior:
            return
        entrada._valor_anterior = actual
        al_cambiar = getattr(entrada, "_al_cambiar", None)
        if al_cambiar:
            al_cambiar(actual, anterior)

    def poner_en_campo(self, atributo, valor):
        campo = getattr(self, atributo, None)
        if campo is None or not campo.winfo_exists():
            return False
        campo.delete(0, "end")
        campo.insert(0, valor)
        campo._valor_anterior = valor
        return True

    def al_cambiar_excel_stock(self, ruta, anterior):
        """Al elegir el Excel de stock de un pedido (p. ej. S40P.xlsx) se ponen también el Excel de Quiter
        (si seguía al de stock) y la propuesta S40PR.xls de la carpeta de rotación, si ya está exportada."""
        if not ruta:
            return
        pedido = Path(ruta).stem
        cambios = []
        quiter = getattr(self, "excel_quiter_entry", None)
        if quiter is not None and quiter.winfo_exists() and quiter.get().strip() in ("", anterior) and quiter.get().strip() != ruta:
            self.poner_en_campo("excel_quiter_entry", ruta)
            cambios.append("el Excel de Quiter")
        carpeta_rotacion = str(self.valor_widget_o_config("folder_rot_entry", "ruta_folder_rotacion", "") or "").strip()
        carpetas = [Path(c) for c in (carpeta_rotacion, str(Path(ruta).parent / "ROTACION")) if c]
        propuesta = next((carpeta / f"{pedido}R{extension}" for carpeta in carpetas for extension in (".xls", ".xlsx")
                          if (carpeta / f"{pedido}R{extension}").is_file()), None)
        if propuesta is not None:
            valor = propuesta.as_posix()
            if self.widget_get("excel_rot_entry").strip() != valor and self.poner_en_campo("excel_rot_entry", valor):
                cambios.append(f"su propuesta {propuesta.name}")
        falta = f"Aún no está la propuesta {pedido}R.xls en {Path(carpeta_rotacion).name}: elígela cuando la exportes." if propuesta is None and carpeta_rotacion else ""
        if cambios:
            self.avisar(f"Pedido {pedido}: he puesto también " + " y ".join(cambios) + "." + (f" {falta}" if falta else ""), "info" if falta else "ok")
        elif falta:
            self.avisar(falta)

    def sugerir_hoja_de_hoy_en_avisos(self):
        """Si la hoja a procesar en Avisos no es la de hoy y la de hoy existe, lo propone con un botón (sin cambiarla sola)."""
        campo = getattr(self, "hoja_avisos_entry", None)
        hoja = self.ruta_hoja_de_hoy()
        if campo is None or not campo.winfo_exists() or hoja is None or not hoja.exists():
            return
        actual = campo.get().strip()
        if actual and Path(actual).resolve() == hoja.resolve():
            return
        texto = f"En Avisos de no suministro tienes la hoja {Path(actual).name}." if actual else "En Avisos de no suministro no hay hoja elegida."
        self.avisar(texto + " ¿Usar la de hoy?", accion=("Usar", self.usar_hoja_de_hoy_en_avisos), duracion=9000)

    def usar_hoja_de_hoy_en_avisos(self):
        hoja = self.ruta_hoja_de_hoy()
        if hoja is None or not hoja.exists():
            self.avisar("Todavía no hay hoja de hoy: pulsa «Generar hoja de hoy».")
            return False
        self.poner_en_campo("hoja_avisos_entry", hoja.as_posix())
        return True

    # ---------- recordatorio de tareas al abrir ----------

    def recordar_tareas(self):
        hoy = datetime.now().date()
        pendientes = [t for t in self.lista_tareas.tareas if not t["hecha"]]
        para_hoy = sum(1 for t in pendientes if tareas.estado_fecha(t["fecha"], hoy) == "hoy")
        vencidas = sum(1 for t in pendientes if tareas.estado_fecha(t["fecha"], hoy) == "vencida")
        if not para_hoy and not vencidas:
            return None
        partes = []
        if para_hoy:
            partes.append(f"{para_hoy} tarea{'s' if para_hoy != 1 else ''} para hoy")
        if vencidas:
            partes.append(f"{vencidas} vencida{'s' if vencidas != 1 else ''}")
        return self.avisar("Tienes " + " y ".join(partes) + ".", "aviso" if vencidas else "info", duracion=9000, accion=("Ver", self.ir_a_tareas))

    # ---------- reenviar las facturas de un día ----------

    def reenviar_facturas_dia(self):
        global madd_running
        if self.hay_tarea_activa():
            return False
        self.guardar_configuracion_ui_internal()
        try:
            fecha = interpretar_fecha_pasada(self.widget_get("fecha_reenvio_madd_entry"))
        except ValueError as e:
            self.avisar(str(e), "aviso")
            return False
        carpeta = self.widget_get("carpeta_facturas_madd_entry")
        email_destino = self.widget_get("email_facturas_madd_entry")
        if not self.validar_carpeta(carpeta, "Selecciona la carpeta donde se guardan las facturas."):
            return False
        if not email_destino or correos_invalidos(email_destino):
            messagebox.showwarning("Aviso", "Revisa el correo de destino de las facturas (varios separados por ;).")
            return False
        rutas = facturas_locales_del_dia(carpeta, fecha)
        dia = fecha.strftime("%d/%m/%Y")
        if rutas:
            if not messagebox.askyesno("Reenviar facturas", f"Se enviará a {email_destino} el correo «Facturas Stellantis del {fecha.isoformat()}» "
                                       f"con las {len(rutas)} facturas del {dia} que ya tienes en la carpeta.\n\n¿Enviar?"):
                return False

            def trabajo():
                global madd_running
                log = lambda t, tag=None: self.after(0, self.log_callback_madd, f"[FACTURAS MADD] {t}", tag)
                error = None
                try:
                    log(f"Reenviando las {len(rutas)} facturas del {dia}...")
                    _madd_enviar_correo_resultado(email_destino, fecha, [str(r) for r in rutas], log)
                except Exception as e:
                    error = str(e)
                    log(f"❌ No se pudo reenviar: {e}", "error")
                    registrar_error_detallado(str(e), traceback.format_exc(), ref="reenviar_facturas_dia")
                finally:
                    madd_running = False
                    self.after(0, self.finalizar_reenvio_facturas, fecha, len(rutas), error)

            hilo = threading.Thread(target=trabajo, daemon=True)
        else:
            usuario, password = self.widget_get("usuario_madd_entry"), self.widget_get("password_madd_entry")
            if not usuario or not password:
                messagebox.showwarning("Aviso", f"No hay facturas del {dia} en la carpeta. Para descargarlas de MADD completa el usuario y la contraseña.")
                return False
            if not messagebox.askyesno("Reenviar facturas", f"No hay facturas del {dia} en la carpeta.\n\n¿Descargarlas de MADD y enviarlas a {email_destino}? "
                                       "(Si ese día no hubo facturas, se enviará el aviso de «no hay facturas».)"):
                return False
            hilo = threading.Thread(target=tarea_facturas_madd, args=(usuario, password, carpeta, email_destino, self.widget_get("email_aviso_fallo_madd_entry")),
                                    kwargs={"fecha_objetivo": fecha,
                                            "log_callback": lambda t, tag=None: self.after(0, self.log_callback_madd, t, tag),
                                            "finalizado_callback": lambda: self.after(0, self.finalizar_reenvio_facturas, fecha, None, None)},
                                    daemon=True)
        self.actualizar_estado(f"reenviando las facturas del {dia}")
        self.btn_reenviar_madd.configure(state="disabled")
        efectos.marcar_ocupado(self.btn_reenviar_madd, "Enviando")
        self.limpiar_caja(self.madd_log_box)
        madd_running = True
        hilo.start()
        return True

    def finalizar_reenvio_facturas(self, fecha, cuantas, error):
        dia = fecha.strftime("%d/%m")
        boton = getattr(self, "btn_reenviar_madd", None)
        if boton is not None and boton.winfo_exists():
            boton.configure(text="Reenviar", state="normal")
        self.actualizar_estado(f"reenvío de las facturas del {dia} terminado")
        if error:
            self.avisar(f"No se pudieron reenviar las facturas del {dia}: revisa el registro.", "aviso")
        else:
            self.avisar(f"Reenvío de las facturas del {dia} terminado: revisa el registro.", "ok")
            if cuantas is not None:  # con descarga, la propia tarea apunta la última vez
                self.registrar_ejecucion("madd", f"reenvío del {dia} ({cuantas} facturas)")
        self.refrescar_ultima_vez("madd")
        self.actualizar_resumen()

    def seleccionar_archivo(self, entry, tipos):
        ruta = filedialog.askopenfilename(filetypes=tipos)
        if ruta:
            entry.delete(0, "end")
            entry.insert(0, ruta)
            self.campo_cambiado(entry)

    def seleccionar_carpeta(self, entry):
        ruta = filedialog.askdirectory()
        if ruta:
            entry.delete(0, "end")
            entry.insert(0, ruta)

    def recortar_log(self, caja):
        """Mantiene el registro en un tamaño razonable: al pasar del máximo borra las líneas más antiguas."""
        lineas = int(caja.index("end-1c").split(".")[0])
        if lineas > LOG_MAX_LINEAS:
            caja.delete("1.0", f"{lineas - LOG_MAX_LINEAS + LOG_LINEAS_A_RECORTAR}.0")

    def agregar_log(self, texto, tag=None):
        self.log_box.insert("end", f"{texto}\n", tag)
        self.recortar_log(self.log_box)
        self.log_box.see("end")

    def actualizar_progreso(self, valor):
        porcentaje = int(valor * 100)
        if porcentaje == self._ultimo_porcentaje and valor < 1:
            return
        self._ultimo_porcentaje = porcentaje
        self.progreso_var.set(valor)
        self.porcentaje_label.configure(text=f"{porcentaje}%")

    def actualizar_estado(self, texto, ruta=""):
        self.estado_label.configure(text=f"Estado: {texto}")
        if ruta:
            self.ultima_ruta_generada = ruta
            self.ruta_label.configure(text=f"Última salida: {ruta}")

    def limpiar_log(self):
        self.log_box.delete("1.0", "end")

    def agregar_a_caja(self, caja, texto, tag=None):
        caja.configure(state="normal")
        caja.insert("end", f"{texto}\n", tag)
        self.recortar_log(caja)
        caja.configure(state="disabled")
        caja.see("end")

    def limpiar_caja(self, caja):
        caja.configure(state="normal")
        caja.delete("1.0", "end")
        caja.configure(state="disabled")

    def log_callback_auto(self, texto, tag=None):
        self.agregar_log(texto, tag)
        self.agregar_a_caja(self.auto_log_box, texto, tag)

    def log_callback_rotacion(self, texto, tag=None):
        self.agregar_log(texto, tag)
        self.agregar_a_caja(self.rotacion_log_box, texto, tag)

    def abrir_carpeta_auto(self):
        ruta = self.widget_get("excel_auto_entry")
        if not ruta:
            self.avisar("Selecciona primero un Excel de stock.")
            return
        self.abrir_directorio(os.path.dirname(ruta) if os.path.isfile(ruta) else ruta)

    def abrir_carpeta_rotacion(self):
        ruta = self.widget_get("folder_rot_entry")
        if not ruta:
            self.avisar("Selecciona primero una carpeta de destino.")
            return
        self.abrir_directorio(ruta)

    def agregar_pedido_generado(self, numero):
        self.pedidos_generados_box.configure(state="normal")
        self.pedidos_generados_box.insert("end", f"{numero}\n")
        self.pedidos_generados_box.configure(state="disabled")
        self.pedidos_generados_box.see("end")

    def limpiar_pedidos_generados(self):
        self.pedidos_generados_box.configure(state="normal")
        self.pedidos_generados_box.delete("1.0", "end")
        self.pedidos_generados_box.configure(state="disabled")

    def agregar_mopar_generado(self, nombre_fichero):
        self.mopar_generados_box.configure(state="normal")
        self.mopar_generados_box.insert("end", f"{nombre_fichero}\n")
        self.mopar_generados_box.configure(state="disabled")
        self.mopar_generados_box.see("end")

    def limpiar_mopar_generados(self):
        self.mopar_generados_box.configure(state="normal")
        self.mopar_generados_box.delete("1.0", "end")
        self.mopar_generados_box.configure(state="disabled")

    def abrir_log_detallado(self):
        if not os.path.exists(LOG_DETALLADO):
            Path(LOG_DETALLADO).touch()
        self.abrir_directorio(LOG_DETALLADO)

    def abrir_directorio(self, ruta):
        if not abrir_ruta_en_sistema(ruta):
            messagebox.showwarning("Aviso", "No se pudo abrir la ruta seleccionada.")

    def abrir_ultima_salida(self):
        if not self.ultima_ruta_generada:
            self.avisar("Todavía no hay una salida generada.")
            return
        if os.path.isfile(self.ultima_ruta_generada):
            ruta_a_abrir = os.path.dirname(self.ultima_ruta_generada)
        else:
            ruta_a_abrir = self.ultima_ruta_generada
        self.abrir_directorio(ruta_a_abrir)

    def widget_get(self, attr_name, default=""):
        widget = getattr(self, attr_name, None)
        if not widget:
            return default
        try:
            return widget.get()
        except Exception:
            return default

    def hay_tarea_activa(self):
        activa = tarea_activa()
        if activa:
            messagebox.showwarning("Aviso", f"Ya hay una tarea en curso ({activa}). Espera a que termine antes de iniciar otra.")
            return True
        return False

    def excel_abierto(self, ruta, descripcion):
        if archivo_bloqueado(ruta):
            messagebox.showwarning("Archivo en uso", f"{descripcion} está abierto en Excel u otro programa:\n{ruta}\n\nCiérralo y vuelve a intentarlo.")
            return True
        return False

    def validar_archivo(self, ruta, mensaje):
        if not ruta or not os.path.exists(ruta):
            messagebox.showwarning("Aviso", mensaje)
            return False
        return True

    def validar_carpeta(self, ruta, mensaje):
        if not ruta:
            messagebox.showwarning("Aviso", mensaje)
            return False
        Path(ruta).mkdir(parents=True, exist_ok=True)
        return True

    def guardar_config_ui(self, silent=False):
        self.guardar_configuracion_ui_internal()
        if not silent:
            self.avisar("Configuración guardada.", "ok")

    def cambiar_modo_apariencia(self, valor=None):
        activa = tarea_activa()
        if activa:
            messagebox.showwarning("Aviso", f"No se puede cambiar la apariencia mientras hay una tarea en curso ({activa}). Espera a que termine.")
            self.modo_option.set("Oscuro" if self.config.get("modo_apariencia", "Light") == "Dark" else "Claro")
            if getattr(self, "tema_option", None) is not None:
                self.tema_option.set(self.tema_actual())
            return
        self.guardar_configuracion_ui_internal()
        self.reconstruir_interfaz()

    def reconstruir_interfaz(self):
        pestana_actual = None
        if hasattr(self, "tabview"):
            try:
                pestana_actual = self.tabview.get()
            except Exception:
                pestana_actual = None
        for widget in self.winfo_children():
            widget.destroy()
        # Se eliminan las referencias a widgets/variables destruidos: las pestañas se reconstruyen bajo demanda
        for nombre, valor in list(vars(self).items()):
            if isinstance(valor, (tkinter.Misc, tkinter.Variable)) and valor is not self:
                delattr(self, nombre)
        self.aplicar_tema_inicial()
        self.setup_ui()
        if pestana_actual:
            try:
                self.tabview.set(pestana_actual)
                self.construir_pestana(pestana_actual)
                self.actualizar_paneles_por_pestana()
            except Exception:
                pass
        self.update_idletasks()  # pinta la interfaz nueva de una vez, sin quedarse a medias

    def toggle_auto(self):
        global automation_running
        if automation_running:
            stop_event_auto.set()
            self.btn_auto.configure(state="disabled")
            efectos.marcar_ocupado(self.btn_auto, "Deteniendo")
        else:
            self.iniciar_auto()

    def iniciar_auto(self, al_terminar=None):
        global automation_running
        if self.hay_tarea_activa():
            return False
        try:
            self.guardar_configuracion_ui_internal()
            ruta = self.widget_get("excel_auto_entry")
            # Usuario, contraseña e hilos viven en la pestaña Configuración, que puede no estar construida todavía
            user = self.valor_widget_o_config("user_entry", "usuario", "")
            pwd = self.valor_widget_o_config("pass_entry", "password", "")
            hilos = int(self.valor_widget_o_config("hilos_option", "num_hilos", "3"))
            precio = float(self.widget_get("precio_entry", "0"))
        except Exception as e:
            registrar_error_detallado(str(e), traceback.format_exc())
            messagebox.showerror("Error", f"No se pudo iniciar la automatización: {e}\n\nComprueba que el precio mínimo (pestaña Automatización) sea un número, por ejemplo 150, y que los hilos paralelos (pestaña Configuración) sean un valor válido.")
            return False
        if precio < 0:
            messagebox.showwarning("Aviso", "El precio mínimo debe ser un número igual o mayor que 0 (por ejemplo 150).")
            return False
        if not self.validar_archivo(ruta, "Selecciona un Excel valido para la automatizacion."):
            return False
        if self.excel_abierto(ruta, "El Excel de la automatización"):
            return False
        if not user or not pwd:
            messagebox.showwarning("Aviso", "Faltan el usuario o la contraseña del portal Citroën. Introdúcelos en la pestaña Configuración y guarda.")
            return False
        self.actualizar_estado("automatización en curso", ruta)
        self.btn_auto.configure(fg_color=self.color_peligro, hover_color=self.color_peligro_hover)
        efectos.marcar_ocupado(self.btn_auto, "Detener")
        self.limpiar_caja(self.auto_log_box)
        stop_event_auto.clear()
        automation_running = True

        def correr():
            global automation_running
            try:
                ejecutar_automatizacion_task(ruta, user, pwd, precio, hilos, lambda t, tag=None: self.after(0, self.log_callback_auto, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_auto), stop_event_auto)
            finally:
                automation_running = False
            if al_terminar:
                self.after(0, al_terminar, "detenida" if stop_event_auto.is_set() else auto_resultado)

        threading.Thread(target=correr, daemon=True).start()
        return True

    def iniciar_flujo_completo(self):
        if self.hay_tarea_activa():
            return
        ruta_auto = self.widget_get("excel_auto_entry")
        ruta_rot = self.widget_get("excel_rot_entry")
        folder_salida = self.widget_get("folder_rot_entry")
        if not self.validar_archivo(ruta_auto, "Selecciona el Excel de stock (se usará para la automatización y para la rotación)."):
            return
        if not self.validar_archivo(ruta_rot, "Selecciona el archivo de propuesta (filtro de rotación)."):
            return
        if not self.validar_carpeta(folder_salida, "Selecciona la carpeta destino del filtro de rotación."):
            return
        try:
            float(self.widget_get("umbral_entry"))
        except Exception:
            messagebox.showwarning("Aviso", "El umbral de rotación debe ser un número, por ejemplo 0.12 (usa punto como separador decimal).")
            return
        if self.excel_abierto(ruta_auto, "El Excel de stock"):
            return
        if self.excel_abierto(str(Path(folder_salida) / f"{Path(ruta_auto).stem}_ROTACION.xlsx"), "El Excel de salida de rotación"):
            return
        if not messagebox.askyesno("Flujo completo", "Se ejecutará la automatización sobre el Excel de stock y, si termina bien, el filtro de rotación sobre ese mismo Excel. ¿Continuar?"):
            return
        self.excel_quiter_entry.delete(0, "end")
        self.excel_quiter_entry.insert(0, ruta_auto)
        self.btn_flujo.configure(state="disabled")
        if not self.iniciar_auto(al_terminar=self.continuar_flujo_tras_auto):
            self.btn_flujo.configure(state="normal")

    def continuar_flujo_tras_auto(self, estado):
        if estado == "incompleta":
            if not messagebox.askyesno("Flujo completo", "La automatización terminó con incidencias: no se han consultado todas las referencias (revisa el registro de la automatización).\n\n¿Quieres lanzar igualmente el filtro de rotación?"):
                self.log_callback_rotacion("Flujo completo detenido por decisión del usuario tras una automatización con incidencias.", "error")
                self.btn_flujo.configure(state="normal")
                return
        elif estado != "completa":
            motivo = {"detenida": "se detuvo", "credenciales": "no pudo entrar al portal (usuario o contraseña incorrectos)"}.get(estado, "falló")
            self.log_callback_rotacion(f"Flujo completo detenido: la automatización {motivo}, no se lanza la rotación.", "error")
            messagebox.showwarning("Flujo completo", f"La automatización {motivo}, así que no se ha lanzado el filtro de rotación. Revisa el registro de la automatización.")
            self.btn_flujo.configure(state="normal")
            return
        self.log_callback_rotacion("Automatización terminada, lanzando el filtro de rotación...", "success")
        self._rotacion_del_flujo = True
        if not self.iniciar_filtro_rotacion():
            self._rotacion_del_flujo = False
            self.btn_flujo.configure(state="normal")

    def finalizar_auto(self):
        acc_color = self.color_principal
        self.actualizar_estado("automatización finalizada", self.excel_auto_entry.get())
        self.btn_auto.configure(text="Iniciar automatización", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        ruta_excel = self.excel_auto_entry.get()
        abrir = ("Abrir", lambda: self.abrir_archivo(ruta_excel))
        if stop_event_auto.is_set() or auto_resultado == "detenida":
            self.avisar("Automatización detenida.")
            self.registrar_ejecucion("auto", "detenida")
        elif auto_resultado == "credenciales":
            self.avisar("El portal ha rechazado el usuario o la contraseña: revísalos en Configuración.", "aviso")
            self.registrar_ejecucion("auto", "usuario o contraseña incorrectos")
        elif auto_resultado == "completa":
            self.avisar("Automatización terminada: se han consultado todas las referencias.", "ok", accion=abrir)
            self.registrar_ejecucion("auto", f"{Path(ruta_excel).stem} completa")
        else:
            self.avisar("La automatización terminó con incidencias: revisa el registro.", "aviso", accion=abrir)
            self.registrar_ejecucion("auto", f"{Path(ruta_excel).stem} con incidencias")

    def toggle_pedidos(self):
        global pedidos_running
        if not pedidos_running:
            self.iniciar_pedidos()
        else:
            stop_event_pedidos.set()

    def iniciar_pedidos(self):
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        ruta = self.excel_pedidos_entry.get()
        folder = self.folder_pedidos_entry.get()
        if not self.validar_archivo(ruta, "Selecciona un Excel válido para generar pedidos."):
            return
        if not self.validar_carpeta(folder, "Selecciona una carpeta de destino para los pedidos."):
            return
        if self.excel_abierto(ruta, "El Excel de pedidos"):
            return
        self.actualizar_estado("generando pedidos", folder)
        self.btn_pedidos.configure(fg_color=self.color_peligro, hover_color=self.color_peligro_hover)
        efectos.marcar_ocupado(self.btn_pedidos, "Detener")
        self.limpiar_pedidos_generados()
        stop_event_pedidos.clear()
        threading.Thread(target=generar_pedidos_task, args=(ruta, folder, lambda t, tag=None: self.after(0, self.agregar_log, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_pedidos), stop_event_pedidos, lambda n: self.after(0, self.agregar_pedido_generado, n)), daemon=True).start()

    def finalizar_pedidos(self):
        acc_color = self.color_principal
        self.actualizar_estado("pedidos generados", self.folder_pedidos_entry.get())
        self.btn_pedidos.configure(text="Generar pedidos", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        carpeta_pedidos = self.folder_pedidos_entry.get()
        self.avisar("Generación de pedidos terminada: revisa el registro.", accion=("Carpeta", lambda: self.abrir_directorio(carpeta_pedidos)))
        self.registrar_ejecucion("pedidos")

    def iniciar_hoja_pedidos(self, forzar=False, automatico=False):
        global hoja_pedidos_running
        if hoja_pedidos_running:
            return
        self.guardar_configuracion_ui_internal()
        plantilla = self.widget_get("plantilla_hoja_entry")
        carpeta = self.widget_get("carpeta_hojas_entry")
        if not self.validar_archivo(plantilla, "Selecciona la plantilla HOJA DE PEDIDOS.xlsx."):
            return
        if not self.validar_carpeta(carpeta, "Selecciona la carpeta HOJAS DE PEDIDO."):
            return
        if not forzar and ruta_hoja_pedidos_del_dia(carpeta).exists():
            if automatico:
                return
            if not messagebox.askyesno("La hoja de hoy ya existe", "Ya existe una hoja de pedidos generada hoy. ¿Quieres sustituirla por una copia en blanco de la plantilla? Se perderá lo que hayas escrito en ella."):
                return
            forzar = True
        self.actualizar_estado("generando hoja de pedidos del día")
        self.btn_hoja_pedidos.configure(state="disabled")
        efectos.marcar_ocupado(self.btn_hoja_pedidos, "Generando")
        threading.Thread(
            target=generar_hoja_pedidos_task,
            args=(plantilla, carpeta),
            kwargs={
                "log_callback": lambda t, tag=None: self.after(0, self.agregar_log, t, tag),
                "progreso_callback": lambda v: self.after(0, self.actualizar_progreso, v),
                "finalizado_callback": lambda: self.after(0, self.finalizar_hoja_pedidos, carpeta),
                "forzar": forzar,
            },
            daemon=True,
        ).start()

    def finalizar_hoja_pedidos(self, carpeta):
        acc_color = self.color_principal
        ruta = ruta_hoja_pedidos_del_dia(carpeta)
        self.actualizar_estado("hoja de pedidos lista", str(ruta))
        self.actualizar_resumen()
        boton = getattr(self, "btn_hoja_pedidos", None)
        if boton is not None:
            boton.configure(text="Generar hoja de hoy", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        if ruta.exists():
            self.avisar(f"Hoja de pedidos de hoy lista: {ruta.name}", "ok", accion=("Abrir", lambda: self.abrir_archivo(ruta)))
            self.registrar_ejecucion("hoja", ruta.stem)

    def comprobar_hoja_pedidos_al_inicio(self):
        plantilla = self.config.get("ruta_plantilla_hojas", "")
        carpeta = self.config.get("ruta_carpeta_hojas", "")
        if not plantilla or not carpeta or not os.path.exists(plantilla) or not os.path.isdir(carpeta):
            return
        if ruta_hoja_pedidos_del_dia(carpeta).exists():
            return
        self.agregar_log("[HOJA PEDIDOS] No existe la hoja de hoy, generándola automáticamente...")
        threading.Thread(
            target=generar_hoja_pedidos_task,
            args=(plantilla, carpeta),
            kwargs={
                "log_callback": lambda t, tag=None: self.after(0, self.agregar_log, t, tag),
                "finalizado_callback": lambda: self.after(0, self.finalizar_hoja_pedidos, carpeta),
            },
            daemon=True,
        ).start()

    def log_callback_avisos(self, texto, tag=None):
        self.agregar_log(texto, tag)
        self.agregar_a_caja(self.avisos_log_box, texto, tag)

    def iniciar_avisos_no_suministro(self):
        global avisos_running
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        ruta_hoja = self.widget_get("hoja_avisos_entry")
        ruta_clientes = self.widget_get("clientes_avisos_entry")
        if not self.validar_archivo(ruta_hoja, "Selecciona la hoja de pedido (dd-mm-aaaa.xlsx) ya rellenada."):
            return
        if not self.validar_archivo(ruta_clientes, "Selecciona el Excel de clientes (clientes.xlsx)."):
            return
        envio_automatico = self.avisos_auto_var.get()
        emails_companeros = self.widget_get("email_companeros_avisos_entry")
        invalidos = correos_invalidos(emails_companeros)
        if invalidos:
            messagebox.showwarning("Aviso", "Estas direcciones de compañeros no parecen correos válidos:\n" + "\n".join(invalidos) + "\n\nSepara varias con punto y coma (;).")
            return
        if envio_automatico:
            mensaje = "Se generará un aviso por cada cliente con referencias en NO SUMINISTRA y otro por cada cliente con referencias SUSTITUIDAS. Los que tengan email registrado en el Excel de clientes se ENVIARÁN automáticamente; los que no tengan email se guardarán como borrador para completarlos a mano."
        else:
            mensaje = "Se creará un borrador en Outlook por cada cliente con referencias en NO SUMINISTRA y otro por cada cliente con referencias SUSTITUIDAS. Se guardan en Borradores sin enviarse."
        if emails_companeros:
            mensaje += f"\n\nAdemás se enviará la hoja de pedido completa a: {emails_companeros}."
        mensaje += "\n\n¿Continuar?"
        if not messagebox.askyesno("Generar avisos", mensaje):
            return
        self.actualizar_estado("generando avisos de no suministro")
        self.btn_avisos.configure(state="disabled")
        efectos.marcar_ocupado(self.btn_avisos, "Generando")
        self.limpiar_caja(self.avisos_log_box)
        avisos_running = True
        threading.Thread(
            target=generar_avisos_no_suministro_task,
            args=(ruta_hoja, ruta_clientes),
            kwargs={
                "enviar_automaticamente": envio_automatico,
                "emails_companeros": emails_companeros,
                "log_callback": lambda t, tag=None: self.after(0, self.log_callback_avisos, t, tag),
                "progreso_callback": lambda v: self.after(0, self.actualizar_progreso, v),
                "finalizado_callback": lambda: self.after(0, self.finalizar_avisos_no_suministro),
                "resultado_callback": lambda r: self.after(0, self.mostrar_resultado_avisos, r),
            },
            daemon=True,
        ).start()

    def finalizar_avisos_no_suministro(self):
        acc_color = self.color_principal
        self.actualizar_estado("avisos de no suministro procesados")
        self.btn_avisos.configure(text="Generar avisos de no suministro", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        self.avisar("Avisos de no suministro procesados: revisa el registro.")
        self.registrar_ejecucion("avisos", Path(self.widget_get("hoja_avisos_entry")).stem)

    def mostrar_resultado_avisos(self, resultado):
        if resultado.get("error"):
            return
        no_suministra = resultado.get("no_suministra", {})
        sustitucion = resultado.get("sustitucion", {})
        if not any(no_suministra.get(k) for k in ("enviados", "borradores", "sin_correo")) and \
           not any(sustitucion.get(k) for k in ("enviados", "borradores", "sin_correo")):
            return
        lineas = []

        def volcar_bloque(titulo, bucket):
            enviados = bucket.get("enviados", [])
            borradores = bucket.get("borradores", [])
            sin_correo = bucket.get("sin_correo", [])
            if not enviados and not borradores and not sin_correo:
                return
            if lineas:
                lineas.append("")
            lineas.append(f"== {titulo} ==")
            if enviados:
                lineas.append(f"Correos enviados: {len(enviados)}")
                for entrada in enviados:
                    lineas.append(f"  • {nombre_cliente_aviso(entrada)}: {', '.join(entrada['detalle'])}")
            if borradores:
                lineas.append(f"Borradores generados: {len(borradores)}")
                for entrada in borradores:
                    lineas.append(f"  • {nombre_cliente_aviso(entrada)}: {', '.join(entrada['detalle'])}")
            if sin_correo:
                lineas.append(f"Sin email en el Excel de clientes ({len(sin_correo)}), borrador creado sin destinatario:")
                for entrada in sin_correo:
                    lineas.append(f"  • Cliente {entrada['cliente']}: {', '.join(entrada['detalle'])}")

        volcar_bloque("No suministra", no_suministra)
        volcar_bloque("Sustituidas", sustitucion)
        messagebox.showinfo("Avisos de no suministro", "\n".join(lineas))

    def toggle_comparador(self):
        global comparador_running
        if not comparador_running:
            self.iniciar_comparador()
        else:
            stop_event_comparador.set()

    def iniciar_comparador(self):
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        ruta_ant = self.excel_old_entry.get()
        ruta_nue = self.excel_new_entry.get()
        folder = self.folder_comp_entry.get()
        color = self.color_optionmenu.get()
        colores_hex = {"Amarillo": "FFFF00", "Rojo": "FF0000", "Verde": "00FF00", "Azul": "0000FF", "Naranja": "FFA500"}
        if not self.validar_archivo(ruta_ant, "Selecciona el Excel antiguo."):
            return
        if not self.validar_archivo(ruta_nue, "Selecciona el Excel nuevo."):
            return
        if not self.validar_carpeta(folder, "Selecciona una carpeta para el resultado comparado."):
            return
        ruta_salida = str(Path(folder) / "resultado_comparacion.xlsx")
        if self.excel_abierto(ruta_salida, "El resultado de la comparación anterior"):
            return
        self.actualizar_estado("comparando excels", ruta_salida)
        self.btn_comparador.configure(fg_color=self.color_peligro, hover_color=self.color_peligro_hover)
        efectos.marcar_ocupado(self.btn_comparador, "Detener")
        stop_event_comparador.clear()
        threading.Thread(target=comparar_excel_task, args=(ruta_ant, ruta_nue, ruta_salida, colores_hex[color], lambda t, tag=None: self.after(0, self.agregar_log, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_comparador), stop_event_comparador), daemon=True).start()

    def finalizar_comparador(self):
        acc_color = self.color_principal
        self.actualizar_estado("comparación finalizada", str(Path(self.folder_comp_entry.get()) / "resultado_comparacion.xlsx"))
        self.btn_comparador.configure(text="Comparar Excels", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        resultado = Path(self.folder_comp_entry.get()) / "resultado_comparacion.xlsx"
        self.avisar("Comparación terminada: resultado_comparacion.xlsx", accion=("Abrir", lambda: self.abrir_archivo(resultado)) if resultado.exists() else None)
        self.registrar_ejecucion("comparador")

    def iniciar_filtro_rotacion(self):
        global filtro_rotacion_running
        if self.hay_tarea_activa():
            return False
        self.guardar_configuracion_ui_internal()
        ruta_quiter = self.excel_quiter_entry.get()
        ruta_rot = self.excel_rot_entry.get()
        folder_salida = self.folder_rot_entry.get()
        try:
            umbral = float(self.umbral_entry.get())
        except Exception:
            messagebox.showwarning("Aviso", "El umbral de rotación debe ser un número, por ejemplo 0.12 (usa punto como separador decimal).")
            return False
        try:
            precio_minimo = float(self.precio_entry.get().replace(",", "."))
        except Exception:
            messagebox.showwarning("Aviso", "Introduce un precio mínimo válido en la pestaña Automatización (se usa también en el filtro de rotación).")
            return False
        limite_revisar = self.limite_revisar()
        if limite_revisar is None:
            messagebox.showwarning("Aviso", "El importe para revisar piezas caras debe ser un número, por ejemplo 1600 (o 0 para no marcarlas).")
            return False
        if not self.validar_archivo(ruta_quiter, "Selecciona el Excel de Quiter."):
            return False
        if not self.validar_archivo(ruta_rot, "Selecciona el archivo de propuesta."):
            return False
        if not self.validar_carpeta(folder_salida, "Selecciona una carpeta de salida para rotación."):
            return False
        ruta_estimacion = str(Path(folder_salida) / f"{Path(ruta_quiter).stem}_ROTACION.xlsx")
        if self.excel_abierto(ruta_estimacion, "El Excel de salida de rotación"):
            return False
        self.actualizar_estado("analizando rotación", ruta_estimacion)
        self.btn_rotacion.configure(state="disabled")
        efectos.marcar_ocupado(self.btn_rotacion, "Analizando")
        self.limpiar_caja(self.rotacion_log_box)
        filtro_rotacion_running = True
        self._inicio_rotacion = time.time()
        threading.Thread(target=filtrar_por_rotacion_task, args=(ruta_quiter, ruta_rot, folder_salida, umbral, "FF0000", precio_minimo, lambda t, tag=None: self.after(0, self.log_callback_rotacion, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_filtro_rotacion)), kwargs={"limite_revisar": limite_revisar}, daemon=True).start()
        return True

    def limite_revisar(self):
        """Importe unitario a partir del cual una pieza sale en azul (0 = no se marcan); None si no es un número."""
        texto = str(self.valor_widget_o_config("limite_revisar_entry", "limite_revisar_rotacion", "1600") or "").strip()
        if not texto:
            return 0.0
        texto = texto.replace("€", "").replace(" ", "")
        if "," in texto or re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):  # «1.600» o «1.600,50»: punto de miles
            texto = texto.replace(".", "").replace(",", ".")
        try:
            return max(0.0, float(texto))
        except ValueError:
            return None

    def finalizar_filtro_rotacion(self):
        del_flujo = getattr(self, "_rotacion_del_flujo", False)
        self._rotacion_del_flujo = False
        ruta_quiter = self.excel_quiter_entry.get()
        generado = None
        if self.folder_rot_entry.get() and ruta_quiter:
            ruta_final = str(Path(self.folder_rot_entry.get()) / f"{Path(ruta_quiter).stem}_ROTACION.xlsx")
            self.actualizar_estado("rotación completada", ruta_final)
            # Solo cuenta como hecho si el Excel se ha escrito en esta ejecución
            if os.path.isfile(ruta_final) and os.path.getmtime(ruta_final) >= getattr(self, "_inicio_rotacion", 0) - 2:
                generado = ruta_final
                self.log_callback_rotacion("Pulsa «Copiar lista para Quiter» para pegar en Claude las líneas a poner a 0.")
        self.btn_rotacion.configure(state="normal", text="Resaltar por rotación")
        self.btn_flujo.configure(state="normal")
        if generado is None:
            self.avisar("El filtro de rotación terminó sin generar el Excel: revisa el registro de rotación.", "aviso")
            return
        self.avisar(f"Rotación completada: {Path(generado).name}", "ok", accion=("Abrir", lambda: self.abrir_archivo(generado)))

        def continuar(aplicado=True):
            try:
                marcadas, total = lineas_marcadas_rotacion(generado)
            except Exception:
                marcadas, total = [], 0
            pedido = Path(generado).stem.removesuffix("_ROTACION")
            self.registrar_ejecucion("rotacion", f"{pedido}: {len(marcadas)} de {total} líneas" if total else pedido)
            if del_flujo:
                self.mostrar_captura(len(marcadas), total, pedido)

        try:
            caras = piezas_caras_rotacion(generado)
        except Exception:
            caras = []
        if caras:
            self.decidir_piezas_caras(generado, caras, al_terminar=continuar)
        else:
            continuar()

    # ---------- piezas caras: pedirlas o ponerlas a 0 (ventana en piezas_caras.py) ----------

    def decidir_piezas_caras(self, ruta, piezas, al_terminar=None):
        """Pregunta qué hacer con cada pieza en azul y lo guarda en el Excel (blanco o rojo)."""
        nombre = Path(ruta).name

        def al_decidir(decisiones):
            if decisiones is None:
                quedan = len(piezas)
                self.avisar(f"Quedan {quedan} pieza{'s' if quedan != 1 else ''} cara{'s' if quedan != 1 else ''} en azul en {nombre}: decídelas antes de pasarlo a Quiter.",
                            accion=("Decidir", lambda: self.decidir_piezas_caras_pendientes(ruta)))
                if al_terminar:
                    al_terminar(False)
                return True
            try:
                pedir, cero = decidir_piezas_caras_en_excel(ruta, decisiones)
            except PermissionError:
                messagebox.showwarning("Piezas caras", f"No se puede guardar {nombre}: está abierto en Excel. Ciérralo y pulsa «Aplicar» otra vez.")
                return False
            except Exception as e:
                registrar_error_detallado(f"Piezas caras: {e}", traceback.format_exc())
                messagebox.showwarning("Piezas caras", f"No se pudo guardar la decisión en {nombre}: {e}")
                return False
            self.log_callback_rotacion(f"Piezas caras decididas: {pedir} para pedir (en blanco) y {cero} a 0 (en rojo).", "success")
            self.avisar(f"Piezas caras: {pedir} para pedir y {cero} a 0. {nombre} actualizado.", "ok")
            if al_terminar:
                al_terminar(True)
            return True

        self.dialogo_piezas_caras = piezas_caras.DialogoPiezasCaras(self, piezas, Path(ruta).stem.removesuffix("_ROTACION"), self.limite_revisar() or 0, al_decidir)
        if os.path.exists(ICONO_PATH):  # customtkinter pone su icono a los 200 ms: se cambia después
            ventana = self.dialogo_piezas_caras.ventana
            ventana.after(260, lambda: ventana.winfo_exists() and ventana.iconbitmap(ICONO_PATH))
        return self.dialogo_piezas_caras

    def decidir_piezas_caras_pendientes(self, ruta):
        try:
            piezas = piezas_caras_rotacion(ruta)
        except Exception as e:
            self.avisar(f"No se pudo leer {Path(ruta).name}: {e}", "aviso")
            return None
        if not piezas:
            self.avisar("No quedan piezas caras por decidir.")
            return None
        return self.decidir_piezas_caras(ruta, piezas)

    # ---------- vista previa del umbral y lista para Quiter ----------

    def iniciar_vista_previa_umbral(self):
        """Revisa cada poco si han cambiado el umbral, el precio mínimo o los Excel, y recalcula la vista previa."""
        if not hasattr(self, "_cache_rotacion"):
            self._cache_rotacion = {}
            self._cargando_vista_previa = False
        self._generacion_vista_previa = getattr(self, "_generacion_vista_previa", 0) + 1
        self._firma_vista_previa = None
        self.revisar_vista_previa_umbral(self._generacion_vista_previa)

    def revisar_vista_previa_umbral(self, generacion):
        etiqueta = getattr(self, "vista_previa_umbral", None)
        if generacion != self._generacion_vista_previa or etiqueta is None or not etiqueta.winfo_exists():
            return
        try:
            self.actualizar_vista_previa_umbral()
        except Exception as e:
            registrar_error_detallado(f"Vista previa del umbral: {e}", traceback.format_exc())
        self.after(700, self.revisar_vista_previa_umbral, generacion)

    def poner_vista_previa(self, texto):
        etiqueta = getattr(self, "vista_previa_umbral", None)
        if etiqueta is not None and etiqueta.winfo_exists() and etiqueta.cget("text") != texto:
            etiqueta.configure(text=texto)

    def datos_rotacion_cacheados(self, clave, ruta, fecha):
        guardado = self._cache_rotacion.get(clave)
        if guardado and guardado[0] == ruta and guardado[1] == fecha:
            return guardado[2]
        return None

    def cargar_datos_rotacion(self, ruta_quiter, fecha_quiter, ruta_propuesta, fecha_propuesta):
        """Lee los Excel que no estén ya en caché; un error se guarda para no reintentarlo hasta que cambie el archivo."""
        try:
            for clave, ruta, fecha, lector in (("quiter", ruta_quiter, fecha_quiter, leer_filas_quiter),
                                               ("propuesta", ruta_propuesta, fecha_propuesta, leer_medias_propuesta)):
                if self.datos_rotacion_cacheados(clave, ruta, fecha) is None:
                    try:
                        datos = lector(ruta)
                        if datos is None:
                            datos = ("sin_columnas",)  # propuesta sin columnas de artículo y media
                    except Exception as e:
                        datos = ("error", str(e))
                    self._cache_rotacion[clave] = (ruta, fecha, datos)
        finally:
            self._cargando_vista_previa = False

    def actualizar_vista_previa_umbral(self, en_segundo_plano=True):
        ruta_quiter = self.widget_get("excel_quiter_entry").strip()
        ruta_propuesta = self.widget_get("excel_rot_entry").strip()
        texto_umbral = self.widget_get("umbral_entry").strip()
        texto_precio = self.widget_get("precio_entry", "0").strip()
        limite_revisar = self.limite_revisar()
        fecha_quiter = os.path.getmtime(ruta_quiter) if os.path.isfile(ruta_quiter) else None
        fecha_propuesta = os.path.getmtime(ruta_propuesta) if os.path.isfile(ruta_propuesta) else None
        firma = (ruta_quiter, fecha_quiter, ruta_propuesta, fecha_propuesta, texto_umbral, texto_precio, limite_revisar)
        if firma == self._firma_vista_previa:
            return
        if fecha_quiter is None or fecha_propuesta is None:
            self._firma_vista_previa = firma
            self.poner_vista_previa("Elige el Excel de Quiter y la propuesta para ver cuántas filas se marcarían.")
            return
        try:
            umbral = float(texto_umbral)  # igual que el filtro: con punto decimal
            precio_minimo = float(texto_precio.replace(",", "."))
            if limite_revisar is None:
                raise ValueError("límite para revisar")
        except ValueError:
            self._firma_vista_previa = firma
            self.poner_vista_previa("Escribe el umbral con punto decimal (por ejemplo 0.12), un precio mínimo y un importe para revisar válidos para ver la vista previa.")
            return
        filas = self.datos_rotacion_cacheados("quiter", ruta_quiter, fecha_quiter)
        medias = self.datos_rotacion_cacheados("propuesta", ruta_propuesta, fecha_propuesta)
        if filas is None or medias is None:
            if not en_segundo_plano:
                self.cargar_datos_rotacion(ruta_quiter, fecha_quiter, ruta_propuesta, fecha_propuesta)
                return self.actualizar_vista_previa_umbral(en_segundo_plano=False)
            if not self._cargando_vista_previa:
                self._cargando_vista_previa = True
                self.poner_vista_previa("Calculando cuántas filas se marcarían...")
                threading.Thread(target=self.cargar_datos_rotacion, args=(ruta_quiter, fecha_quiter, ruta_propuesta, fecha_propuesta), daemon=True).start()
            return  # la siguiente revisión mostrará el resultado
        self._firma_vista_previa = firma
        if isinstance(filas, tuple):
            self.poner_vista_previa("No se puede leer el Excel de Quiter para la vista previa (puede estar abierto o guardándose).")
        elif medias == ("sin_columnas",):
            self.poner_vista_previa("La propuesta no tiene las columnas de artículo y media: el filtro no podrá ejecutarse.")
        elif isinstance(medias, tuple):
            self.poner_vista_previa("No se puede leer el archivo de propuesta para la vista previa.")
        else:
            conteo = contar_rotacion(filas, medias, umbral, precio_minimo, limite_revisar)
            azules = f"{conteo['revisar']} en azul para revisar · " if limite_revisar else ""
            texto = (f"Con {texto_umbral}: {conteo['rojo']} en rojo · {conteo['amarillo']} en amarillo · {azules}"
                     f"{conteo['precio']} sin marcar por precio · {conteo['blanco']} en blanco")
            if not any(fila[2] for fila in filas):
                texto += " (el amarillo aparece tras la automatización)"
            self.poner_vista_previa(texto)

    def ruta_salida_rotacion(self):
        ruta_quiter = self.widget_get("excel_quiter_entry").strip()
        carpeta = self.widget_get("folder_rot_entry").strip()
        if not ruta_quiter or not carpeta:
            return None
        return str(Path(carpeta) / f"{Path(ruta_quiter).stem}_ROTACION.xlsx")

    def copiar_lista_quiter(self):
        """Copia al portapapeles las líneas en rojo y amarillo del último Excel de rotación, listas para la skill de Quiter."""
        ruta = self.ruta_salida_rotacion()
        if not ruta or not os.path.isfile(ruta):
            self.avisar("Todavía no hay Excel de rotación para este Excel de Quiter. Ejecuta primero «Resaltar por rotación».")
            return None
        try:
            caras = piezas_caras_rotacion(ruta)
            marcadas, total = lineas_marcadas_rotacion(ruta)
        except Exception as e:
            messagebox.showwarning("Lista para Quiter", f"No se pudo leer {Path(ruta).name}: {e}")
            return None
        if caras:  # a Quiter no llega nada en azul: primero se decide
            self.decidir_piezas_caras(ruta, caras, al_terminar=lambda aplicado: self.copiar_lista_quiter() if aplicado else
                                      self.avisar("Decide primero las piezas caras para copiar la lista."))
            return None
        pedido = Path(ruta).stem.removesuffix("_ROTACION")
        texto = texto_lista_quiter(pedido, marcadas, total)
        self.clipboard_clear()
        self.clipboard_append(texto)
        self.update_idletasks()
        generado = datetime.fromtimestamp(os.path.getmtime(ruta)).strftime("%d/%m %H:%M")
        self.log_callback_rotacion(f"Lista copiada: {len(marcadas)} de {total} líneas de {Path(ruta).name} (generado el {generado}). Pégala en Claude.", "success")
        self.avisar(f"Lista copiada: {len(marcadas)} de {total} líneas a poner a 0. Pégala en Claude.", "ok")
        return texto

    def toggle_mopar(self):
        global mopar_running
        if mopar_running:
            stop_event_mopar.set()
            self.btn_mopar.configure(state="disabled")
            efectos.marcar_ocupado(self.btn_mopar, "Deteniendo")
        else:
            self.iniciar_mopar()

    def iniciar_mopar(self):
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        ruta = self.excel_mopar_entry.get()
        folder = self.folder_mopar_entry.get()
        if not self.validar_archivo(ruta, "Por favor selecciona un archivo Excel válido."):
            return
        if not self.validar_carpeta(folder, "Por favor selecciona una carpeta de destino."):
            return
        self.actualizar_estado("generando fichero Mopar", folder)
        if not os.path.exists(ruta):
            messagebox.showwarning("Aviso", "Por favor selecciona un archivo Excel válido.")
            return
        if not os.path.exists(folder):
            messagebox.showwarning("Aviso", "Por favor selecciona una carpeta de destino.")
            return
        self.btn_mopar.configure(fg_color=self.color_peligro, hover_color=self.color_peligro_hover)
        efectos.marcar_ocupado(self.btn_mopar, "Detener")
        self.limpiar_mopar_generados()
        stop_event_mopar.clear()
        threading.Thread(target=generar_fichero_mopar_task, args=(ruta, folder, lambda t, tag=None: self.after(0, self.agregar_log, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_mopar), stop_event_mopar, lambda n: self.after(0, self.agregar_mopar_generado, n)), daemon=True).start()

    def toggle_stock_fiat(self):
        global stock_fiat_running
        if stock_fiat_running:
            stop_event_stock_fiat.set()
            self.btn_stock_fiat.configure(state="disabled")
            efectos.marcar_ocupado(self.btn_stock_fiat, "Deteniendo")
        else:
            self.iniciar_stock_fiat()

    def iniciar_stock_fiat(self):
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        ruta = self.excel_stock_fiat_entry.get()
        folder = self.folder_stock_fiat_entry.get()
        if not self.validar_archivo(ruta, "Por favor selecciona un archivo Excel válido para Stock FIAT."):
            return
        if not self.validar_carpeta(folder, "Por favor selecciona una carpeta de destino para Stock FIAT."):
            return
        self.actualizar_estado("generando Stock FIAT", folder)
        if not os.path.exists(ruta):
            messagebox.showwarning("Aviso", "Por favor selecciona un archivo Excel válido para Stock FIAT.")
            return
        if not os.path.exists(folder):
            messagebox.showwarning("Aviso", "Por favor selecciona una carpeta de destino para Stock FIAT.")
            return
        self.btn_stock_fiat.configure(fg_color=self.color_peligro, hover_color=self.color_peligro_hover)
        efectos.marcar_ocupado(self.btn_stock_fiat, "Detener")
        stop_event_stock_fiat.clear()
        threading.Thread(target=generar_fichero_stock_fiat_task, args=(ruta, folder, lambda t, tag=None: self.after(0, self.agregar_log, t, tag), lambda v: self.after(0, self.actualizar_progreso, v), lambda: self.after(0, self.finalizar_stock_fiat), stop_event_stock_fiat), daemon=True).start()

    def log_callback_madd(self, texto, tag=None):
        self.agregar_log(texto, tag)
        self.agregar_a_caja(self.madd_log_box, texto, tag)

    def iniciar_facturas_madd(self):
        global madd_running
        if self.hay_tarea_activa():
            return
        self.guardar_configuracion_ui_internal()
        usuario = self.widget_get("usuario_madd_entry")
        password = self.widget_get("password_madd_entry")
        carpeta = self.widget_get("carpeta_facturas_madd_entry")
        email_destino = self.widget_get("email_facturas_madd_entry")
        email_aviso = self.widget_get("email_aviso_fallo_madd_entry")
        filtro_dia = self.widget_get("filtro_dia_madd_option", "Ayer")
        if not usuario or not password:
            messagebox.showwarning("Aviso", "Completa el usuario y la contraseña de MADD antes de continuar.")
            return
        if not self.validar_carpeta(carpeta, "Selecciona la carpeta donde guardar las facturas."):
            return
        if not email_destino:
            messagebox.showwarning("Aviso", "Introduce el correo de destino de las facturas.")
            return
        invalidos = correos_invalidos(email_destino) + correos_invalidos(email_aviso)
        if invalidos:
            messagebox.showwarning("Aviso", "Estas direcciones no parecen correos válidos:\n" + "\n".join(invalidos) + "\n\nSepara varias con punto y coma (;).")
            return
        if not messagebox.askyesno("Ejecutar prueba manual", "Esto iniciará sesión en MADD, descargará las facturas del día seleccionado y ENVIARÁ un correo real con los adjuntos (o el aviso de 'sin facturas'). ¿Continuar?"):
            return
        self.actualizar_estado("descargando facturas MADD")
        self.btn_madd.configure(state="disabled")
        efectos.marcar_ocupado(self.btn_madd, "Ejecutando")
        self.limpiar_caja(self.madd_log_box)
        madd_running = True
        threading.Thread(
            target=tarea_facturas_madd,
            args=(usuario, password, carpeta, email_destino, email_aviso),
            kwargs={
                "filtro_dia": filtro_dia,
                "log_callback": lambda t, tag=None: self.after(0, self.log_callback_madd, t, tag),
                "finalizado_callback": lambda: self.after(0, self.finalizar_facturas_madd),
            },
            daemon=True,
        ).start()

    def finalizar_facturas_madd(self):
        acc_color = self.color_principal
        self.actualizar_estado("descarga de facturas MADD finalizada", self.widget_get("carpeta_facturas_madd_entry"))
        self.btn_madd.configure(text="Ejecutar ahora (prueba manual)", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        carpeta_facturas = self.widget_get("carpeta_facturas_madd_entry")
        self.avisar("Descarga de facturas MADD terminada: revisa el registro.", accion=("Carpeta", lambda: self.abrir_directorio(carpeta_facturas)))
        self.refrescar_ultima_vez("madd")
        self.actualizar_resumen()

    def activar_tarea_programada_madd(self):
        if not getattr(sys, "frozen", False):
            messagebox.showwarning("Aviso", "La tarea programada solo se puede activar desde el .exe compilado (dist/), no ejecutando el código fuente con python.")
            return
        if not messagebox.askyesno("Activar tarea programada", "Esto crea una tarea en el Programador de tareas de Windows que ejecutará la descarga de facturas MADD todos los días laborables a las 07:00, despertando el PC de la suspensión si hace falta (no de un apagado completo). ¿Confirmas que ya has probado la ejecución manual y ha funcionado bien?"):
            return
        ruta_exe = sys.executable
        carpeta_exe = os.path.dirname(ruta_exe)
        ps_script = (
            f"$accion = New-ScheduledTaskAction -Execute '{ruta_exe}' -Argument '--tarea-madd' -WorkingDirectory '{carpeta_exe}'; "
            "$disparador = New-ScheduledTaskTrigger -Weekly -DaysOfWeek Monday,Tuesday,Wednesday,Thursday,Friday -At 07:00; "
            "$config = New-ScheduledTaskSettingsSet -WakeToRun -StartWhenAvailable; "
            "Register-ScheduledTask -TaskName 'PalausaFacturasMADD' -Action $accion -Trigger $disparador -Settings $config -Force"
        )
        try:
            resultado = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps_script], capture_output=True, text=True)
            if resultado.returncode == 0:
                messagebox.showinfo(
                    "Éxito",
                    "Tarea programada activada: se ejecutará todos los días laborables a las 07:00, despertando el PC si está en suspensión (no si está apagado del todo).\n\n"
                    "Importante: comprueba también que en Windows está activado \"Permitir temporizadores de reactivación\" para tu plan de energía "
                    "(Panel de control > Opciones de energía > Cambiar configuración avanzada > Suspender > Permitir temporizadores de reactivación). "
                    "Sin eso, algunos equipos ignoran la reactivación aunque la tarea esté bien configurada. Te recomiendo probarlo una noche antes de confiar en ello del todo."
                )
            else:
                messagebox.showerror("Error", f"No se pudo crear la tarea programada:\n{resultado.stderr or resultado.stdout}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo crear la tarea programada: {e}")

    def desactivar_tarea_programada_madd(self):
        try:
            resultado = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", "Unregister-ScheduledTask -TaskName 'PalausaFacturasMADD' -Confirm:$false"], capture_output=True, text=True)
            if resultado.returncode == 0:
                self.avisar("Tarea programada desactivada.", "ok")
            else:
                self.avisar("No había ninguna tarea programada activa (o ya se había desactivado).")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo desactivar la tarea: {e}")

    def valor_widget_o_config(self, atributo, clave, defecto):
        """Valor actual del widget; si su pestaña aún no se ha construido (o se destruyó) se usa el valor guardado."""
        widget = getattr(self, atributo, None)
        if widget is not None:
            try:
                return widget.get()
            except Exception:
                pass
        return self.config.get(clave, defecto)

    def guardar_configuracion_ui_internal(self):
        config = {}
        for clave, atributo, defecto, *convertir in CAMPOS_CONFIG:
            valor = self.valor_widget_o_config(atributo, clave, defecto)
            config[clave] = convertir[0](valor) if convertir else valor
        variable_avisos = getattr(self, "avisos_auto_var", None)
        config["avisos_envio_automatico"] = bool(variable_avisos.get()) if variable_avisos is not None else self.config.get("avisos_envio_automatico", True)
        variable_animacion = getattr(self, "animacion_var", None)
        config["animacion_inicio"] = bool(variable_animacion.get()) if variable_animacion is not None else self.config.get("animacion_inicio", True)
        selector_tema = getattr(self, "tema_option", None)
        try:
            config["tema_color"] = selector_tema.get() if selector_tema is not None else self.tema_actual()
        except Exception:
            config["tema_color"] = self.tema_actual()
        selector_modo = getattr(self, "modo_option", None)
        if selector_modo is not None:
            config["modo_apariencia"] = "Dark" if selector_modo.get() == "Oscuro" else "Light"
        else:
            config["modo_apariencia"] = self.config.get("modo_apariencia", "Light")
        if hasattr(self, "reclamaciones_power_line"):
            config.update(self.reclamaciones_power_line.obtener_configuracion())
        # Se fusiona con lo que hay en disco para no pisar claves que escriben otros procesos (contadores de Mopar/Stock FIAT)
        completa = cargar_configuracion()
        completa.update(config)
        guardar_configuracion(completa)
        self.config.update(completa)

    def finalizar_mopar(self):
        acc_color = self.color_principal
        self.actualizar_estado("fichero Mopar generado", self.folder_mopar_entry.get())
        self.btn_mopar.configure(text="Generar fichero Mopar", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        carpeta_mopar = self.folder_mopar_entry.get()
        self.avisar("Fichero Mopar terminado: revisa el registro.", accion=("Carpeta", lambda: self.abrir_directorio(carpeta_mopar)))
        self.registrar_ejecucion("mopar")

    def finalizar_stock_fiat(self):
        acc_color = self.color_principal
        self.actualizar_estado("Stock FIAT generado", self.folder_stock_fiat_entry.get())
        self.btn_stock_fiat.configure(text="Generar Stock FIAT", fg_color=acc_color, hover_color=self.color_hover, state="normal")
        carpeta_fiat = self.folder_stock_fiat_entry.get()
        self.avisar("Stock FIAT terminado: revisa el registro.", accion=("Carpeta", lambda: self.abrir_directorio(carpeta_fiat)))
        self.registrar_ejecucion("stock_fiat")

def probar_navegador():
    """Comprueba que el Chromium incluido arranca (sirve para verificar un .exe recién compilado).
    El resultado se escribe en diagnostico_navegador.txt, junto al programa."""
    ruta = ruta_datos("diagnostico_navegador.txt", migrar=False)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--disable-gpu", "--disable-dev-shm-usage"])
            try:
                page = browser.new_page()
                page.set_content("<title>navegador-ok</title><p>hola</p>")
                resultado = f"OK {page.title()}"
            finally:
                browser.close()
    except Exception as e:
        resultado = f"ERROR {e}"
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(resultado + "\n")
    return resultado

if __name__ == "__main__":
    if "--tarea-madd" in sys.argv:
        sys.exit(ejecutar_tarea_madd_standalone())
    if "--probar-navegador" in sys.argv:
        probar_navegador()
        sys.exit(0)
    mutex_instancia = instancia_unica.tomar()  # se conserva hasta salir
    if mutex_instancia is None:
        instancia_unica.traer_al_frente()
        sys.exit(0)
    app = AppIntegrada()
    app.mostrar_animacion_inicio()
    app.comprobar_actualizaciones_al_inicio()
    app.after(2500, app.recordar_tareas)
    app.mainloop()
